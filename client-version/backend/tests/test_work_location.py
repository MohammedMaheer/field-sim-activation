"""Work-location gates use genuine scoped sessions; tests supply explicit GPS fixtures."""
import csv
import io
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta, timezone
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from sqlalchemy import delete, select

import test_client_scope as shared
from app.db import Agent, DB, Outlet, Permission, Role, Shift, User, now
from app.location_tracking import (
    FenceAssignment, LatestWorkLocation, LocationPolicy, Point, WorkFence, WorkLocation,
    WorkSession, begin_work, retain_locations,
)

client, database, login = shared.client, shared.database, shared.login
URL = "/api/location-tracking"


@pytest.fixture(autouse=True)
def clean_geo(database):
    def clear():
        with DB() as db:
            for model in (LatestWorkLocation, WorkLocation, WorkSession, FenceAssignment, WorkFence, LocationPolicy):
                db.execute(delete(model))
            db.commit()
    clear()
    yield
    clear()


def gps(**changes):
    return {"latitude": 25.08, "longitude": 55.141, "accuracy_m": 10,
            "captured_at": now().replace(tzinfo=timezone.utc).isoformat(),
            "device": "Explicit isolated test device", "mocked": False, **changes}


def configure(client, account="agent1", enabled=True):
    login(client, account)
    person = client.get("/api/auth/me").json()
    # /auth/me uses the account shape directly in this API.
    if "user" in person:
        person = person["user"]
    login(client)
    body = {"name": "Isolated QA worksite", "branch_id": person.get("branch_id"),
            "latitude": 25.08, "longitude": 55.141, "radius_m": 500, "enabled": True}
    response = client.post(URL + "/fences", json=body)
    assert response.status_code == 201, response.text
    if not body["branch_id"]:
        assert client.put(URL + "/assignments/" + person["id"], json={"fence_id": response.json()["id"]}).status_code == 200
    if enabled:
        assert client.patch(URL + "/settings", json={"enabled": True}).status_code == 200
    login(client, account)
    return person, response.json()


def ready_for_next_point():
    with DB() as db:
        for row in db.scalars(select(LatestWorkLocation)):
            row.received_at = now() - timedelta(seconds=3)
        db.commit()


@pytest.mark.parametrize("account", ["admin", "ops", "leader", "leader2", "cluster", "compliance", "inventory", "salesmanager", "tele", "welcome", "agent1", "agent2"])
def test_all_members_can_explicitly_track_and_end_own_work(client, account):
    login(client, account)
    before = client.get(URL + "/me")
    assert before.status_code == 200 and before.json()["enabled"] is False
    assert before.json()["work_session"] is None
    result = client.post(URL + "/start", json=gps())
    assert result.status_code == 200, result.text
    session = result.json()["work_session"]["id"]
    assert client.post(URL + "/end", json={"session_id": session}).json()["work_session"] is None
    assert client.post(URL + "/points", json={**gps(), "session_id": session}).status_code == 409
    with DB() as db:
        row = db.scalar(select(WorkLocation))
        assert row and '25.08' not in row.coordinates_encrypted


@pytest.mark.parametrize("changes,state", [({"latitude": 25.1}, "OUTSIDE"),
                                          ({"latitude": 25.08445, "accuracy_m": 30}, "UNCERTAIN"),
                                          ({"accuracy_m": 150}, "UNCERTAIN")])
def test_enabled_fence_blocks_outside_or_uncertain_start(client, changes, state):
    configure(client)
    response = client.post(URL + "/start", json=gps(**changes))
    assert response.status_code == 422
    assert "boundary" in response.json()["detail"]
    assert client.get(URL + "/me").json()["work_session"] is None
    with DB() as db:
        assert db.scalar(select(WorkLocation)) is None


@pytest.mark.parametrize("case", ["stale", "future", "mocked"])
def test_stale_future_and_mocked_locations_are_never_accepted(client, case):
    configure(client)
    changes = {"mocked": True} if case == "mocked" else {
        "captured_at": (now() + timedelta(minutes=-3 if case == "stale" else 3)).replace(tzinfo=timezone.utc).isoformat()
    }
    assert client.post(URL + "/start", json=gps(**changes)).status_code == 422


def test_enabled_unassigned_member_and_no_gps_cannot_capture(client):
    configure(client, "agent1")
    login(client, "leader2")
    assert client.post(URL + "/start", json=gps()).status_code == 422
    login(client, "agent1")
    person = client.get("/api/auth/me").json()
    response = client.post("/api/transactions", json={"agent_id": person["agent_id"], "operation_id": str(uuid4())})
    assert response.status_code == 403 and "Start work tracking" in response.json()["detail"]
    response = client.post("/api/sales-management/sales", json={"agent_id": person["agent_id"], "order_type": "NEW", "customer_name": "Synthetic QA", "plan_name": "Synthetic plan"})
    assert response.status_code == 403


def test_fresh_inside_capture_and_shift_outside_stop_gate(client):
    person, _ = configure(client)
    started = client.post(URL + "/start", json=gps()).json()
    assert started["can_capture"] is True and started["latest_location"]["state"] == "INSIDE"
    session = started["work_session"]["id"]
    with DB() as db:
        assert db.scalar(select(Shift).where(Shift.agent_id == person["agent_id"], Shift.ended_at.is_(None)))
    ready_for_next_point()
    outside = client.post(URL + "/points", json={**gps(latitude=25.1), "session_id": session})
    assert outside.status_code == 200 and outside.json()["can_capture"] is False
    legacy = client.post(f"/api/agents/{person['agent_id']}/shift", json={"action": "start"})
    assert legacy.status_code == 403
    # A clear image reaches the location gate before admission or Tesseract.
    from PIL import Image
    import base64
    data = io.BytesIO()
    Image.new("RGB", (640, 400), "white").save(data, "PNG")
    encoded = base64.b64encode(data.getvalue()).decode()
    for path, body in [("/api/kyc-captures/read-order", {"image_base64": encoded}),
                       ("/api/kyc-captures/ocr-jobs", {"operation_id": str(uuid4()), "stage": "ORDER", "image_base64": encoded})]:
        response = client.post(path, json=body)
        assert response.status_code == 403 and "outside" in response.json()["detail"].lower()
    stopped = client.post(URL + "/end", json={"session_id": session})
    assert stopped.status_code == 200 and stopped.json()["can_capture"] is False
    with DB() as db:
        assert not db.scalar(select(Shift).where(Shift.agent_id == person["agent_id"], Shift.ended_at.is_(None)))


def test_stale_latest_gps_fails_even_if_ui_still_claims_inside(client):
    configure(client)
    client.post(URL + "/start", json=gps())
    with DB() as db:
        latest = db.scalar(select(LatestWorkLocation))
        latest.captured_at = now() - timedelta(minutes=3)
        db.commit()
    result = client.get(URL + "/me").json()
    assert result["latest_location"]["state"] == "STALE" and result["can_capture"] is False


def test_device_binding_session_ownership_monotonic_and_replay(client):
    configure(client)
    first = gps()
    started = client.post(URL + "/start", json=first).json()
    session = started["work_session"]["id"]
    assert client.post(URL + "/points", json={**first, "session_id": session}).status_code == 200
    with DB() as db:
        assert len(list(db.scalars(select(WorkLocation)))) == 1
    assert client.post(URL + "/start", json=gps(device="Another browser")).status_code == 409
    assert client.post(URL + "/points", json={**gps(device="Another browser"), "session_id": session}).status_code == 409
    older = gps(captured_at=(now() - timedelta(seconds=5)).replace(tzinfo=timezone.utc).isoformat())
    assert client.post(URL + "/points", json={**older, "session_id": session}).status_code == 409
    login(client, "agent2")
    assert client.post(URL + "/points", json={**gps(), "session_id": session}).status_code == 409
    assert client.post(URL + "/end", json={"session_id": session}).status_code == 404


def test_concurrent_start_creates_one_active_session_and_one_exact_point(client):
    configure(client)
    person = client.get("/api/auth/me").json()
    point = Point(**gps())
    def start():
        with DB() as db:
            begin_work(db, db.get(User, person["id"]), point)
            db.commit()
    with ThreadPoolExecutor(max_workers=2) as executor:
        list(executor.map(lambda _: start(), range(2)))
    with DB() as db:
        assert len(list(db.scalars(select(WorkSession).where(WorkSession.user_id == person["id"], WorkSession.ended_at.is_(None))))) == 1
        assert len(list(db.scalars(select(WorkLocation)))) == 1


@pytest.mark.parametrize("account", ["agent1", "leader", "cluster", "compliance", "inventory", "salesmanager", "tele", "welcome"])
def test_only_authorized_management_can_change_policy_fences_assignments(client, account):
    login(client, account)
    assert client.patch(URL + "/settings", json={"enabled": False}).status_code == 403
    assert client.post(URL + "/fences", json={"name": "Forbidden", "latitude": 25, "longitude": 55, "radius_m": 100}).status_code == 403
    assert client.put(URL + "/assignments/anything", json={"fence_id": None}).status_code == 403


def test_permissions_revocation_retains_read_without_collect_and_maximum_roles(client):
    login(client, "leader")
    person = client.get("/api/auth/me").json()
    with DB() as db:
        grant = db.scalar(select(Permission).where(Permission.role_id == db.get(User, person["id"]).role_id, Permission.name == "location.self"))
        grant_id, role_id = grant.id, grant.role_id
        db.delete(grant)
        db.commit()
    try:
        assert client.get(URL + "/members").status_code == 200
        assert client.get(URL + "/me").json()["can_start"] is False
        assert client.post(URL + "/start", json=gps()).status_code == 403
    finally:
        with DB() as db:
            db.add(Permission(id=grant_id, role_id=role_id, name="location.self"))
            db.commit()


def test_reassignment_preserves_history_without_giving_new_leader_old_gps(client):
    person, _ = configure(client)
    client.post(URL + "/start", json=gps())
    with DB() as db:
        agent = db.get(Agent, person["agent_id"])
        old = agent.leader_id
        new = db.scalar(select(User).where(User.email == "leader2@relay.demo"))
        agent.leader_id = new.id
        db.commit()
    try:
        login(client, "leader2")
        member = next(row for row in client.get(URL + "/members").json()["items"] if row["user_id"] == person["id"])
        assert member["latest_location"] is None
        report = client.get(URL + "/export?format=csv&user_id=" + person["id"])
        assert len(list(csv.reader(io.StringIO(report.content.decode("utf-8-sig"))))) == 1
        login(client, "leader")
        assert person["id"] not in {row["user_id"] for row in client.get(URL + "/members").json()["items"]}
        login(client, "agent1")
        assert client.get(URL + "/me").json()["can_capture"] is False
        ready_for_next_point()
        session = client.get(URL + "/me").json()["work_session"]["id"]
        assert client.post(URL + "/points", json={**gps(), "session_id": session}).status_code == 200
        login(client, "leader2")
        member = next(row for row in client.get(URL + "/members").json()["items"] if row["user_id"] == person["id"])
        assert member["latest_location"]["state"] == "INSIDE"
    finally:
        with DB() as db:
            db.get(Agent, person["agent_id"]).leader_id = old
            db.commit()


def test_reports_scoped_record_fence_snapshots_and_map_filter(client):
    person, fence = configure(client)
    client.post(URL + "/start", json=gps())
    login(client)
    client.patch(URL + "/fences/" + fence["id"], json={**{key: value for key, value in fence.items() if key != "id"}, "name": "Updated worksite", "radius_m": 900})
    report = client.get(URL + "/export?format=xlsx&user_id=" + person["id"])
    assert report.status_code == 200
    values = list(load_workbook(io.BytesIO(report.content)).active.values)
    assert values[1][7] == "Isolated QA worksite" and values[1][12] == 500
    assert values[1][13] == 0
    assert "Boundary radius (m)" in values[0]
    empty = client.get(URL + "/export?format=csv&latitude=0&longitude=0&radius_m=500")
    assert len(list(csv.reader(io.StringIO(empty.content.decode("utf-8-sig"))))) == 1
    login(client, "agent2")
    report = client.get(URL + "/export?format=csv&user_id=" + person["id"])
    assert len(list(csv.reader(io.StringIO(report.content.decode("utf-8-sig"))))) == 1
    assert client.get(URL + "/export?latitude=25&radius_m=500").status_code == 422
    assert client.get(URL + "/export?from_date=2026-01-01&to_date=2026-04-01").status_code == 422


def test_boundary_alert_deduplicates_same_transition_and_remains_scoped(client):
    configure(client)
    session = client.post(URL + "/start", json=gps()).json()["work_session"]["id"]
    ready_for_next_point()
    client.post(URL + "/points", json={**gps(latitude=25.1), "session_id": session})
    first = [row for row in client.get("/api/notifications").json()["items"] if row["id"].startswith("work-location:")]
    assert len(first) == 1 and first[0]["mobile_path"] == "/work-location"
    ready_for_next_point()
    client.post(URL + "/points", json={**gps(latitude=25.1001), "session_id": session})
    second = [row for row in client.get("/api/notifications").json()["items"] if row["id"].startswith("work-location:")]
    assert second[0]["id"] == first[0]["id"]
    login(client, "agent2")
    assert not [row for row in client.get("/api/notifications").json()["items"] if row["id"].startswith("work-location:")]


def test_bounded_retention_does_not_remove_accounts_or_shifts(client):
    configure(client)
    client.post(URL + "/start", json=gps())
    with DB() as db:
        for row in db.scalars(select(WorkLocation)):
            row.created_at = now() - timedelta(days=400)
        for row in db.scalars(select(LatestWorkLocation)):
            row.received_at = now() - timedelta(days=400)
        db.commit()
    retain_locations()
    with DB() as db:
        assert db.scalar(select(WorkLocation)) is None
        assert db.scalar(select(LatestWorkLocation)) is None
        assert db.scalar(select(User)) and db.scalar(select(Shift))


@pytest.mark.parametrize("account", ["agent1", "tele", "welcome", "compliance", "inventory"])
def test_non_manager_roles_only_see_their_own_member_record(client, account):
    person = login(client, account)["user"]
    rows = client.get(URL + "/members").json()["items"]
    assert [row["user_id"] for row in rows] == [person["id"]]
    assert client.get(URL + "/members?branch_id=outside-authorized-scope").json()["items"] == []


@pytest.mark.parametrize("account", ["admin", "ops", "salesmanager", "cluster", "leader", "leader2"])
def test_manager_member_visibility_is_current_and_role_scoped(client, account):
    actor = login(client, account)["user"]
    rows = client.get(URL + "/members").json()["items"]
    with DB() as db:
        if account in {"admin", "ops", "salesmanager"}:
            inactive = set(db.scalars(select(Agent.user_id).where(Agent.employment_status != "ACTIVE")))
            expected = set(db.scalars(select(User.id))) - inactive
        elif account == "cluster":
            branch_users = set(db.scalars(select(User.id).where(User.branch_id == actor["branch_id"])))
            assigned_agents = set(db.scalars(select(Agent.user_id).join(Outlet, Outlet.id == Agent.outlet_id).where(Outlet.branch_id == actor["branch_id"], Agent.employment_status == "ACTIVE")))
            all_agents = set(db.scalars(select(Agent.user_id)))
            expected = (branch_users - all_agents) | assigned_agents | {actor["id"]}
        else:
            expected = set(db.scalars(select(Agent.user_id).where(Agent.leader_id == actor["id"], Agent.employment_status == "ACTIVE"))) | {actor["id"]}
    assert {row["user_id"] for row in rows} == expected


def test_old_branch_manager_and_new_branch_manager_cannot_see_reassigned_old_gps(client):
    person, _ = configure(client)
    client.post(URL + "/start", json=gps())
    with DB() as db:
        agent = db.get(Agent, person["agent_id"])
        original = agent.outlet_id
        previous_branch = db.get(Outlet, original).branch_id
        different = db.scalar(select(Outlet).where(Outlet.branch_id != previous_branch))
        agent.outlet_id = different.id
        manager = db.scalar(select(User).where(User.email == "cluster@relay.demo"))
        manager_branch = manager.branch_id
        manager.branch_id = different.branch_id
        db.commit()
    try:
        login(client, "cluster")
        member = next(row for row in client.get(URL + "/members").json()["items"] if row["user_id"] == person["id"])
        assert member["latest_location"] is None
        exported = client.get(URL + "/export?format=csv&user_id=" + person["id"])
        assert len(list(csv.reader(io.StringIO(exported.content.decode("utf-8-sig"))))) == 1
        login(client)
        exported = client.get(URL + "/export?format=xlsx&branch_id=" + previous_branch)
        assert len(list(load_workbook(io.BytesIO(exported.content)).active.values)) == 2
    finally:
        with DB() as db:
            db.get(Agent, person["agent_id"]).outlet_id = original
            db.get(User, manager.id).branch_id = manager_branch
            db.commit()


def test_overgranted_database_role_cannot_manage_locations(client):
    login(client, "agent1")
    with DB() as db:
        role = db.scalar(select(Role).where(Role.name == "Field Agent"))
        grant = Permission(role_id=role.id, name="location.manage")
        db.add(grant)
        db.commit()
        grant_id = grant.id
    try:
        assert client.patch(URL + "/settings", json={"enabled": False}).status_code == 403
        assert client.get(URL + "/me").json()["permissions"]["manage"] is False
    finally:
        with DB() as db:
            db.delete(db.get(Permission, grant_id))
            db.commit()
