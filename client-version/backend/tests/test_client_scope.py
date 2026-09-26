import os
import base64
import io
import tempfile
from pathlib import Path
import pytest
from fastapi.testclient import TestClient
from sqlalchemy import select
from openpyxl import Workbook

temp = tempfile.TemporaryDirectory(prefix="relay-client-tests-")
os.environ["DATABASE_URL"] = "sqlite:///" + str(Path(temp.name) / "client.db")
os.environ["DEMO_PASSWORD"] = "test-client-password"
from app.db import Base, engine, DB, Order, Audit
from app.main import app, limits
from app.seed import seed


@pytest.fixture(scope="module", autouse=True)
def database():
    Base.metadata.create_all(engine)
    seed()
    yield
    engine.dispose()


@pytest.fixture
def client():
    limits.clear()
    with TestClient(app) as c:
        yield c


def login(client, account="admin"):
    r = client.post(
        "/api/auth/login",
        json={"email": account + "@relay.demo", "password": "test-client-password", "native": True},
    )
    assert r.status_code == 200
    client.headers["Authorization"] = "Bearer " + r.json()["access_token"]
    return r.json()


@pytest.mark.parametrize(
    "account,count",
    [
        ("admin", 12),
        ("ops", 12),
        ("leader", 9),
        ("leader2", 3),
        ("cluster", 9),
        ("agent1", 1),
        ("compliance", 12),
        ("inventory", 12),
    ],
)
def test_roles_and_scope(client, account, count):
    data = login(client, account)
    assert "activation.write" not in data["user"]["permissions"]
    assert ("inventory.write" in data["user"]["permissions"]) == (
        account in {"admin", "ops", "inventory"}
    )
    assert len(client.get("/api/resources/agents").json()) == count
    assert client.get("/api/dashboard").status_code == 200


def test_proposal_tasks_are_scoped_audited_and_transition_once(client):
    login(client)
    agent = next(
        a for a in client.get("/api/resources/agents").json() if a["employee_id"] == "RLY-1041"
    )
    created = client.post(
        "/api/field-tasks",
        json={
            "agent_id": agent["id"],
            "title": "Verify outlet stock",
            "note": "Synthetic task",
            "due_date": "2026-09-24",
        },
    )
    assert created.status_code == 201, created.text
    task_id = created.json()["id"]
    login(client, "agent1")
    own = client.get("/api/field-tasks").json()
    assert own and all(t["agent_id"] == own[0]["agent_id"] for t in own)
    if agent["id"] == own[0]["agent_id"]:
        done = client.patch(f"/api/field-tasks/{task_id}", json={"status": "DONE"})
        assert done.status_code == 200 and done.json()["completed_at"]
        assert (
            client.patch(f"/api/field-tasks/{task_id}", json={"status": "OPEN"}).status_code == 409
        )
        statuses = [item["status"] for item in client.get("/api/field-tasks").json()]
        assert statuses == sorted(statuses, key=lambda status: status == "DONE")
    assert (
        client.post(
            "/api/field-tasks",
            json={"agent_id": agent["id"], "title": "Unauthorized", "due_date": "2026-09-24"},
        ).status_code
        == 403
    )
    login(client, "agent2")
    assert client.patch(f"/api/field-tasks/{task_id}", json={"status": "DONE"}).status_code == 404
    with DB() as db:
        assert db.scalar(select(Audit).where(Audit.entity == task_id))


def test_support_ticket_is_scoped_and_review_is_audited(client):
    login(client, "agent1")
    agent_id = client.get("/api/auth/me").json()["agent_id"]
    opened = client.post(
        "/api/support-tickets",
        json={
            "agent_id": agent_id,
            "subject": "Device sync delay",
            "message": "Synthetic field issue for testing",
        },
    )
    assert opened.status_code == 201, opened.text
    ticket_id = opened.json()["id"]
    assert opened.json()["status"] == "OPEN"
    assert (
        client.patch(
            f"/api/support-tickets/{ticket_id}", json={"status": "RESOLVED", "response": "Resolved"}
        ).status_code
        == 403
    )
    login(client, "agent2")
    assert all(row["id"] != ticket_id for row in client.get("/api/support-tickets").json())
    login(client, "leader2")
    assert (
        client.patch(
            f"/api/support-tickets/{ticket_id}", json={"status": "RESOLVED", "response": "Resolved"}
        ).status_code
        == 404
    )
    login(client, "ops")
    resolved = client.patch(
        f"/api/support-tickets/{ticket_id}",
        json={"status": "RESOLVED", "response": "Restarted sync process"},
    )
    assert resolved.status_code == 200 and resolved.json()["status"] == "RESOLVED"
    assert (
        client.patch(
            f"/api/support-tickets/{ticket_id}", json={"status": "RESOLVED", "response": "Again"}
        ).status_code
        == 409
    )
    with DB() as db:
        assert db.scalar(
            select(Audit).where(Audit.entity == ticket_id, Audit.action == "Support Ticket Updated")
        )


def test_proposal_incentive_import_validates_whole_batch_and_preserves_amount(client):
    login(client)
    agent = client.get("/api/resources/agents").json()[0]
    manual = client.post(
        "/api/incentives",
        json={
            "agent_id": agent["id"],
            "period": "2026-09",
            "amount": "125.50",
            "note": "Demo target",
        },
    )
    assert manual.status_code == 201 and manual.json()["amount"] == "125.50"
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(["employee_id", "period", "amount", "note"])
    sheet.append([agent["employee_id"], "2026-09", "30.25", "Synthetic example"])
    sheet.append(["INVALID-EMPLOYEE", "2026-09", "50.00", "Bad row"])
    output = io.BytesIO()
    workbook.save(output)
    payload = {
        "filename": "incentives.xlsx",
        "content_base64": base64.b64encode(output.getvalue()).decode(),
    }
    before = len(client.get("/api/incentives").json())
    invalid = client.post("/api/incentives/import", json=payload)
    assert invalid.status_code == 422 and len(client.get("/api/incentives").json()) == before
    sheet.cell(3, 1).value = agent["employee_id"]
    output = io.BytesIO()
    workbook.save(output)
    payload["content_base64"] = base64.b64encode(output.getvalue()).decode()
    imported = client.post("/api/incentives/import", json=payload)
    assert imported.status_code == 201 and imported.json()["count"] == 2
    retry = client.post("/api/incentives/import", json=payload)
    assert retry.status_code == 201 and retry.json()["already_imported"] is True
    assert len(client.get("/api/incentives").json()) == before + 2
    assert client.get("/api/incentives/export?period=2026-09").status_code == 200
    login(client, "agent1")
    own = client.get("/api/incentives").json()
    assert own and all(x["agent_id"] == own[0]["agent_id"] for x in own)
    assert (
        client.post(
            "/api/incentives", json={"agent_id": agent["id"], "period": "2026-09", "amount": "20"}
        ).status_code
        == 403
    )


def test_proposal_customer_lookup_and_stock_move(client):
    login(client)
    assert client.get("/api/resources/customers").status_code == 200
    agents = client.get("/api/resources/agents").json()
    sims = client.get("/api/resources/inventory").json()
    source = next(
        s for s in sims if s["status"] == "AVAILABLE" and s["agent_id"] == agents[0]["id"]
    )
    target = next(a for a in agents if a["id"] != source["agent_id"])
    moved = client.post(
        f"/api/inventory/{source['id']}/move",
        json={"status": "AVAILABLE", "agent_id": target["id"], "reason": "Demo stock transfer"},
    )
    assert moved.status_code == 200 and moved.json()["agent_id"] == target["id"]
    with DB() as db:
        assert db.scalar(
            select(Audit).where(Audit.entity == source["id"], Audit.action == "Stock Transferred")
        )


def test_agent_can_report_own_stock_but_not_reassign_or_change_others(client):
    login(client, "agent1")
    own_agent = client.get("/api/auth/me").json()["agent_id"]
    own_sim = next(
        s
        for s in client.get("/api/resources/inventory").json()
        if s["agent_id"] == own_agent and s["status"] == "AVAILABLE"
    )
    login(client, "agent2")
    assert (
        client.post(
            f"/api/inventory/{own_sim['id']}/move",
            json={"status": "DAMAGED", "reason": "Broken in field"},
        ).status_code
        == 404
    )
    login(client, "agent1")
    assert (
        client.post(
            f"/api/inventory/{own_sim['id']}/move",
            json={"status": "WAREHOUSE", "reason": "Not allowed"},
        ).status_code
        == 403
    )
    moved = client.post(
        f"/api/inventory/{own_sim['id']}/move",
        json={"status": "DAMAGED", "reason": "Damaged connector"},
    )
    assert moved.status_code == 200 and moved.json()["status"] == "DAMAGED"
    assert (
        client.post(
            f"/api/inventory/{own_sim['id']}/move",
            json={"status": "RETURNED", "reason": "Another change"},
        ).status_code
        == 409
    )
    with DB() as db:
        assert db.scalar(
            select(Audit).where(Audit.entity == own_sim["id"], Audit.action == "Stock Transferred")
        )


@pytest.mark.parametrize(
    "path",
    [
        "/api/orders/draft",
        "/api/orders/x/submit",
        "/api/plans/x",
        "/api/roles",
        "/api/resources/plans",
        "/api/resources/orders",
    ],
)
def test_extra_capabilities_unavailable(client, path):
    login(client)
    for method in ["GET", "POST", "PATCH"]:
        assert client.request(method, path, json={}).status_code in {404, 405}


def test_unrelated_identity_simulation_is_not_in_proposal_api(client):
    login(client, "agent1")
    assert client.post("/api/ekyc/verify", json={}).status_code == 404
    assert "/api/ekyc/verify" not in client.get("/openapi.json").json()["paths"]


@pytest.mark.parametrize(
    "report",
    ["daily", "monthly", "agent", "team", "ekyc", "branch", "inventory", "failed", "audit"],
)
def test_csv_and_compliance_pdf_scope(client, report):
    login(client)
    assert client.get("/api/reports/" + report + "?format=csv").status_code == 200
    r = client.get("/api/reports/" + report + "?format=pdf")
    assert r.status_code == 200 and r.content.startswith(b"%PDF")


def test_readonly_records_and_api_contract(client):
    assert client.get("/api/dashboard").status_code == 401
    login(client)
    row = client.get("/api/resources/activations").json()[0]
    r = client.get("/api/activations/" + row["id"])
    assert r.status_code == 200 and r.json()["events"]
    paths = client.get("/openapi.json").json()["paths"]
    assert not any("/orders" in p or "/plans" in p or "/roles" in p for p in paths)


def test_location_collection_is_removed(client):
    data = login(client, "agent1")
    assert "location.write" not in data["user"]["permissions"]
    a = client.get("/api/resources/agents").json()[0]
    assert not {"lat", "lng", "accuracy", "geofence", "distance"} & a.keys()
    assert (
        client.post(
            f"/api/agents/{a['id']}/location", json={"lat": 25, "lng": 55, "accuracy": 10}
        ).status_code
        == 404
    )
    for resource in ("locations", "territories"):
        assert client.get("/api/resources/" + resource).status_code == 404
    assert client.get("/api/reports/geofence").status_code == 404
    assert not any(
        "/location" in p or "/territories" in p for p in client.get("/openapi.json").json()["paths"]
    )


def test_branch_filters_intersect_role_scope_and_charts_reconcile(client):
    login(client)
    branches = client.get("/api/resources/branches").json()
    assert len(branches) == 2
    total = 0
    for b in branches:
        d = client.get("/api/dashboard", params={"branch_id": b["id"]}).json()
        assert all(a["branch_id"] == b["id"] for a in d["agents"])
        assert all(t["branch_id"] == b["id"] for t in d["teams"])
        assert sum(r["value"] for r in d["capture_statuses"]) == d["capture_total"]
        assert sum(t["agents"] for t in d["teams"]) == len(d["agents"])
        total += len(d["agents"])
        csv = client.get("/api/reports/team", params={"branch_id": b["id"]}).text
        assert b["name"] in csv
    assert total == 12
    login(client, "agent1")
    own = client.get("/api/resources/agents").json()[0]
    foreign = next(b for b in branches if b["id"] != own["branch_id"])
    d = client.get("/api/dashboard", params={"branch_id": foreign["id"]}).json()
    assert d["agents"] == [] and d["capture_total"] == 0 and d["today"] == 0
    assert (
        client.get("/api/resources/team-leaders", params={"branch_id": foreign["id"]}).json() == []
    )


def test_multibyte_password_does_not_cause_server_error(client):
    r = client.post("/api/auth/login", json={"email": "admin@relay.demo", "password": "🔒" * 60})
    assert r.status_code == 401


def test_team_kpis_are_weighted_by_actual_records(client):
    from app.db import Ekyc, business_date

    login(client)
    teams = client.get("/api/resources/team-leaders").json()
    agents = client.get("/api/resources/agents").json()
    with DB() as db:
        for team in teams:
            members = {
                a["id"]
                for a in agents
                if a["leader_id"] == team["leader_id"] and a["branch_id"] == team["branch_id"]
            }
            checks = list(db.scalars(select(Ekyc).where(Ekyc.agent_id.in_(members))))
            orders = [
                o
                for o in db.scalars(
                    select(Order).where(Order.agent_id.in_(members), Order.status == "ACTIVATED")
                )
                if business_date(o.created_at) == business_date()
            ]
            assert team["aht"] == round(
                sum(o.handling_seconds for o in orders) / max(1, len(orders)) / 60, 1
            )
            assert team["ekyc_rate"] == round(
                sum(e.status == "VERIFIED" for e in checks) / max(1, len(checks)) * 100, 1
            )


def test_capture_lifecycle_access_and_excel(client, monkeypatch):
    import base64
    import io
    from PIL import Image
    from openpyxl import load_workbook
    from app import captures
    from app.db import KycCapture

    monkeypatch.setattr(
        captures.extractor,
        "extract",
        lambda _: {
            "engine": "Test",
            "lines": [{"text": "TXN-001 Jordan Demo Data Plan", "confidence": 96.2}],
            "rows": [],
            "history": [],
        },
    )
    login(client, "agent1")
    agent = client.get("/api/auth/me").json()["agent_id"]
    image = io.BytesIO()
    Image.new("RGB", (400, 200), "white").save(image, format="PNG")
    body = {
        "agent_id": agent,
        "operation_id": str(__import__("uuid").uuid4()),
        "source_reference": "DEMO-CAPTURE",
        "image_base64": base64.b64encode(image.getvalue()).decode(),
    }
    created = client.post("/api/kyc-captures", json=body)
    assert created.status_code == 201, created.text
    row = created.json()
    path = "/api/kyc-captures/" + row["id"]
    assert client.post("/api/kyc-captures", json=body).json()["id"] == row["id"]
    assert client.post(path + "/submit", json={"version": row["version"]}).status_code == 409
    captures.process_capture()
    row = client.get(path).json()
    assert row["status"] == "EXTRACTED"
    assert client.get(path + "/original").content == image.getvalue()
    with DB() as db:
        stored = db.get(KycCapture, row["id"])
        assert "TXN-001" not in stored.payload_encrypted
        assert stored.image_encrypted.startswith("gAAAA")
    items = [
        {
            "reference": "=HYPERLINK(1)",
            "customer": "Jordan Demo",
            "account": "0000123",
            "details": "Data plan",
            "source_line": 0,
        }
    ]
    assert (
        client.patch(
            path + "/rows",
            json={"version": row["version"], "rows": items + items, "reason": "Reviewed image"},
        ).status_code
        == 422
    )
    saved = client.patch(
        path + "/rows", json={"version": row["version"], "rows": items, "reason": "Reviewed image"}
    )
    assert saved.status_code == 200, saved.text
    assert (
        client.patch(
            path + "/rows",
            json={"version": row["version"], "rows": items, "reason": "Reviewed image"},
        ).status_code
        == 409
    )
    row = saved.json()
    sheet = load_workbook(io.BytesIO(client.get(path + "/excel").content)).active
    assert sheet["A2"].data_type == "s" and sheet["C2"].value == "0000123"
    row = client.post(path + "/submit", json={"version": row["version"]}).json()
    assert row["status"] == "SUBMITTED"
    assert (
        client.post(
            path + "/review",
            json={"version": row["version"], "outcome": "VERIFIED", "reason": "Checked original"},
        ).status_code
        == 403
    )
    login(client, "agent2")
    assert client.get(path).status_code == 404
    login(client, "inventory")
    assert client.get(path).status_code == 403
    login(client, "compliance")
    reviewed = client.post(
        path + "/review",
        json={
            "version": row["version"],
            "outcome": "VERIFIED",
            "reason": "Checked original screenshot",
        },
    )
    assert reviewed.status_code == 200, reviewed.text
    assert reviewed.json()["status"] == "VERIFIED"
    login(client, "agent1")
    assert client.get(path).json()["status"] == "VERIFIED"
    assert (
        client.post(
            "/api/kyc-captures",
            json={
                **body,
                "operation_id": str(__import__("uuid").uuid4()),
                "image_base64": base64.b64encode(b"not an image").decode(),
            },
        ).status_code
        == 422
    )


def test_capture_label_mapping_does_not_invent_fields():
    from app.capture_ocr import organize_lines

    lines = [
        {"text": text}
        for text in [
            "Transaction Reference: DEMO-01",
            "Customer: Jordan Demo",
            "Account: 0000123",
            "Plan: Essential 125",
            "Transaction ID: DEMO-02",
            "Customer: Avery Demo",
        ]
    ]
    rows = organize_lines(lines)
    assert len(rows) == 2
    assert rows[0]["account"] == "0000123" and rows[0]["customer"] == "Jordan Demo"
    assert rows[1]["reference"] == "DEMO-02" and rows[1]["account"] == ""
    assert organize_lines([{"text": "Unrecognized layout"}]) == []


def test_admin_management_validation_conflicts_and_permissions(client):
    login(client)
    a = next(
        row
        for row in client.get("/api/resources/agents").json()
        if row["employee_id"] == "RLY-1041"
    )
    path = f"/api/agents/{a['id']}/management"
    options = client.get(path).json()
    body = {
        "target": a["target"] + 1,
        "outlet_id": a["outlet_id"],
        "leader_id": a["leader_id"],
        "expected_target": a["target"],
        "expected_outlet_id": a["outlet_id"],
        "expected_leader_id": a["leader_id"],
        "reason": "Admin control audit test",
    }
    assert client.patch(path, json={**body, "target": 0}).status_code == 422
    wrong = next(leader for leader in options["leaders"] if leader["branch_id"] != a["branch_id"])
    assert client.patch(path, json={**body, "leader_id": wrong["id"]}).status_code == 422
    other = next(
        o
        for o in options["outlets"]
        if o["id"] != a["outlet_id"] and o["branch_id"] == a["branch_id"]
    )
    assert client.patch(path, json={**body, "outlet_id": other["id"]}).status_code == 409
    result = client.patch(path, json=body)
    assert result.status_code == 200, result.text
    assert result.json()["target"] == body["target"]
    roster = client.get("/api/resources/agents").json()
    assert [r["employee_id"] for r in roster] == sorted(r["employee_id"] for r in roster)
    assert client.patch(path, json=body).status_code == 409
    assert any(
        x["action"] == "Agent Assignment Changed" for x in client.get("/api/resources/audit").json()
    )
    assert (
        client.patch(
            path, json={**body, "expected_target": body["target"], "target": a["target"]}
        ).status_code
        == 200
    )
    for account in ["agent1", "leader", "ops", "compliance"]:
        login(client, account)
        assert client.get(path).status_code == 403
        assert client.patch(path, json=body).status_code == 403


def test_capture_filters_validate_and_preserve_scope(client):
    login(client)
    assert client.get("/api/kyc-captures?status=INVALID").status_code == 422
    assert client.get("/api/kyc-captures?offset=-1").status_code == 422
    all_rows = client.get("/api/kyc-captures?limit=100").json()
    first = client.get("/api/kyc-captures?limit=1").json()
    second = client.get("/api/kyc-captures?limit=1&offset=1").json()
    assert first == all_rows[:1] and second == all_rows[1:2]
    if all_rows:
        row = all_rows[0]
        match = client.get(
            "/api/kyc-captures", params={"search": row["source_reference"], "status": row["status"]}
        ).json()
        assert row in match
    assert client.get("/api/kyc-captures?search=NO-SUCH-CAPTURE-987654321").json() == []
    login(client, "agent2")
    own = client.get("/api/resources/agents").json()[0]["id"]
    rows = client.get("/api/kyc-captures?status=VERIFIED&limit=100").json()
    assert all(r["agent_id"] == own for r in rows)


def test_capture_history_reaches_beyond_fifty_without_leaking_other_agents(client):
    import uuid
    from sqlalchemy import delete
    from app.db import KycCapture
    from app.security import cipher

    login(client)
    agents = client.get("/api/resources/agents").json()
    own = next(a for a in agents if a["employee_id"] == "RLY-1041")
    other = next(a for a in agents if a["employee_id"] == "RLY-1042")
    ids = []
    with DB() as db:
        for i in range(62):
            row = KycCapture(
                agent_id=own["id"] if i < 61 else other["id"],
                creator_id=own["user_id"],
                operation_id=str(uuid.uuid4()),
                source_reference=f"AUDIT-PAGE-{i:02}",
                status="SUBMITTED",
                image_type="image/png",
                image_hash=str(i).zfill(64),
                image_encrypted=cipher.encrypt(b"synthetic").decode(),
            )
            db.add(row)
            db.flush()
            ids.append(row.id)
        db.commit()
    try:
        pages = [
            client.get(f"/api/kyc-captures?search=AUDIT-PAGE&limit=20&offset={offset}").json()
            for offset in (0, 20, 40, 60)
        ]
        assert [len(page) for page in pages] == [20, 20, 20, 2]
        assert len({r["id"] for page in pages for r in page}) == 62
        login(client, "agent1")
        rows = client.get("/api/kyc-captures?search=AUDIT-PAGE&limit=100").json()
        assert len(rows) == 61 and all(r["agent_id"] == own["id"] for r in rows)
    finally:
        with DB() as db:
            db.execute(delete(KycCapture).where(KycCapture.id.in_(ids)))
            db.commit()


def test_organization_setup_is_atomic_scoped_and_audited(client):
    login(client, "agent1")
    assert client.get("/api/organization").status_code == 403
    assert client.post("/api/organization/branches", json={"name": "Forbidden"}).status_code == 403
    login(client)
    branch = client.post("/api/organization/branches", json={"name": "Synthetic setup branch"})
    assert branch.status_code == 201, branch.text
    branch_id = branch.json()["id"]
    assert (
        client.post(
            "/api/organization/branches", json={"name": "SYNTHETIC SETUP BRANCH"}
        ).status_code
        == 409
    )
    assert any(b["id"] == branch_id for b in client.get("/api/resources/branches").json())
    team = client.post(
        "/api/organization/teams",
        json={
            "name": "Setup Leader",
            "branch_id": branch_id,
            "email": "setup.leader@relay.demo",
            "password": "test-client-password",
        },
    )
    assert team.status_code == 201, team.text
    leader_id = team.json()["id"]
    assert any(
        t["leader_id"] == leader_id and t["agents"] == 0
        for t in client.get("/api/resources/team-leaders").json()
    )
    outlet = client.post(
        "/api/organization/outlets", json={"name": "Setup outlet", "branch_id": branch_id}
    )
    assert outlet.status_code == 201
    payload = {
        "name": "Setup Agent",
        "branch_id": branch_id,
        "leader_id": leader_id,
        "outlet_id": outlet.json()["id"],
        "employee_id": "SETUP-01",
        "email": "setup.agent@relay.demo",
        "password": "test-client-password",
        "target": 25,
    }
    other = next(
        o for o in client.get("/api/organization").json()["outlets"] if o["branch_id"] != branch_id
    )
    assert (
        client.post(
            "/api/organization/agents", json={**payload, "outlet_id": other["id"]}
        ).status_code
        == 422
    )
    assert (
        client.post("/api/organization/agents", json={**payload, "password": "short"}).status_code
        == 422
    )
    created = client.post("/api/organization/agents", json=payload)
    assert created.status_code == 201, created.text
    assert client.post("/api/organization/agents", json=payload).status_code == 409
    with DB() as db:
        event = db.scalar(select(Audit).where(Audit.entity == created.json()["id"]))
        assert event is not None
        assert "password" not in str(event.new_value)
    login(client, "setup.agent")
    agents = client.get("/api/resources/agents").json()
    assert len(agents) == 1 and agents[0]["id"] == created.json()["id"]
    login(client, "setup.leader")
    assert len(client.get("/api/resources/agents").json()) == 1
    login(client, "leader")
    assert not any(
        a["id"] == created.json()["id"] for a in client.get("/api/resources/agents").json()
    )


def test_demo_refresh_is_opt_in_idempotent_and_preserves_history(monkeypatch):
    from app.demo_refresh import refresh_demo
    from app.db import Sim, Movement

    monkeypatch.delenv("RELAY_DEMO_REFRESH", raising=False)
    with pytest.raises(RuntimeError, match="explicit"):
        refresh_demo()
    with DB() as db:
        original = {o.id: (o.status, o.created_at) for o in db.scalars(select(Order)).all()}
    monkeypatch.setenv("RELAY_DEMO_REFRESH", "1")
    refresh_demo()
    with DB() as db:
        added = db.scalars(select(Order).where(Order.operation_id.like("demo-day-%"))).all()
        assert added
        added_ids = {o.id for o in added}
        for order in added:
            assert db.get(Sim, order.sim_id).agent_id == order.agent_id
            assert db.scalar(select(Movement.id).where(Movement.sim_id == order.sim_id))
            assert db.scalar(select(Audit.id).where(Audit.entity == order.id))
    refresh_demo()
    with DB() as db:
        assert {
            o.id for o in db.scalars(select(Order).where(Order.operation_id.like("demo-day-%")))
        } == added_ids
        assert all(
            (db.get(Order, key).status, db.get(Order, key).created_at) == value
            for key, value in original.items()
        )
