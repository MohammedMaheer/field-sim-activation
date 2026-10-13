"""Reject forged XLSX dimensions before expanding cells or writing any rows."""

# ruff: noqa: F401,F811
import base64
import io
import re
from uuid import uuid4
from zipfile import ZipFile

import pytest
from openpyxl import Workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet._read_only import ReadOnlyWorksheet
from sqlalchemy import func, select

from test_client_scope import client, database, login
from app.db import Branch, DB, FieldAsset, Role, User, business_date
from app.final_incentives import FinalIncentive, HEADERS as FINAL_HEADERS
from app.asset_import import HEADERS as ASSET_HEADERS
from app.organization_import import HEADERS as ORGANIZATION_HEADERS


def forged_sheet(headers, dimension):
    book = Workbook()
    book.active.append(headers)
    book.active.append(["Synthetic value"] * len(headers))
    original = io.BytesIO()
    book.save(original)
    output = io.BytesIO()
    with ZipFile(io.BytesIO(original.getvalue())) as source, ZipFile(output, "w") as target:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/worksheets/sheet1.xml":
                data, replacements = re.subn(
                    rb'<dimension ref="[^"]+"',
                    f'<dimension ref="{dimension}"'.encode(),
                    data,
                )
                assert replacements == 1
            target.writestr(entry, data)
    return {
        "filename": "synthetic-dimension.xlsx",
        "content_base64": base64.b64encode(output.getvalue()).decode(),
        "apply": True,
    }


@pytest.mark.parametrize(
    "endpoint,headers",
    [
        ("/api/incentives/final/import", FINAL_HEADERS),
        ("/api/field-assets/import", ASSET_HEADERS),
        *[
            (f"/api/organization/import/{kind}", headers)
            for kind, headers in ORGANIZATION_HEADERS.items()
        ],
    ],
)
@pytest.mark.parametrize("inflated", ["columns", "rows"])
def test_inflated_workbook_is_rejected_without_materializing_or_writing(
    client, monkeypatch, endpoint, headers, inflated
):
    login(client)
    dimension = (
        "A1:XFD2" if inflated == "columns" else f"A1:{get_column_letter(len(headers))}1048576"
    )
    body = forged_sheet(headers, dimension)
    models = (Branch, User, FieldAsset, FinalIncentive)
    with DB() as db:
        before = [db.scalar(select(func.count()).select_from(model)) for model in models]
        if endpoint == "/api/field-assets/import":
            body["branch_id"] = db.scalar(select(Branch.id).limit(1))
    iterations = []
    original = ReadOnlyWorksheet.iter_rows

    def observed(self, *args, **kwargs):
        iterations.append((args, kwargs))
        return original(self, *args, **kwargs)

    monkeypatch.setattr(ReadOnlyWorksheet, "iter_rows", observed)
    response = client.post(endpoint, json=body)
    assert response.status_code == 422, response.text
    assert iterations == [], "Reject the dimension before any cell expansion"
    with DB() as db:
        assert [db.scalar(select(func.count()).select_from(model)) for model in models] == before


@pytest.mark.parametrize("kind", ["final", "assets", "branches", "teams", "agents"])
def test_ordinary_excel_template_remains_previewable(client, kind):
    login(client)
    unique = str(uuid4())[:8]
    with DB() as db:
        branch = db.scalar(select(Branch).limit(1))
        if kind == "final":
            endpoint, headers = "/api/incentives/final/import", FINAL_HEADERS
            row = ["agent1@relay.demo", business_date().isoformat(), 12.5, unique, "Approved"]
        elif kind == "assets":
            endpoint, headers = "/api/field-assets/import", ASSET_HEADERS
            row = ["UNIFORM", "Synthetic uniform", "", 2, ""]
        else:
            endpoint = f"/api/organization/import/{kind}"
            headers = ORGANIZATION_HEADERS[kind]
            if kind == "branches":
                row = [f"Synthetic branch {unique}"]
            elif kind == "teams":
                row = [
                    f"Synthetic leader {unique}",
                    f"{unique}@example.test",
                    "synthetic-password",
                    branch.name,
                ]
            else:
                leader = db.scalar(
                    select(User)
                    .join(Role)
                    .where(Role.name == "Team Leader", User.branch_id.is_not(None))
                    .limit(1)
                )
                branch = db.get(Branch, leader.branch_id)
                row = [
                    f"Synthetic agent {unique}",
                    f"{unique}@example.test",
                    "synthetic-password",
                    branch.name,
                    leader.email,
                    f"X-{unique}",
                    12,
                ]
        branch_id = branch.id
    book = Workbook()
    book.active.append(headers)
    book.active.append(row)
    stream = io.BytesIO()
    book.save(stream)
    body = {
        "filename": "ordinary-template.xlsx",
        "content_base64": base64.b64encode(stream.getvalue()).decode(),
    }
    if kind == "assets":
        body["branch_id"] = branch_id
    response = client.post(endpoint, json=body)
    assert response.status_code == 200, response.text
    assert response.json().get("new_rows", response.json().get("count")) == 1
