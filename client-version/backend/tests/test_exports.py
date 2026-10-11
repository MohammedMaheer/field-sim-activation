"""Download integrity, completeness, safe formatting and scope regressions."""
# Pytest discovers these imported fixtures; function parameters intentionally share their names.
# ruff: noqa: F811

import base64
import csv
import io
from datetime import datetime, timezone
from decimal import Decimal
from uuid import uuid4
from zipfile import ZipFile

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader

from app.exports import csv_value, filename, receipt_pdf_bytes, table_pdf_bytes, workbook_bytes
from test_client_scope import client, database, login, review_regression_cleanup  # noqa: F401


def test_workbook_keeps_identifiers_zero_amounts_typed_dates_and_safe_text():
    columns = ["sr_number", "employee_id", "monthly_cost", "prepayment", "created_at", "note"]
    rows = [["00000000123456789012", "0007", Decimal("350.50"), 0, datetime(2026, 10, 11, 8, tzinfo=timezone.utc), "  =HYPERLINK(1)\nSecond line"]]
    exported = workbook_bytes(columns, rows, "Sales records", metadata={"branch": "Marina Branch", "from_date": "2026-10-11"})
    book = load_workbook(io.BytesIO(exported))
    sheet = book.active
    assert sheet["A2"].value == "00000000123456789012" and sheet["A2"].number_format == "@"
    assert sheet["B2"].value == "0007"
    assert sheet["C2"].value == 350.5 and sheet["C2"].data_type == "n"
    assert sheet["D2"].value == 0 and sheet["D2"].data_type == "n"
    assert isinstance(sheet["E2"].value, datetime) and sheet["E2"].number_format == "dd mmm yyyy hh:mm"
    assert sheet["F2"].value.startswith("  =HYPERLINK") and sheet["F2"].data_type == "s"
    assert sheet["F2"].alignment.wrap_text and sheet.column_dimensions["F"].width >= 38
    assert sheet.freeze_panes == "A2" and sheet.auto_filter.ref == "A1:F2"
    assert book.properties.creator == "Qanawat Connect Solutions LLC"
    assert "Marina Branch" in str(list(book["Export details"].values))
    with ZipFile(io.BytesIO(exported)) as archive:
        assert any(name.startswith("xl/media/") for name in archive.namelist())


@pytest.mark.parametrize("text", ["=1+1", "+CMD()", "-CMD()", "@SUM(1)", " \t=HYPERLINK(1)", "\tcell", "\rformula"])
def test_csv_escapes_formula_vectors_in_every_text_column(text):
    assert csv_value(text).startswith("'")
    assert csv_value(-2) == -2
    assert csv_value(0) == 0


def test_filename_rejects_header_characters_and_has_brand_and_date():
    name = filename('Sale receipt"\r\nX-Evil: 1', "pdf", "SR/0001", datetime(2026, 10, 11))
    assert name.startswith("qanawat-sale-receipt") and name.endswith("2026-10-11.pdf")
    assert not any(value in name for value in ('"', "\n", "\r", "/"))


def test_pdf_keeps_every_column_and_complete_long_values():
    row = {"reference": "CHECK-EXPORT-001", **{f"detail_{index}": f"UNIQUEVALUE{index:02}" for index in range(17)}, "final_note": "A complete note with three lines\nsecond line stays visible\nFINAL-LINE-STAYS"}
    result = table_pdf_bytes("Complete sales report", [row], "From 11 Oct 2026 to 11 Oct 2026")
    document = PdfReader(io.BytesIO(result))
    text = "\n".join(page.extract_text() for page in document.pages)
    assert len(document.pages) >= 4
    for value in row.values():
        assert all(part in text for part in value.splitlines())
    assert text.count("CHECK-EXPORT-001") >= 4
    assert "Qanawat Connect Solutions LLC" in text and "Fields 4 of 4" in text


def test_receipt_pdf_is_compact_and_keeps_all_sections_and_missing_fields():
    sections = [{"title": title, "fields": [{"label": f"Field {index}", "value": f"VALUE-{title}-{index}" if index else "Not recorded"} for index in range(8)]} for title in ("Customer", "Order", "Backend review", "SR verification")]
    result = receipt_pdf_bytes("Sale receipt", sections, "Activated - pending SR verification")
    document = PdfReader(io.BytesIO(result))
    assert len(document.pages) == 1
    text = document.pages[0].extract_text()
    for section in sections:
        for field in section["fields"]:
            assert field["value"] in text
    assert "Activated - pending SR verification" in text


@pytest.mark.parametrize("report", ["daily", "monthly", "agent", "team", "branch", "ekyc", "inventory", "failed", "audit"])
def test_general_reports_offer_csv_xlsx_pdf_with_matching_filtered_rows(client, report):
    login(client)
    # An empty future date isolates the schema and no-record output from mutable fixtures.
    params = {"start": "2099-10-11", "end": "2099-10-11"}
    csv_result = client.get(f"/api/reports/{report}", params={**params, "format": "csv"})
    xlsx_result = client.get(f"/api/reports/{report}", params={**params, "format": "xlsx"})
    pdf_result = client.get(f"/api/reports/{report}", params={**params, "format": "pdf"})
    assert csv_result.status_code == xlsx_result.status_code == pdf_result.status_code == 200
    rows = list(csv.reader(io.StringIO(csv_result.content.decode("utf-8-sig"))))
    book = load_workbook(io.BytesIO(xlsx_result.content))
    assert book.active.max_row == len(rows)
    assert len(book.active[1]) == len(rows[0])
    assert "2099-10-11" in str(list(book["Export details"].values))
    assert PdfReader(io.BytesIO(pdf_result.content)).pages
    for response in (csv_result, xlsx_result, pdf_result):
        assert 'filename="qanawat-' in response.headers["content-disposition"]
        assert "no-store" in response.headers["cache-control"]


def test_sales_export_preserves_zero_charges_and_import_column_contract(client, review_regression_cleanup):
    login(client, "ops")
    agent = next(row for row in client.get("/api/resources/agents").json() if row["employment_status"] == "ACTIVE")
    reference = "EXPORT-ZERO-" + str(uuid4())
    sr_number = "0000000" + str(uuid4().int)[:15]
    created = client.post("/api/sales-management/sales", json={"agent_id": agent["id"], "order_type": "NEW", "customer_name": "Export Customer", "plan_name": "Recorded plan", "document_number": "784199912345671", "request_id": reference, "monthly_cost": "0", "prepayment": "0", "sr_number": sr_number})
    assert created.status_code == 201, created.text
    # Screenshot captures retain recorded charges in details; management entry has no charge inputs.
    from app.db import DB, SalesRecord
    with DB() as db:
        sale = db.get(SalesRecord, created.json()["id"])
        sale.details = {**sale.details, "monthly_cost": 0, "prepayment": 0}
        db.commit()
    for format in ("csv", "xlsx"):
        response = client.get("/api/sales-management/export", params={"format": format, "agent_id": agent["id"]})
        assert response.status_code == 200
        if format == "xlsx":
            records = list(load_workbook(io.BytesIO(response.content)).active.values)
        else:
            records = list(csv.reader(io.StringIO(response.content.decode("utf-8-sig"))))
        columns = list(records[0])
        assert columns[:3] == ["sale_id", "current_status", "new_status"]
        row = next(row for row in records[1:] if row[columns.index("request_id")] == reference)
        assert str(row[columns.index("monthly_cost")]) in {"0", "0.0"}
        assert str(row[columns.index("prepayment")]) in {"0", "0.0"}
        assert row[columns.index("sr_number")] == sr_number
    pdf = client.get("/api/sales-management/export", params={"format": "pdf", "agent_id": agent["id"]})
    assert pdf.status_code == 200
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf.content)).pages)
    assert reference in "".join(text.split()) and sr_number in "".join(text.split())
    assert "784199912345671" not in text
    assert "Monthlycharge(AED)" in "".join(text.split())
    login(client, "agent2")
    if client.get("/api/auth/me").json().get("agent_id") == agent["id"]:
        login(client, "agent1")
    forbidden = client.get("/api/sales-management/export", params={"format": "xlsx", "agent_id": agent["id"], "branch_id": agent["branch_id"]})
    assert forbidden.status_code == 200
    book = load_workbook(io.BytesIO(forbidden.content))
    assert book.active.max_row == 1 and reference not in str(list(book.active.values))
    assert agent["name"] not in str(list(book["Export details"].values))


def test_branded_import_templates_keep_roundtrip_headers_and_text_identifiers(client):
    login(client)
    for route, expected in [("/api/inventory/bulk-template", ["ICCID", "SIM Serial", "SIM Type", "Business Category"]), ("/api/sales-management/sr/template?business_date=2026-10-11", ["sr_number", "status", "request_id"]), ("/api/sales-management/targets/template?period=2099-10", ["employee_id", "period", "order_type", "current_daily", "current_monthly", "daily_target", "monthly_target"])]:
        response = client.get(route)
        assert response.status_code == 200
        book = load_workbook(io.BytesIO(response.content))
        assert list(next(book.active.values)) == expected
        assert book.active.freeze_panes == "A2" and "Export details" in book.sheetnames
        assert 'filename="qanawat-' in response.headers["content-disposition"]
        assert book.active["A2"].number_format == "@"
    response = client.get("/api/sales-management/sr/template")
    book = load_workbook(io.BytesIO(response.content))
    book.active["A2"], book.active["B2"], book.active["C2"] = "000012345", "CLOSED", "000006789"
    binary = io.BytesIO()
    book.save(binary)
    from app.sr_verification import ReportBody, read_report
    entries, errors, _ = read_report(ReportBody(filename="branded.xlsx", business_date="2026-10-11", content_base64=base64.b64encode(binary.getvalue()).decode()))
    assert not errors and entries["000012345"]["request_id"] == "000006789"


@pytest.mark.parametrize("format", ["csv", "xlsx", "pdf"])
def test_commission_downloads_are_scoped_branded_and_multiformat(client, format):
    login(client)
    response = client.get("/api/incentives/export", params={"format": format})
    assert response.status_code == 200 and 'filename="qanawat-sales-commissions' in response.headers["content-disposition"]
    if format == "xlsx":
        book = load_workbook(io.BytesIO(response.content))
        assert "Sales agent" in list(next(book.active.values))
    elif format == "pdf":
        assert PdfReader(io.BytesIO(response.content)).pages
    login(client, "tele")
    assert client.get("/api/incentives/export", params={"format": format}).status_code == 403


@pytest.mark.parametrize("kind", ["summary", "movements", "stock", "requests"])
def test_stock_pdf_exports_keep_the_requested_columns_and_brand(client, kind):
    login(client)
    result = client.get("/api/field-assets/report/export", params={"kind": kind, "format": "pdf"})
    assert result.status_code == 200
    text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(result.content)).pages)
    assert "Qanawat Connect Solutions LLC" in text
    assert f'filename="qanawat-stock-{kind}-' in result.headers["content-disposition"]
    foreign = client.get("/api/field-assets/report/export", params={"kind": kind, "format": "xlsx", "branch_id": "outside-access"})
    assert foreign.status_code == 200 and load_workbook(io.BytesIO(foreign.content)).active.max_row == 1


@pytest.mark.parametrize("format", ["csv", "xlsx", "pdf"])
def test_calculated_commission_export_matches_authoritative_summary_and_keeps_hold_state(client, format):
    login(client, "agent1")
    summary = client.get("/api/commissions/summary", params={"period": "2099-10"})
    assert summary.status_code == 200
    expected = summary.json()["rows"]
    result = client.get("/api/commissions/export", params={"period": "2099-10", "format": format})
    assert result.status_code == 200 and 'filename="qanawat-commission-calculations-' in result.headers["content-disposition"]
    assert len(expected) == 1 and expected[0]["amount"] is None
    if format == "xlsx":
        book = load_workbook(io.BytesIO(result.content))
        assert book.active.max_row == 2
        values = list(book.active.values)
        assert values[1][0] == expected[0]["name"]
        assert values[1][values[0].index("Amount (AED)")] == "Awaiting configuration"
        assert "Calculation components" in book.sheetnames
        assert book["Calculation components"]["E2"].value == "Awaiting configuration"
    elif format == "csv":
        records = list(csv.reader(io.StringIO(result.content.decode("utf-8-sig"))))
        assert len(records) == 2 and "Awaiting configuration" in records[1]
        assert "policy_for_period" in str(records) or "policy_assignment" in str(records)
    else:
        text = "\n".join(page.extract_text() for page in PdfReader(io.BytesIO(result.content)).pages)
        assert "Awaiting configuration" in text and expected[0]["name"] in text
    login(client, "agent2")
    assert client.get("/api/commissions/export", params={"period": "2099-10", "user_id": expected[0]["user_id"], "format": format}).status_code == 404
    login(client, "compliance")
    assert client.get("/api/commissions/export", params={"period": "2099-10", "format": format}).status_code == 403


def test_export_permissions_respect_revoked_grants_without_broadening_sales_agent_scope(client):
    from app.db import DB, Permission, Role
    from sqlalchemy import delete, select
    login(client)
    with DB() as db:
        role_id = db.scalar(select(Role.id).where(Role.name == "Administrator"))
        grant = db.scalar(select(Permission).where(Permission.role_id == role_id, Permission.name == "report.read"))
        assert grant
        grant_id = grant.id
        db.execute(delete(Permission).where(Permission.id == grant_id))
        db.commit()
    try:
        assert client.get("/api/auth/me").status_code == 200
        for route in ("/api/reports/daily?format=xlsx", "/api/sales-management/export?format=pdf", "/api/sales-management/targets/template", "/api/field-assets/report/export?format=xlsx", "/api/commissions/export?format=xlsx", "/api/incentives/export"):
            assert client.get(route).status_code == 403, route
    finally:
        with DB() as db:
            db.add(Permission(id=grant_id, role_id=role_id, name="report.read"))
            db.commit()
    login(client, "agent1")
    # A Sales agent's own downloads do not gain broad report privileges.
    assert client.get("/api/commissions/export?format=xlsx").status_code == 200
    assert client.get("/api/reports/daily?format=xlsx").status_code == 403
