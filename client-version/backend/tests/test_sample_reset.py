from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from app.db import Base, User, Plan, KycCapture, Branch, FieldTask, Agent, SalesRecord, SalesTarget, Audit
from app import sample_reset, seed as seed_module


def test_reset_preserves_accounts_and_edited_plans(tmp_path, monkeypatch):
    from app import commissions, sr_verification, email_delivery  # noqa: F401
    engine = create_engine('sqlite:///' + str(tmp_path / 'fresh.db'))
    Base.metadata.create_all(engine)
    db_factory = sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(seed_module, 'DB', db_factory)
    monkeypatch.setattr(sample_reset, 'DB', db_factory)
    seed_module.seed()
    with db_factory.begin() as db:
        plan = db.scalar(select(Plan))
        plan.name, plan.monthly_cost = 'Owner edited plan', 399
        # Preserved real account emails must never receive synthetic SR alerts.
        for index, person in enumerate(db.scalars(select(User))):
            person.email = f'preserved-account-{index}@real-company.com'
        db.scalar(select(User).where(User.name == 'Alex Morgan')).email = 'admin@relay.demo'
    def identities():
        with db_factory() as db:
            return (sorted((u.id, u.email, u.name, u.password_hash, u.role_id) for u in db.scalars(select(User))),
                    sorted((p.id, p.name, str(p.monthly_cost), p.active) for p in db.scalars(select(Plan))))
    before = identities()
    with db_factory() as db:
        assignments = sorted((a.id, a.outlet_id, a.leader_id) for a in db.scalars(select(Agent)))
        branches = sorted((b.id, b.name) for b in db.scalars(select(Branch)))
        audit_ids = {a.id for a in db.scalars(select(Audit))}
    monkeypatch.setenv('RELAY_RESET_OPERATIONAL_SAMPLES', 'KEEP_ACCOUNTS_AND_PLANS')
    sample_reset.reset_samples()
    assert identities() == before
    with db_factory() as db:
        assert len(list(db.scalars(select(KycCapture)))) == 14
        assert sorted((b.id, b.name) for b in db.scalars(select(Branch))) == branches
        assert sorted((a.id, a.outlet_id, a.leader_id) for a in db.scalars(select(Agent))) == assignments
        assert audit_ids <= {a.id for a in db.scalars(select(Audit))}
        assert len(list(db.scalars(select(SalesRecord)))) == 14
        outbox = list(db.scalars(select(email_delivery.EmailOutbox)))
        assert outbox
        assert all(item.status == 'SKIPPED' and item.attempts == 0 for item in outbox)
        from app.sales_management import ORDER_TYPES
        assert {t.order_type for t in db.scalars(select(SalesTarget))} == set(ORDER_TYPES)
        assert not list(db.scalars(select(FieldTask)))
        for capture in db.scalars(select(KycCapture)):
            from app.captures import payload, Intake
            data = payload(capture)
            assert data['document_kind'] == 'SALE_SCREENSHOTS'
            assert data['intake']['capture_mode'] == 'SCREENSHOT_SALE'
            assert data['intake']['sr_number']
            assert data['intake']['order_image']
            assert data['intake']['document_image']
            assert data['synthetic'] is True
            Intake.model_validate(data['intake']).complete()
    # A second explicit reset clears every dependent table and leaves one small
    # current dataset, rather than growing queues or duplicate claims each run.
    sample_reset.reset_samples()
    assert identities() == before
    with db_factory() as db:
        assert len(list(db.scalars(select(SalesRecord)))) == 14
    engine.dispose()


def test_fresh_records_match_current_workflow_and_stock(tmp_path, monkeypatch):
    import io
    from openpyxl import load_workbook
    from app import commissions, sr_verification, email_delivery  # noqa: F401
    from app.captures import payload, Intake
    from app.db import CallAttempt, FieldAsset, FieldAssetMovement, FieldAssetRequest, SavedCaptureDraft, Sim, Movement
    from app.sales_management import ORDER_TYPES, call_timing, call_stage_access, permitted
    from app.security import cipher
    engine = create_engine('sqlite:///' + str(tmp_path / 'workflows.db'))
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False)
    monkeypatch.setattr(seed_module, 'DB', factory)
    seed_module.seed()
    monkeypatch.setenv('RELAY_RESET_OPERATIONAL_SAMPLES', 'KEEP_ACCOUNTS_AND_PLANS')
    from app.fresh_samples import refresh_samples
    refresh_samples(factory)
    with factory() as db:
        sales = list(db.scalars(select(SalesRecord)))
        assert {s.order_type for s in sales} == set(ORDER_TYPES)
        assert {s.agent_id for s in sales} == {a.id for a in db.scalars(select(Agent))}
        assert {s.status for s in sales} == {'CLOSED', 'CANCELLED', 'IN_PROGRESS'}
        assert {s.status for s in db.scalars(select(sr_verification.SrCheck))} == {'MATCHED', 'MISMATCH', 'PENDING_SR_VERIFICATION'}
        assert any(s.status == 'CLOSED' and db.scalar(select(sr_verification.SrCheck).where(sr_verification.SrCheck.sale_id == s.id)).status == 'MISMATCH' for s in sales)
        for sale in sales:
            if sale.status in {'CLOSED', 'CANCELLED'} and sale.order_type != 'ELIFE':
                sim = db.get(Sim, sale.details['stock_deducted_sim_id'])
                assert sim.status == 'ACTIVATED'
                assert sim.serial == sale.details['sim_serial']
                intake = payload(db.get(KycCapture, sale.capture_id))['intake']
                assert intake['sim_type'] == ('ESIM' if sim.sim_type == 'eSIM' else 'PHYSICAL')
                assert sim.business_category == {'HW': 'Home Wireless', 'WASEL': 'Wasel / Prepaid', 'VISITOR': 'Visitor'}.get(sale.order_type, 'Postpaid')
                assert len(list(db.scalars(select(Movement).where(Movement.sim_id == sim.id, Movement.new_status == 'ACTIVATED')))) == 1
            if sale.order_type == 'ELIFE':
                assert not sale.details.get('sim_serial')
            if sale.order_type in {'HW', 'ELIFE'}:
                assert sale.plan_name in {'Home Wireless', 'eLife'}
                assert sale.details['monthly_cost'] == ''
        draft = db.scalar(select(SavedCaptureDraft))
        data = payload(draft)
        assert data['saved_draft_id'] == draft.id
        assert Intake.model_validate(data).name == 'Jamie Cole'
        assert data['step'] == 0
        assert list(db.scalars(select(FieldAssetRequest)))
        assets = list(db.scalars(select(FieldAsset)))
        assert {a.category for a in assets} == {'ROUTER', 'TABLET', 'UNIFORM'}
        assert len(list(db.scalars(select(FieldAssetMovement)))) == len(assets)
        assert any(c.stage == 'WELCOME_CALL' for c in db.scalars(select(CallAttempt)))
        for attempt in db.scalars(select(CallAttempt)):
            actor = db.get(User, attempt.actor_id)
            sale = db.get(SalesRecord, attempt.sale_id)
            call_stage_access(db, actor, attempt.stage)
            assert permitted(db, actor, sale), 'Sample call actor must have access to this sale'
        for batch in db.scalars(select(sr_verification.SrBatch)):
            book = load_workbook(io.BytesIO(cipher.decrypt(batch.report_encrypted.encode())), read_only=True)
            assert book.active.cell(1, 1).value == 'SR number'
            for check in db.scalars(select(sr_verification.SrCheck).where(sr_verification.SrCheck.batch_id == batch.id)):
                assert check.business_date == batch.business_date
        deferred = next(s for s in sales if s.details.get('tele_deferred_until'))
        assert deferred.details['tele_defer_reason']
        assert call_timing(deferred, 'TELE_VERIFICATION')['deferred']
    engine.dispose()


def test_flexible_fields_keep_unknown_labels_and_multiline_values():
    from app.capture_ocr import organize_lines
    rows = organize_lines([{'text': 'Terminal location  Marina kiosk'}, {'text': 'Authorization code:'}, {'text': 'ABC123'}, {'text': 'Unclassified note'}])
    fields = rows[0]['fields']
    assert fields[0]['label'] == 'Terminal location'
    assert fields[0]['value'] == 'Marina kiosk'
    assert fields[1]['value'] == 'ABC123'
    assert fields[2]['value'] == 'Unclassified note'
