"""Imported final commission amounts, separate from policy-derived estimates."""

import base64
import csv
import hashlib
import io
import json
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import islice
from typing import Literal
from zipfile import ZipFile
from fastapi import APIRouter, Depends, HTTPException, Request, Response
from pydantic import Field
from openpyxl import Workbook, load_workbook
from sqlalchemy import Date, ForeignKey, Numeric, String, select, text
from sqlalchemy.orm import Mapped, mapped_column
from sqlalchemy.exc import IntegrityError
from .db import Agent, Branch, Entity, Outlet, User, get_db, business_date
from .validation import BusinessInput
from .security import principal, permissions
from .commissions import allowed_subjects, role_name
from .services import audit

router = APIRouter(prefix="/api/incentives/final", tags=["Final commission amounts"])
HEADERS = ["user_email", "earned_on", "amount", "reference", "note"]


class FinalIncentive(Entity):
    __tablename__ = "final_incentives"
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id"), index=True)
    branch_id: Mapped[str | None] = mapped_column(
        ForeignKey("branches.id"), nullable=True, index=True
    )
    leader_id: Mapped[str | None] = mapped_column(ForeignKey("users.id"), nullable=True, index=True)
    earned_on: Mapped[date] = mapped_column(Date, index=True)
    amount: Mapped[Decimal] = mapped_column(Numeric(12, 2))
    reference: Mapped[str] = mapped_column(String(120))
    note: Mapped[str] = mapped_column(String(300), default="")
    source_role: Mapped[str] = mapped_column(String(50))
    import_key: Mapped[str] = mapped_column(String(64), unique=True)
    imported_by: Mapped[str] = mapped_column(ForeignKey("users.id"))
    filename: Mapped[str] = mapped_column(String(180))


def write_access(db, user):
    if role_name(db, user) not in {"Administrator", "Operations Manager"} or not {
        "read",
        "incentive.write",
    } <= permissions(db, user):
        raise HTTPException(403, "Only authorized management may import final commission amounts")


def scoped(db, user):
    people = allowed_subjects(db, user)
    return select(FinalIncentive).where(FinalIncentive.user_id.in_([p.id for p in people]))


def view(row, people, branches, agents):
    person = people[row.user_id]
    branch = branches.get(row.branch_id)
    leader = people.get(row.leader_id)
    agent = agents.get(row.user_id)
    return {
        "id": row.id,
        "user_id": row.user_id,
        "name": person.name,
        "user_email": person.email,
        "role": row.source_role,
        "branch_id": row.branch_id,
        "branch": branch.name if branch else "Not recorded",
        "leader_id": row.leader_id,
        "team_leader": leader.name if leader else "Not recorded",
        "agent_id": agent.id if agent else None,
        "employee_id": agent.employee_id if agent else "",
        "earned_on": row.earned_on.isoformat(),
        "amount": str(row.amount),
        "reference": row.reference,
        "note": row.note,
        "filename": row.filename,
        "imported_at": row.created_at,
        "source": "Imported final amount",
    }


def views(db, rows):
    """Use a fixed number of reads for the list and its scoped exports."""
    if not rows:
        return []
    subject_ids = {row.user_id for row in rows}
    person_ids = subject_ids | {row.leader_id for row in rows if row.leader_id}
    branch_ids = {row.branch_id for row in rows if row.branch_id}
    people = {
        person.id: person for person in db.scalars(select(User).where(User.id.in_(person_ids)))
    }
    branches = {
        branch.id: branch for branch in db.scalars(select(Branch).where(Branch.id.in_(branch_ids)))
    }
    agents = {
        agent.user_id: agent
        for agent in db.scalars(select(Agent).where(Agent.user_id.in_(subject_ids)))
    }
    return [view(row, people, branches, agents) for row in rows]


@router.get("")
def listing(
    from_date: date | None = None,
    to_date: date | None = None,
    branch_id: str = "",
    user_id: str = "",
    leader_id: str = "",
    user=Depends(principal),
    db=Depends(get_db),
):
    first, last = from_date or business_date().replace(day=1), to_date or business_date()
    if first > last or (last - first).days > 366:
        raise HTTPException(422, "Choose an ordered date range of up to one year")
    query = scoped(db, user).where(
        FinalIncentive.earned_on >= first, FinalIncentive.earned_on <= last
    )
    for column, value in (
        (FinalIncentive.branch_id, branch_id),
        (FinalIncentive.user_id, user_id),
        (FinalIncentive.leader_id, leader_id),
    ):
        if value:
            query = query.where(column == value)
    rows = list(
        db.scalars(query.order_by(FinalIncentive.earned_on.desc(), FinalIncentive.id).limit(5001))
    )
    if len(rows) > 5000:
        raise HTTPException(
            422, "Select a smaller date range or branch to keep the report within 5,000 entries"
        )
    return views(db, rows)


@router.get("/template")
def template(user=Depends(principal), db=Depends(get_db)):
    write_access(db, user)
    from .exports import style_sheet, add_metadata, XLSX_MEDIA

    book = Workbook()
    sheet = book.active
    sheet.title = "Final amounts"
    sheet.append(HEADERS)
    style_sheet(sheet, HEADERS, machine_headers=True)
    add_metadata(
        book,
        "Final commission import",
        {"currency": "AED", "maximum_rows": 500},
        [
            "Enter one final approved amount per user, date and reference. Use a known sales agent, team leader or Sales Manager email. earned_on uses YYYY-MM-DD; amount has at most two decimals. The same reference/date/user cannot be imported twice. Ask the Super admin to review conflicting amounts; corrections cannot be added as a second import. Formula cells are rejected. Calculated estimates stay separate."
        ],
    )
    output = io.BytesIO()
    book.save(output)
    return Response(
        output.getvalue(),
        media_type=XLSX_MEDIA,
        headers={
            "Content-Disposition": 'attachment; filename="qanawat-final-commissions-template.xlsx"',
            "Cache-Control": "private, no-store",
        },
    )


class Upload(BusinessInput):
    filename: str = Field(min_length=4, max_length=180)
    content_base64: str = Field(min_length=4, max_length=2_800_000)
    apply: bool = False


def read_file(body):
    try:
        raw = base64.b64decode(body.content_base64, validate=True)
        if len(raw) > 2_000_000:
            raise ValueError()
        if body.filename.lower().endswith(".csv"):
            rows = list(islice(csv.reader(io.StringIO(raw.decode("utf-8-sig"))), 502))
        elif body.filename.lower().endswith(".xlsx"):
            with ZipFile(io.BytesIO(raw)) as archive:
                if (
                    len(archive.infolist()) > 100
                    or sum(item.file_size for item in archive.infolist()) > 10_000_000
                ):
                    raise ValueError()
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                sheet = book.active
                # Read-only worksheets synthesize empty cells up to their declared
                # dimensions. Reject inflated ranges before materializing rows.
                if sheet.max_column != len(HEADERS) or not 2 <= (sheet.max_row or 0) <= 501:
                    raise ValueError()
                rows = list(sheet.iter_rows(max_col=len(HEADERS), max_row=501, values_only=True))
                while rows and all(value is None for value in rows[-1]):
                    rows.pop()
            finally:
                book.close()
        else:
            raise ValueError()
    except Exception:
        raise HTTPException(422, "Upload a valid CSV or Excel .xlsx file under 2 MB") from None
    if (
        not rows
        or [str(x or "").strip().lower() for x in rows[0]] != HEADERS
        or not 1 <= len(rows) - 1 <= 500
    ):
        raise HTTPException(
            422,
            "Use the template columns user_email, earned_on, amount, reference, note with 1–500 rows",
        )
    return rows[1:]


@router.post("/import")
def import_file(body: Upload, request: Request, user=Depends(principal), db=Depends(get_db)):
    write_access(db, user)
    people = {p.email.casefold(): p for p in allowed_subjects(db, user)}
    prepared = []
    keys = set()
    for number, row in enumerate(read_file(body), 2):
        if len(row) != 5:
            raise HTTPException(422, f"Row {number}: expected five columns")
        email, earned, amount, reference, note = row
        person = people.get(str(email or "").strip().casefold())
        if not person:
            raise HTTPException(422, f"Row {number}: unknown commission account email")
        try:
            day = earned.date() if isinstance(earned, datetime) else date.fromisoformat(str(earned))
            value = Decimal(str(amount))
            reference = str(reference or "").strip()
            note = str(note or "").strip()
            if (
                day > business_date()
                or not value.is_finite()
                or not 0 <= value <= 100000
                or value.as_tuple().exponent < -2
                or not reference
                or len(reference) > 120
                or len(note) > 300
                or reference.startswith(("=", "+", "@"))
                or note.startswith("=")
            ):
                raise ValueError()
        except (ValueError, InvalidOperation):
            raise HTTPException(
                422,
                f"Row {number}: use a non-future date, final AED amount, text reference and note",
            ) from None
        key = hashlib.sha256(
            json.dumps([person.id, day.isoformat(), reference.casefold()]).encode()
        ).hexdigest()
        if key in keys:
            raise HTTPException(422, f"Row {number}: duplicate user/date/reference in this file")
        keys.add(key)
        agent = db.scalar(select(Agent).where(Agent.user_id == person.id))
        branch_id = db.get(Outlet, agent.outlet_id).branch_id if agent else person.branch_id
        prepared.append(
            {
                "user_id": person.id,
                "branch_id": branch_id,
                "leader_id": agent.leader_id
                if agent
                else person.id
                if role_name(db, person) == "Team Leader"
                else None,
                "earned_on": day,
                "amount": value,
                "reference": reference,
                "note": note,
                "source_role": role_name(db, person),
                "import_key": key,
                "imported_by": user.id,
                "filename": body.filename,
            }
        )
    if body.apply and db.bind.dialect.name == "postgresql":
        if not db.scalar(text("SELECT pg_try_advisory_xact_lock(82831021)")):
            raise HTTPException(409, "Another commission import is running. Retry shortly.")
    prior = {
        row.import_key: row
        for row in db.scalars(select(FinalIncentive).where(FinalIncentive.import_key.in_(keys)))
    }
    for row in prepared:
        old = prior.get(row["import_key"])
        if old and (old.amount != row["amount"] or old.note != row["note"]):
            raise HTTPException(
                409,
                "An imported user/date/reference already has a different amount or note. Review it before importing.",
            )
    additions = [row for row in prepared if row["import_key"] not in prior]
    if not body.apply:
        return {
            "applied": False,
            "new_rows": len(additions),
            "already_recorded": len(prior),
            "total_amount": str(sum((row["amount"] for row in additions), Decimal("0"))),
            "rows": [
                {
                    "name": db.get(User, row["user_id"]).name,
                    "earned_on": row["earned_on"],
                    "amount": str(row["amount"]),
                    "reference": row["reference"],
                }
                for row in additions
            ],
        }
    for values in additions:
        row = FinalIncentive(**values)
        db.add(row)
        db.flush()
        audit(
            db,
            user,
            "Final Commission Imported",
            row.id,
            new={
                "user_id": row.user_id,
                "branch_id": row.branch_id,
                "earned_on": str(row.earned_on),
                "amount": str(row.amount),
                "reference": row.reference,
            },
            request=request,
        )
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            409, "This commission was imported concurrently. Preview the file again."
        ) from None
    return {"applied": True, "count": len(additions), "already_recorded": len(prior)}


@router.get("/export")
def export(
    from_date: date | None = None,
    to_date: date | None = None,
    branch_id: str = "",
    user_id: str = "",
    leader_id: str = "",
    format: Literal["csv", "xlsx", "pdf"] = "xlsx",
    request: Request = None,
    user=Depends(principal),
    db=Depends(get_db),
):
    if not (
        "report.read" in permissions(db, user)
        or role_name(db, user) == "Field Agent"
        and "read" in permissions(db, user)
    ):
        raise HTTPException(403, "Report access required")
    rows = listing(from_date, to_date, branch_id, user_id, leader_id, user, db)
    from .exports import table_download

    audit(
        db,
        user,
        "Final Commissions Exported",
        "date-range",
        new={"from_date": str(from_date), "to_date": str(to_date), "rows": len(rows)},
        request=request,
    )
    db.commit()
    return table_download(
        [
            "earned_on",
            "name",
            "role",
            "employee_id",
            "branch",
            "team_leader",
            "amount",
            "reference",
            "note",
            "filename",
            "imported_at",
        ],
        rows,
        "Final commission tracking",
        format,
        metadata={
            "currency": "AED",
            "from_date": from_date or business_date().replace(day=1),
            "to_date": to_date or business_date(),
            "basis": "Imported final amounts; calculated estimates excluded",
        },
        machine_headers=False,
    )
