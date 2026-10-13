"""Bounded, atomic imports for branch-linked staff, using normal provisioning rules."""

import base64
import csv
import io
import re
from itertools import islice
from typing import Literal
from zipfile import ZipFile

from fastapi import APIRouter, Depends, HTTPException, Request, Response
from openpyxl import Workbook, load_workbook
from pydantic import Field, ValidationError
from sqlalchemy import select, text
from .db import Agent, Branch, Role, User, get_db
from .organization import Creation, administrator, provision
from .security import principal
from .validation import BusinessInput

router = APIRouter(prefix="/api/organization/import", tags=["Organization Excel import"])
HEADERS = {
    "branches": ["name"],
    "teams": ["name", "email", "password", "branch_name"],
    "agents": [
        "name",
        "email",
        "password",
        "branch_name",
        "leader_email",
        "employee_id",
        "daily_target",
    ],
}
Kind = Literal["branches", "teams", "agents"]


def access(db, user, kind):
    administrator(db, user)
    if kind != "branches" and db.get(Role, user.role_id).name != "Administrator":
        raise HTTPException(403, "Only the Super admin can import staff accounts")


@router.get("/{kind}/template")
def template(kind: Kind, user=Depends(principal), db=Depends(get_db)):
    access(db, user, kind)
    from .exports import XLSX_MEDIA, add_metadata, style_sheet

    book = Workbook()
    book.active.title = "Import"
    book.active.append(HEADERS[kind])
    style_sheet(book.active, HEADERS[kind], machine_headers=True)
    add_metadata(
        book,
        "Branch and staff import",
        {"maximum_rows": 50},
        [
            "Create branches first, then team leaders, then sales agents. Branch names and leader emails must match existing records.",
            "Staff passwords require 10–72 bytes. Keep this file private and share passwords securely. Existing accounts are never silently changed. Daily target is the overall daily target; set product-specific targets in Sales management.",
            "Preview before applying. Duplicate or invalid rows reject the whole file. Exact repeated imports do not add duplicate records.",
        ],
    )
    output = io.BytesIO()
    book.save(output)
    return Response(
        output.getvalue(),
        media_type=XLSX_MEDIA,
        headers={
            "Content-Disposition": f'attachment; filename="qanawat-{kind}-template.xlsx"',
            "Cache-Control": "private, no-store",
        },
    )


class Upload(BusinessInput):
    filename: str = Field(min_length=4, max_length=180)
    content_base64: str = Field(min_length=4, max_length=2_800_000)
    apply: bool = False


def rows_from(body, kind):
    try:
        raw = base64.b64decode(body.content_base64, validate=True)
        if len(raw) > 2_000_000:
            raise ValueError()
        if body.filename.lower().endswith(".csv"):
            rows = list(islice(csv.reader(io.StringIO(raw.decode("utf-8-sig"))), 52))
        elif body.filename.lower().endswith(".xlsx"):
            with ZipFile(io.BytesIO(raw)) as archive:
                if (
                    len(archive.infolist()) > 100
                    or sum(item.file_size for item in archive.infolist()) > 10_000_000
                ):
                    raise ValueError()
            book = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                sheet = book.active
                # Bound columns as well as rows before read-only openpyxl expands
                # the workbook's declared dimensions into empty cells.
                if sheet.max_column != len(HEADERS[kind]) or not 2 <= (sheet.max_row or 0) <= 51:
                    raise ValueError()
                rows = list(
                    sheet.iter_rows(max_col=len(HEADERS[kind]), max_row=51, values_only=True)
                )
                while rows and all(value is None for value in rows[-1]):
                    rows.pop()
            finally:
                book.close()
        else:
            raise ValueError()
        return rows
    except Exception:
        raise HTTPException(422, "Choose a valid Excel or CSV file under 2 MB") from None


@router.post("/{kind}")
def upload(kind: Kind, body: Upload, request: Request, user=Depends(principal), db=Depends(get_db)):
    access(db, user, kind)
    rows = rows_from(body, kind)
    if not rows or list(rows[0]) != HEADERS[kind] or not 1 <= len(rows) - 1 <= 50:
        raise HTTPException(422, "Use this import's template with 1–50 rows")
    if body.apply and db.bind.dialect.name == "postgresql":
        if not db.scalar(text("SELECT pg_try_advisory_xact_lock(82831023)")):
            raise HTTPException(409, "Another organization import is running. Retry shortly.")
    branches = {b.name.casefold(): b for b in db.scalars(select(Branch))}
    users = {p.email.casefold(): p for p in db.scalars(select(User))}
    employees = {a.employee_id.casefold(): a for a in db.scalars(select(Agent))}
    prepared, seen, already = [], set(), 0
    employee_ids = set()
    for number, row in enumerate(rows[1:], 2):
        if len(row) != len(HEADERS[kind]):
            raise HTTPException(422, f"Row {number}: incorrect number of columns")
        data = {key: str(value or "").strip() for key, value in zip(HEADERS[kind], row)}
        if any(
            value.startswith(("=", "+", "@")) for key, value in data.items() if key != "password"
        ):
            raise HTTPException(422, f"Row {number}: formula cells are not supported")
        key = data["name"].casefold() if kind == "branches" else data["email"].casefold()
        if key in seen:
            raise HTTPException(422, f"Row {number}: duplicate name or email in this file")
        seen.add(key)
        values = {"name": data["name"]}
        existing = branches.get(key) if kind == "branches" else users.get(key)
        if kind != "branches":
            branch = branches.get(data["branch_name"].casefold())
            if not branch or branch.lifecycle_status != "ACTIVE":
                raise HTTPException(422, f"Row {number}: choose an existing active branch")
            if (
                not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", data["email"])
                or not 10 <= len(data["password"]) <= 72
                or len(data["password"].encode()) > 72
            ):
                raise HTTPException(
                    422, f"Row {number}: use a valid email and a 10–72 byte password"
                )
            values.update(email=data["email"], password=data["password"], branch_id=branch.id)
            if kind == "agents":
                if data["employee_id"].casefold() in employee_ids:
                    raise HTTPException(422, f"Row {number}: duplicate employee ID in this file")
                employee_ids.add(data["employee_id"].casefold())
                leader = users.get(data["leader_email"].casefold())
                if (
                    not leader
                    or db.get(Role, leader.role_id).name != "Team Leader"
                    or leader.branch_id != branch.id
                ):
                    raise HTTPException(
                        422, f"Row {number}: leader must be assigned to the chosen branch"
                    )
                if (
                    not re.fullmatch(r"[A-Za-z0-9_-]{2,40}", data["employee_id"])
                    or not data["daily_target"].isdigit()
                ):
                    raise HTTPException(
                        422, f"Row {number}: enter an employee ID and a whole daily target"
                    )
                values.update(
                    leader_id=leader.id,
                    employee_id=data["employee_id"],
                    target=int(data["daily_target"]),
                )
                employee = employees.get(data["employee_id"].casefold())
                if employee and (not existing or employee.user_id != existing.id):
                    raise HTTPException(
                        409, f"Row {number}: employee ID already belongs to another account"
                    )
        try:
            item = Creation.model_validate(values)
        except ValidationError:
            raise HTTPException(
                422, f"Row {number}: check name, password, employee ID and target"
            ) from None
        if existing:
            same = existing.name == item.name
            if kind != "branches":
                same = (
                    same
                    and existing.branch_id == item.branch_id
                    and db.get(Role, existing.role_id).name
                    == ("Team Leader" if kind == "teams" else "Field Agent")
                )
            if kind == "agents":
                agent = db.scalar(select(Agent).where(Agent.user_id == existing.id))
                same = (
                    same
                    and agent is not None
                    and agent.employee_id.casefold() == item.employee_id.casefold()
                    and agent.leader_id == item.leader_id
                    and agent.target == item.target
                )
            if not same:
                raise HTTPException(
                    409,
                    f"Row {number}: record already exists with different details. Edit it using the management form.",
                )
            already += 1
        else:
            prepared.append(item)
    if not body.apply:
        return {
            "count": len(prepared),
            "already_recorded": already,
            "rows": [
                {
                    "name": p.name,
                    "email": p.email,
                    "branch": db.get(Branch, p.branch_id).name if p.branch_id else p.name,
                }
                for p in prepared
            ],
        }
    for item in prepared:
        provision(kind, item, request, db, user, commit=False)
    db.commit()
    return {"count": len(prepared), "already_recorded": already, "applied": True}
