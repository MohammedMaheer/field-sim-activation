"""Permission scoped work-location tracking, independent of historical sale ownership.

Tracking is explicit and ends with the work session. Precise member coordinates
are encrypted at rest. Device GPS is evidence, not an anti-spoofing guarantee.
"""
import csv
import io
import json
import os
from datetime import date, datetime, time, timedelta, timezone
from typing import Literal
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from pydantic import BaseModel, ConfigDict, Field, field_validator
from sqlalchemy import Boolean, DateTime, Float, ForeignKey, Index, Integer, String, Text, delete, select, text
from sqlalchemy.orm import Mapped, mapped_column
from sqlalchemy.exc import OperationalError

from .db import Agent, Branch, DB, Entity, Outlet, Permission, Role, Shift, User, get_db, now
from .security import cipher, permissions, principal, require
from .services import audit, distance_m

router = APIRouter(prefix="/api/location-tracking", tags=["Work location"])
FRESHNESS_SECONDS = 90
MAX_ACCURACY_M = 100
RETENTION_DAYS = int(os.getenv("WORK_LOCATION_RETENTION_DAYS", "30"))
if not 1 <= RETENTION_DAYS <= 365:
    raise ValueError("Work location retention must be between 1 and 365 days")
READER_ROLES = {"Administrator", "Operations Manager", "Sales Manager", "Branch Manager", "Team Leader"}
MANAGER_ROLES = {"Administrator", "Operations Manager"}


class LocationPolicy(Entity):
    __tablename__ = "work_location_policy"
    enabled: Mapped[bool] = mapped_column(Boolean, default=False)


class WorkFence(Entity):
    __tablename__ = "work_geofences"
    name: Mapped[str] = mapped_column(String(120))
    branch_id: Mapped[str | None] = mapped_column(ForeignKey("branches.id"), unique=True, nullable=True)
    latitude: Mapped[float] = mapped_column(Float)
    longitude: Mapped[float] = mapped_column(Float)
    radius_m: Mapped[int] = mapped_column(Integer)
    enabled: Mapped[bool] = mapped_column(Boolean, default=True)
    updated_at: Mapped[datetime] = mapped_column(DateTime, default=now)


class FenceAssignment(Entity):
    __tablename__ = "work_geofence_assignments"
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id"), unique=True)
    fence_id: Mapped[str] = mapped_column(ForeignKey("work_geofences.id"))


class WorkSession(Entity):
    __tablename__ = "work_location_sessions"
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id"), index=True)
    device: Mapped[str] = mapped_column(String(100), default="")
    ended_at: Mapped[datetime | None] = mapped_column(DateTime, nullable=True)
    __table_args__ = (Index("uq_work_session_active_user", "user_id", unique=True,
                           postgresql_where=text("ended_at IS NULL"), sqlite_where=text("ended_at IS NULL")),)


class WorkLocation(Entity):
    __tablename__ = "work_locations"
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id"))
    session_id: Mapped[str] = mapped_column(ForeignKey("work_location_sessions.id"))
    coordinates_encrypted: Mapped[str] = mapped_column(Text)
    captured_at: Mapped[datetime] = mapped_column(DateTime)
    state: Mapped[str] = mapped_column(String(20))
    fence_id: Mapped[str | None] = mapped_column(ForeignKey("work_geofences.id"), nullable=True)
    __table_args__ = (Index("ix_work_location_user_time", "user_id", "captured_at"),
                     Index("ix_work_location_session_time", "session_id", "captured_at"),)


class LatestWorkLocation(Entity):
    __tablename__ = "work_locations_latest"
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id"), unique=True)
    session_id: Mapped[str] = mapped_column(ForeignKey("work_location_sessions.id"))
    coordinates_encrypted: Mapped[str] = mapped_column(Text)
    captured_at: Mapped[datetime] = mapped_column(DateTime)
    received_at: Mapped[datetime] = mapped_column(DateTime, default=now)
    state: Mapped[str] = mapped_column(String(20))
    fence_id: Mapped[str | None] = mapped_column(ForeignKey("work_geofences.id"), nullable=True)
    state_changed_at: Mapped[datetime] = mapped_column(DateTime, default=now)


class Input(BaseModel):
    model_config = ConfigDict(extra="forbid", allow_inf_nan=False)


class Point(Input):
    latitude: float = Field(ge=-90, le=90)
    longitude: float = Field(ge=-180, le=180)
    accuracy_m: float = Field(gt=0, le=10000)
    captured_at: datetime
    mocked: bool = False
    device: str = Field(default="", max_length=100)

    @field_validator("captured_at")
    @classmethod
    def utc_stamp(cls, value):
        if value.tzinfo is None:
            raise ValueError("Location time must include a timezone")
        return value.astimezone(timezone.utc).replace(tzinfo=None)


class SessionPoint(Point):
    session_id: str = Field(min_length=1, max_length=36)


class SessionEnd(Input):
    session_id: str = Field(min_length=1, max_length=36)


class FenceInput(Input):
    name: str = Field(min_length=2, max_length=120)
    branch_id: str | None = None
    latitude: float = Field(ge=-90, le=90)
    longitude: float = Field(ge=-180, le=180)
    radius_m: int = Field(ge=25, le=100000)
    enabled: bool = True

    @field_validator("name")
    @classmethod
    def clean_name(cls, value):
        value = value.strip()
        if len(value) < 2:
            raise ValueError("Enter a worksite name")
        return value


class AssignmentInput(Input):
    fence_id: str | None = None


class PolicyInput(Input):
    enabled: bool


def policy_enabled(db):
    row = db.get(LocationPolicy, "policy")
    return bool(row and row.enabled)


def branch_for(db, user):
    agent = db.scalar(select(Agent).where(Agent.user_id == user.id))
    outlet = db.get(Outlet, agent.outlet_id) if agent else None
    return outlet.branch_id if outlet else user.branch_id


def fence_for(db, user):
    assigned = db.scalar(select(FenceAssignment).where(FenceAssignment.user_id == user.id))
    if assigned:
        return db.get(WorkFence, assigned.fence_id)
    branch_id = branch_for(db, user)
    return db.scalar(select(WorkFence).where(WorkFence.branch_id == branch_id)) if branch_id else None


def scoped_users(db, actor):
    """Current assignments only: prior sales never grant access to live GPS."""
    grants = permissions(db, actor)
    role = db.get(Role, actor.role_id).name
    if "location.read" not in grants or role not in READER_ROLES:
        return [actor]
    query = select(User)
    if role == "Team Leader":
        ids = list(db.scalars(select(Agent.user_id).where(Agent.leader_id == actor.id,
                                                        Agent.employment_status == "ACTIVE")))
        return list(db.scalars(query.where(User.id.in_([actor.id, *ids])).order_by(User.name)))
    rows = list(db.scalars(query.order_by(User.name)))
    if role == "Branch Manager":
        agent_branches = dict(db.execute(select(Agent.user_id, Outlet.branch_id).join(Outlet, Outlet.id == Agent.outlet_id)).all())
        rows = [person for person in rows if person.id == actor.id or
                (actor.branch_id and agent_branches.get(person.id, person.branch_id) == actor.branch_id)]
    # Do not keep tracking exited sales agents simply because the account exists.
    inactive = set(db.scalars(select(Agent.user_id).where(Agent.employment_status != "ACTIVE")))
    return [person for person in rows if person.id == actor.id or person.id not in inactive]


def manage(db, actor):
    require(db, actor, "location.manage")
    if db.get(Role, actor.role_id).name not in MANAGER_ROLES:
        raise HTTPException(403, "Your role cannot manage work boundaries")


def view_access(db, actor):
    if not permissions(db, actor) & {"location.self", "location.read", "location.manage"}:
        raise HTTPException(403, "Your role cannot access work locations")


def active_session(db, user_id):
    return db.scalar(select(WorkSession).where(WorkSession.user_id == user_id,
                                             WorkSession.ended_at.is_(None)))


def lock_member(db, user_id):
    if db.bind.dialect.name == "sqlite":
        connection = db.connection().connection.driver_connection
        if not connection.in_transaction:
            db.execute(text("BEGIN IMMEDIATE"))
    # Short point/start writes queue for their member rather than failing with
    # an unrelated assignment error. PostgreSQL bounds contention at two seconds.
    try:
        if db.bind.dialect.name == "postgresql":
            db.execute(text("SET LOCAL lock_timeout = '2s'"))
        return db.scalar(select(User).where(User.id == user_id).with_for_update()
                         .execution_options(populate_existing=True))
    except OperationalError as error:
        if getattr(error.orig, "sqlstate", None) not in {"55P03", "40P01"}:
            raise
        db.rollback()
        raise HTTPException(409, "Your work session is being updated. Try again shortly.") from None


def fence_view(row):
    if not row:
        return None
    return {"id": row.id, "name": row.name, "branch_id": row.branch_id,
            "latitude": row.latitude, "longitude": row.longitude,
              "radius_m": row.radius_m, "enabled": row.enabled,
              "updated_at": row.updated_at.isoformat() if row.updated_at else None}


def classify(point, fence, stamp=None):
    stamp = stamp or now()
    if point.mocked:
        return "MOCKED", None
    if point.captured_at < stamp - timedelta(seconds=FRESHNESS_SECONDS) or point.captured_at > stamp + timedelta(seconds=10):
        return "STALE", None
    if not fence or not fence.enabled:
        return "UNASSIGNED", None
    distance = distance_m(point.latitude, point.longitude, fence.latitude, fence.longitude)
    if point.accuracy_m > MAX_ACCURACY_M or distance + point.accuracy_m > fence.radius_m >= distance - point.accuracy_m:
        return "UNCERTAIN", distance
    return ("INSIDE" if distance + point.accuracy_m <= fence.radius_m else "OUTSIDE"), distance


def unpack(row):
    return json.loads(cipher.decrypt(row.coordinates_encrypted.encode()))


def current_assignment(db, user):
    agent = db.scalar(select(Agent).where(Agent.user_id == user.id))
    return {"branch_id": branch_for(db, user), "leader_id": agent.leader_id if agent else None}


def current_point(db, user, row):
    """A new branch or leader never acquires a previous assignment's GPS."""
    snapshot = unpack(row).get("_snapshot", {})
    assignment = current_assignment(db, user)
    return all(snapshot.get(key) == value for key, value in assignment.items())


def point_view(row, fence):
    if not row:
        return None
    data = unpack(row)
    data.pop("_snapshot", None)
    point = Point(**data, captured_at=row.captured_at.replace(tzinfo=timezone.utc))
    state, distance = classify(point, fence)
    return {**data, "captured_at": row.captured_at.replace(tzinfo=timezone.utc), "received_at": getattr(row, "received_at", row.created_at).replace(tzinfo=timezone.utc),
            "state": state, "distance_m": round(distance, 1) if distance is not None else None,
            "fresh": state != "STALE"}


def view_context(db, people):
    """Batch related rows so a fleet refresh does not issue queries per member."""
    ids = [person.id for person in people]
    agents = {row.user_id: row for row in db.scalars(select(Agent).where(Agent.user_id.in_(ids)))}
    outlets = {row.id: row for row in db.scalars(select(Outlet).where(Outlet.id.in_({a.outlet_id for a in agents.values()})))}
    branches = {person.id: outlets[agents[person.id].outlet_id].branch_id if person.id in agents else person.branch_id for person in people}
    fences = {row.id: row for row in db.scalars(select(WorkFence))}
    branch_fences = {row.branch_id: row for row in fences.values() if row.branch_id}
    assignments = {row.user_id: row.fence_id for row in db.scalars(select(FenceAssignment).where(FenceAssignment.user_id.in_(ids)))}
    sessions = {row.user_id: row for row in db.scalars(select(WorkSession).where(WorkSession.user_id.in_(ids), WorkSession.ended_at.is_(None)))}
    latest = {row.user_id: row for row in db.scalars(select(LatestWorkLocation).where(LatestWorkLocation.user_id.in_(ids)))}
    roles = {row.id: row.name for row in db.scalars(select(Role))}
    from .role_policy import ROLE_PERMISSIONS
    grants = {}
    for role_id, name in db.execute(select(Permission.role_id, Permission.name).where(Permission.name.in_(["location.self", "location.read", "location.manage"]))):
        if name in ROLE_PERMISSIONS.get(roles.get(role_id), []):
            grants.setdefault(role_id, set()).add(name)
    return {"agents": agents, "branches": branches, "fences": fences, "branch_fences": branch_fences,
            "assignments": assignments, "sessions": sessions, "latest": latest, "roles": roles,
            "grants": grants, "enabled": policy_enabled(db)}


def status(db, user, context=None):
    context = context or view_context(db, [user])
    assigned = context["assignments"].get(user.id)
    fence = context["fences"].get(assigned) if assigned else context["branch_fences"].get(context["branches"][user.id])
    session = context["sessions"].get(user.id)
    latest = context["latest"].get(user.id)
    agent = context["agents"].get(user.id)
    assignment = {"branch_id": context["branches"][user.id], "leader_id": agent.leader_id if agent else None}
    snapshot = unpack(latest).get("_snapshot", {}) if latest else {}
    matching = all(snapshot.get(key) == value for key, value in assignment.items())
    location = point_view(latest, fence) if latest and matching else None
    enabled = context["enabled"]
    inside = bool(location and location["state"] == "INSIDE")
    active_point = bool(session and latest and latest.session_id == session.id and inside)
    reason = ""
    if enabled:
        if not fence or not fence.enabled:
            reason = "An administrator must assign an active work boundary before you can start work or capture a sale."
        elif not session:
            reason = "Start work tracking inside your assigned boundary before capturing a sale."
        elif not active_point:
            state = location["state"] if location else "NO_LOCATION"
            reason = {"OUTSIDE": "You are outside your assigned work boundary. Return inside to capture a sale.",
                      "UNCERTAIN": "Location accuracy overlaps the boundary. Move further inside and refresh your location.",
                      "STALE": "Your location is out of date. Refresh your location before capturing a sale.",
                      "MOCKED": "Mock location cannot be used to capture a sale.",
                      "UNASSIGNED": "Your work boundary is unavailable. Ask an administrator to update it."}.get(state,
                      "A fresh device location is required before capturing a sale.")
    grants = context["grants"].get(user.role_id, set())
    return {"enabled": enabled, "boundary": fence_view(fence),
            "work_session": {"id": session.id, "started_at": session.created_at.replace(tzinfo=timezone.utc), "ended_at": session.ended_at.replace(tzinfo=timezone.utc) if session.ended_at else None,
                             "device": session.device} if session else None,
            "latest_location": location, "can_start": "location.self" in grants and (not enabled or inside),
            "can_capture": not enabled or active_point, "reason": reason,
            "permissions": {"read_members": "location.read" in grants, "manage": "location.manage" in grants}}


def enforce_capture(db, user):
    if policy_enabled(db):
        require(db, user, "location.self")
        lock_member(db, user.id)
        result = status(db, user)
        if not result["can_capture"]:
            raise HTTPException(403, result["reason"])


def save_point(db, user, session, point):
    fence = fence_for(db, user)
    state, distance = classify(point, fence)
    if state == "STALE":
        raise HTTPException(422, "Refresh your device location. The supplied location is out of date.")
    if point.mocked:
        raise HTTPException(422, "Mock device locations are not accepted")
    latest = db.scalar(select(LatestWorkLocation).where(LatestWorkLocation.user_id == user.id))
    if latest and latest.session_id == session.id and point.captured_at < latest.captured_at:
        raise HTTPException(409, "A newer location has already been received")
    data = {"latitude": point.latitude, "longitude": point.longitude, "accuracy_m": point.accuracy_m, "mocked": point.mocked}
    assignment = current_assignment(db, user)
    data["_snapshot"] = {**assignment, "boundary": fence_view(fence),
                         "distance_m": round(distance, 1) if distance is not None else None}
    if latest and latest.session_id == session.id:
        previous = unpack(latest)
        if point.captured_at == latest.captured_at and previous == data:
            return  # An exact network retry is idempotent; do not inflate history.
        if now() - latest.received_at < timedelta(seconds=2) and previous == data:
            return  # Capture and the automatic watcher can share one fresh reading.
        if now() - latest.received_at < timedelta(seconds=1):
            raise HTTPException(429, "Location updates are arriving too quickly. Retry shortly.", headers={"Retry-After": "1"})
    encrypted = cipher.encrypt(json.dumps(data).encode()).decode()
    row = WorkLocation(user_id=user.id, session_id=session.id, coordinates_encrypted=encrypted,
                       captured_at=point.captured_at, state=state, fence_id=fence.id if fence else None)
    db.add(row)
    if not latest:
        latest = LatestWorkLocation(user_id=user.id)
        db.add(latest)
    if latest.state != state or latest.session_id != session.id or latest.fence_id != (fence.id if fence else None):
        latest.state_changed_at = now()
    latest.session_id, latest.coordinates_encrypted = session.id, encrypted
    latest.captured_at, latest.received_at = point.captured_at, now()
    latest.state, latest.fence_id = state, fence.id if fence else None
    db.flush()


def begin_work(db, user, point, request=None):
    require(db, user, "location.self")
    # Serialize concurrent starts and points for the authenticated member.
    lock_member(db, user.id)
    state, _ = classify(point, fence_for(db, user))
    if state == "STALE" or point.mocked:
        raise HTTPException(422, "A fresh, genuine device location is required to start work")
    if policy_enabled(db) and state != "INSIDE":
        raise HTTPException(422, {"OUTSIDE": "You are outside your assigned work boundary. Move inside before starting work.",
                                 "UNCERTAIN": "Location accuracy overlaps the boundary. Move further inside before starting work."}.get(
                                 state, "Ask an administrator to assign an active work boundary before starting work."))
    session = active_session(db, user.id)
    if session and point.device != session.device:
        raise HTTPException(409, "Work tracking is already active on another device. End that session first.")
    if not session:
        session = WorkSession(user_id=user.id, device=point.device)
        db.add(session)
        db.flush()
        audit(db, user, "Work tracking started", session.id, request=request)
    save_point(db, user, session, point)
    agent = db.scalar(select(Agent).where(Agent.user_id == user.id))
    if agent:
        if not db.scalar(select(Shift).where(Shift.agent_id == agent.id, Shift.ended_at.is_(None))):
            db.add(Shift(agent_id=agent.id))
        agent.status = "ACTIVE"
        agent.last_sync = now()
    db.flush()
    return session


def end_work(db, user, session_id, request=None, *, authorized_logout=False):
    if not authorized_logout:
        require(db, user, "location.self")
    lock_member(db, user.id)
    session = db.scalar(select(WorkSession).where(WorkSession.id == session_id, WorkSession.user_id == user.id))
    if not session:
        raise HTTPException(404, "Work session not found")
    if not session.ended_at:
        session.ended_at = now()
        audit(db, user, "Work tracking ended", session.id, request=request)
        agent = db.scalar(select(Agent).where(Agent.user_id == user.id))
        if agent:
            for shift in db.scalars(select(Shift).where(Shift.agent_id == agent.id, Shift.ended_at.is_(None))):
                shift.ended_at = now()
            agent.status = "OFFLINE"
            agent.last_sync = now()
    db.flush()


@router.get("/me")
def me(user=Depends(principal), db=Depends(get_db)):
    view_access(db, user)
    return status(db, user)


@router.get("/settings")
def settings(user=Depends(principal), db=Depends(get_db)):
    view_access(db, user)
    return {"enabled": policy_enabled(db), "freshness_seconds": FRESHNESS_SECONDS,
            "max_accuracy_m": MAX_ACCURACY_M, "retention_days": RETENTION_DAYS}


@router.patch("/settings")
def update_settings(body: PolicyInput, request: Request, user=Depends(principal), db=Depends(get_db)):
    manage(db, user)
    if body.enabled and not db.scalar(select(WorkFence.id).where(WorkFence.enabled.is_(True))):
        raise HTTPException(422, "Configure an active work boundary before enabling enforcement")
    row = db.get(LocationPolicy, "policy")
    if not row:
        row = LocationPolicy(id="policy", enabled=False)
        db.add(row)
    old = row.enabled
    row.enabled = body.enabled
    audit(db, user, "Work boundary enforcement updated", "policy", old={"enabled": old}, new={"enabled": row.enabled}, request=request)
    db.commit()
    return settings(user, db)


@router.post("/start")
def start(body: Point, request: Request, user=Depends(principal), db=Depends(get_db)):
    begin_work(db, user, body, request)
    db.commit()
    return status(db, user)


@router.post("/points")
def points(body: SessionPoint, user=Depends(principal), db=Depends(get_db)):
    require(db, user, "location.self")
    lock_member(db, user.id)
    session = active_session(db, user.id)
    if not session or session.id != body.session_id:
        raise HTTPException(409, "Start work tracking before sharing your location")
    if body.device != session.device:
        raise HTTPException(409, "This work session belongs to another device. End that session before starting here.")
    save_point(db, user, session, body)
    db.commit()
    return status(db, user)


@router.post("/end")
def end(body: SessionEnd, request: Request, user=Depends(principal), db=Depends(get_db)):
    end_work(db, user, body.session_id, request)
    db.commit()
    return status(db, user)


@router.get("/members")
def members(branch_id: str = "", user=Depends(principal), db=Depends(get_db)):
    view_access(db, user)
    people = scoped_users(db, user)
    context = view_context(db, people)
    if branch_id:
        people = [person for person in people if context["branches"][person.id] == branch_id]
    branch_names = {row.id: row for row in db.scalars(select(Branch).where(Branch.id.in_(set(context["branches"].values()))))}
    items, branches = [], {}
    for person in people:
        branch = branch_names.get(context["branches"][person.id])
        if branch:
            branches[branch.id] = {"id": branch.id, "name": branch.name}
        items.append({"user_id": person.id, "name": person.name, "email": person.email,
                      "role": context["roles"][person.role_id], "branch_id": branch.id if branch else None,
                      "branch": branch.name if branch else "Not assigned", "assignment_fence_id": context["assignments"].get(person.id),
                      **status(db, person, context)})
    return {"items": items, "branches": list(branches.values()), "enabled": policy_enabled(db)}


@router.get("/fences")
def fences(user=Depends(principal), db=Depends(get_db)):
    view_access(db, user)
    if "location.manage" in permissions(db, user):
        return [fence_view(row) for row in db.scalars(select(WorkFence).order_by(WorkFence.name))]
    people = scoped_users(db, user)
    context = view_context(db, people)
    ids = {row["boundary"]["id"] for person in people if (row := status(db, person, context))["boundary"]}
    return [fence_view(row) for row in db.scalars(select(WorkFence).where(WorkFence.id.in_(ids)).order_by(WorkFence.name))]


def lock_boundary_changes(db):
    """Serialize short boundary edits/assignments, including cross-member sharing."""
    if db.bind.dialect.name == "sqlite":
        connection = db.connection().connection.driver_connection
        if not connection.in_transaction:
            db.execute(text("BEGIN IMMEDIATE"))
    elif db.bind.dialect.name == "postgresql":
        try:
            db.execute(text("SET LOCAL lock_timeout = '2s'"))
            db.execute(text("SELECT pg_advisory_xact_lock(82821015)"))
        except OperationalError as error:
            if getattr(error.orig, "sqlstate", None) not in {"55P03", "40P01"}:
                raise
            db.rollback()
            raise HTTPException(409, "Work boundaries are being updated. Try again shortly.") from None


def save_fence(db, user, body, row, request, *, commit=True):
    manage(db, user)
    lock_boundary_changes(db)
    if body.branch_id:
        branch = db.get(Branch, body.branch_id)
        if not branch or branch.lifecycle_status != "ACTIVE":
            raise HTTPException(422, "Choose an active branch")
        duplicate = db.scalar(select(WorkFence.id).where(WorkFence.branch_id == body.branch_id,
                                                       WorkFence.id != (row.id if row else "")))
        if duplicate:
            raise HTTPException(409, "This branch already has a work boundary. Edit that boundary instead.")
    if not row:
        row = WorkFence()
        db.add(row)
    old = fence_view(row) if row.id else {}
    for key, value in body.model_dump().items():
        setattr(row, key, value)
    row.updated_at = now()
    db.flush()
    audit(db, user, "Work boundary saved", row.id, old=old, new=fence_view(row), request=request)
    if commit:
        db.commit()
    return fence_view(row)


@router.post("/fences", status_code=201)
def add_fence(body: FenceInput, request: Request, user=Depends(principal), db=Depends(get_db)):
    return save_fence(db, user, body, None, request)


@router.patch("/fences/{fence_id}")
def edit_fence(fence_id: str, body: FenceInput, request: Request, user=Depends(principal), db=Depends(get_db)):
    manage(db, user)
    lock_boundary_changes(db)
    row = db.get(WorkFence, fence_id)
    if not row:
        raise HTTPException(404, "Work boundary not found")
    return save_fence(db, user, body, row, request)


@router.put("/assignments/{user_id}")
def assign(user_id: str, body: AssignmentInput, request: Request, user=Depends(principal), db=Depends(get_db)):
    manage(db, user)
    lock_boundary_changes(db)
    target = lock_member(db, user_id)
    if not target:
        raise HTTPException(404, "Member not found")
    if body.fence_id:
        fence = db.get(WorkFence, body.fence_id)
        if not fence or not fence.enabled:
            raise HTTPException(422, "Choose an active work boundary")
    row = db.scalar(select(FenceAssignment).where(FenceAssignment.user_id == user_id))
    old = row.fence_id if row else None
    if body.fence_id:
        if not row:
            row = FenceAssignment(user_id=user_id)
            db.add(row)
        row.fence_id = body.fence_id
    elif row:
        db.delete(row)
    audit(db, user, "Member work boundary updated", user_id, old={"fence_id": old}, new={"fence_id": body.fence_id}, request=request)
    db.commit()
    return {"user_id": user_id, "assignment_fence_id": body.fence_id, **status(db, target)}


@router.put("/assignments/{user_id}/boundary")
def individual_boundary(user_id: str, body: FenceInput, request: Request,
                        user=Depends(principal), db=Depends(get_db)):
    """Save and assign a personal circle atomically without moving the member."""
    manage(db, user)
    lock_boundary_changes(db)
    target = lock_member(db, user_id)
    if not target:
        raise HTTPException(404, "Member not found")
    if body.branch_id or not body.enabled:
        raise HTTPException(422, "An individual boundary must be active and independent of a branch")
    assignment = db.scalar(select(FenceAssignment).where(FenceAssignment.user_id == user_id))
    previous = assignment.fence_id if assignment else None
    current = db.get(WorkFence, previous) if previous else None
    # An exclusive member worksite can be edited. A shared worksite or branch
    # boundary must stay intact for other members, so create a personal copy.
    exclusive = current and not current.branch_id and len(db.scalars(
        select(FenceAssignment.user_id).where(FenceAssignment.fence_id == current.id)).all()) == 1
    saved = save_fence(db, user, body, current if exclusive else None, request, commit=False)
    if not assignment:
        assignment = FenceAssignment(user_id=user_id)
        db.add(assignment)
    assignment.fence_id = saved["id"]
    audit(db, user, "Member work boundary updated", user_id,
          old={"fence_id": previous}, new={"fence_id": saved["id"]}, request=request)
    db.commit()
    return {"user_id": user_id, "assignment_fence_id": saved["id"], **status(db, target)}


@router.get("/export")
def export(format: Literal["csv", "xlsx"] = "xlsx", from_date: date | None = None, to_date: date | None = None,
           branch_id: str = "", user_id: str = "", state: str = "", latitude: float | None = None,
           longitude: float | None = None, radius_m: float | None = None,
           user=Depends(principal), db=Depends(get_db)):
    view_access(db, user)
    today = datetime.now(ZoneInfo("Asia/Dubai")).date()
    first, last = from_date or today, to_date or today
    if first > last or (last - first).days > 30:
        raise HTTPException(422, "Choose a date range of up to 31 days")
    if any(value is not None for value in (latitude, longitude, radius_m)):
        if latitude is None or longitude is None or radius_m is None or not (-90 <= latitude <= 90 and -180 <= longitude <= 180 and 25 <= radius_m <= 100000):
            raise HTTPException(422, "Choose a valid map centre and a radius from 25 to 100000 metres")
    states = {"INSIDE", "OUTSIDE", "UNCERTAIN", "STALE", "MOCKED", "UNASSIGNED"}
    if state and state not in states:
        raise HTTPException(422, "Choose a valid location state")
    people = [person for person in scoped_users(db, user) if not user_id or person.id == user_id]
    allowed = {person.id: person for person in people}
    zone = ZoneInfo("Asia/Dubai")
    begin = datetime.combine(first, time(), zone).astimezone(timezone.utc).replace(tzinfo=None)
    finish = datetime.combine(last + timedelta(days=1), time(), zone).astimezone(timezone.utc).replace(tzinfo=None)
    query = select(WorkLocation).where(WorkLocation.user_id.in_(allowed), WorkLocation.captured_at >= begin,
                                       WorkLocation.captured_at < finish).order_by(WorkLocation.captured_at)
    if state:
        query = query.where(WorkLocation.state == state)
    locations = list(db.scalars(query.limit(5001)))
    if len(locations) > 5000:
        raise HTTPException(422, "This export is too large. Select a shorter date range or one member.")
    headers = ["Member", "Role", "Recorded branch", "Recorded at (UTC)", "Latitude", "Longitude", "Accuracy (m)", "Recorded boundary", "Recorded state", "Session", "Boundary latitude", "Boundary longitude", "Boundary radius (m)", "Distance from centre (m)"]
    rows = []
    for location in locations:
        data = unpack(location)
        snapshot = data.get("_snapshot", {})
        if branch_id and snapshot.get("branch_id") != branch_id:
            continue
        role = db.get(Role, user.role_id).name
        if role == "Team Leader" and location.user_id != user.id and snapshot.get("leader_id") != user.id:
            continue
        if role == "Branch Manager" and location.user_id != user.id and snapshot.get("branch_id") != user.branch_id:
            continue
        if radius_m is not None and distance_m(data["latitude"], data["longitude"], latitude, longitude) > radius_m:
            continue
        person = allowed[location.user_id]
        branch = db.get(Branch, snapshot.get("branch_id")) if snapshot.get("branch_id") else None
        fence = snapshot.get("boundary") or {}
        rows.append([person.name, db.get(Role, person.role_id).name, branch.name if branch else "Not assigned",
                     location.captured_at.isoformat() + "Z", data["latitude"], data["longitude"], data["accuracy_m"],
                     fence.get("name", "Not assigned"), location.state, location.session_id,
                     fence.get("latitude", ""), fence.get("longitude", ""), fence.get("radius_m", ""),
                     snapshot.get("distance_m", "")])
    # Text beginning with spreadsheet formulas is always literal, including names.
    def safe(value):
        return "'" + value if isinstance(value, str) and value.lstrip().startswith(("=", "+", "-", "@")) else value
    filename = f"qanawat-work-location-{first}-{last}.{format}"
    if format == "csv":
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerows([[safe(value) for value in row] for row in rows])
        data = io.BytesIO(("\ufeff" + output.getvalue()).encode())
        media = "text/csv; charset=utf-8"
    else:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Work locations"
        sheet.append(headers)
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="862451")
            cell.font = Font(color="FFFFFF", bold=True)
        for row in rows:
            sheet.append([safe(value) for value in row])
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = min(40, max(16, max(len(str(cell.value or "")) for cell in column) + 2))
        data = io.BytesIO()
        workbook.save(data)
        data.seek(0)
        media = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    return StreamingResponse(data, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}"', "Cache-Control": "private, no-store"})


def retain_locations():
    """Bound precise work-location retention without altering sales or account data."""
    with DB() as db:
        cutoff = now() - timedelta(days=RETENTION_DAYS)
        # Short, bounded transactions let collection continue while old history
        # is cleaned. The next cycle resumes any unusually large backlog.
        for model, stamp in ((WorkLocation, WorkLocation.created_at), (LatestWorkLocation, LatestWorkLocation.received_at)):
            for _ in range(10):
                ids = list(db.scalars(select(model.id).where(stamp < cutoff).order_by(stamp).limit(5000)))
                if not ids:
                    break
                db.execute(delete(model).where(model.id.in_(ids)))
                db.commit()
