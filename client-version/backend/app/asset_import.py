"""Atomic branch stock Excel import, independent of customer transactions."""

import base64
import csv
import hashlib
import io
import json
from itertools import islice
from zipfile import ZipFile
from fastapi import APIRouter, Depends, HTTPException, Request, Response
from pydantic import Field
from openpyxl import Workbook, load_workbook
from sqlalchemy import select, text
from sqlalchemy.exc import IntegrityError
from .db import Branch, FieldAsset, FieldAssetMovement, get_db
from .field_assets import may_manage, AssetCreate
from .security import principal
from .services import audit
from .validation import BusinessInput

router = APIRouter(prefix="/api/field-assets", tags=["Branch stock import"])
HEADERS = ["category", "label", "serial", "quantity", "note"]


@router.get("/import-template")
def template(user=Depends(principal), db=Depends(get_db)):
    may_manage(db, user)
    from .exports import XLSX_MEDIA, style_sheet, add_metadata

    book = Workbook()
    sheet = book.active
    sheet.title = "Branch assets"
    sheet.append(HEADERS)
    style_sheet(sheet, HEADERS, machine_headers=True)
    add_metadata(
        book,
        "Branch stock import",
        {"maximum_rows": 500},
        [
            "Select the receiving branch in the application. Categories: GRABBA_DEVICE, ROUTER, STAMP, ID_CARD, UNIFORM, OTHER. Serialized items require quantity 1. Unserialized supplies can have a positive quantity. Existing serials and repeated files cannot double-add stock. Formula cells are rejected."
        ],
    )
    output = io.BytesIO()
    book.save(output)
    return Response(
        output.getvalue(),
        media_type=XLSX_MEDIA,
        headers={
            "Content-Disposition": 'attachment; filename="qanawat-branch-stock-template.xlsx"'
        },
    )


class Upload(BusinessInput):
    branch_id: str
    filename: str = Field(min_length=4, max_length=180)
    content_base64: str = Field(min_length=4, max_length=2_800_000)
    apply: bool = False


@router.post("/import")
def upload(body: Upload, request: Request, user=Depends(principal), db=Depends(get_db)):
    may_manage(db, user)
    from .branch_lifecycle import active_branch

    active_branch(db, body.branch_id)
    try:
        data = base64.b64decode(body.content_base64, validate=True)
        if len(data) > 2_000_000:
            raise ValueError()
        if body.filename.lower().endswith(".csv"):
            rows = list(islice(csv.reader(io.StringIO(data.decode("utf-8-sig"))), 502))
        elif body.filename.lower().endswith(".xlsx"):
            with ZipFile(io.BytesIO(data)) as archive:
                if (
                    len(archive.infolist()) > 100
                    or sum(item.file_size for item in archive.infolist()) > 10_000_000
                ):
                    raise ValueError()
            book = load_workbook(io.BytesIO(data), read_only=True, data_only=False)
            try:
                sheet = book.active
                # Sheet dimensions are untrusted; avoid generating thousands of
                # empty columns before the template header is checked.
                if sheet.max_column != len(HEADERS) or not 2 <= (sheet.max_row or 0) <= 501:
                    raise ValueError()
                rows = list(sheet.iter_rows(max_col=len(HEADERS), max_row=501, values_only=True))
                while rows and all(value is None for value in rows[-1]):
                    rows.pop()
            finally:
                book.close()
        else:
            raise ValueError()
    except Exception:
        raise HTTPException(422, "Choose a valid Excel or CSV file under 2 MB") from None
    if not rows or list(rows[0]) != HEADERS or not 1 <= len(rows) - 1 <= 500:
        raise HTTPException(422, "Use the branch stock template with 1–500 rows")
    prepared = []
    serials = set()
    for number, row in enumerate(rows[1:], 2):
        try:
            if len(row) != 5 or any(
                isinstance(v, str) and v.startswith(("=", "+", "@")) for v in row
            ):
                raise ValueError()
            category, label, serial, quantity, note = row
            if str(quantity).isdigit() is False:
                raise ValueError()
            item = AssetCreate(
                category=str(category),
                label=str(label or ""),
                serial=str(serial or "").strip(),
                quantity=int(quantity),
                note=str(note or ""),
                branch_id=body.branch_id,
            )
            if (
                item.category
                not in {"GRABBA_DEVICE", "ROUTER", "STAMP", "ID_CARD", "UNIFORM", "OTHER"}
                or item.serial
                and (item.quantity != 1 or item.serial in serials)
            ):
                raise ValueError()
            if item.serial:
                serials.add(item.serial)
            prepared.append(item.model_dump())
        except Exception:
            raise HTTPException(
                422,
                f"Row {number}: check category, label, serial and positive quantity (one for serialized items)",
            ) from None
    batch = "IMPORT:" + hashlib.sha256(json.dumps(prepared, sort_keys=True).encode()).hexdigest()
    if body.apply and db.bind.dialect.name == "postgresql":
        if not db.scalar(text("SELECT pg_try_advisory_xact_lock(82831022)")):
            raise HTTPException(409, "Another branch stock import is running. Retry shortly.")
    old = list(db.scalars(select(FieldAsset).where(FieldAsset.batch == batch)))
    if old:
        return {"applied": True, "already_imported": True, "count": len(old)}
    if serials and db.scalar(select(FieldAsset.id).where(FieldAsset.serial.in_(serials)).limit(1)):
        raise HTTPException(409, "A serial is already recorded. No stock was added.")
    if not body.apply:
        return {
            "applied": False,
            "count": len(prepared),
            "quantity": sum(item["quantity"] for item in prepared),
            "branch": db.get(Branch, body.branch_id).name,
            "rows": prepared,
        }
    try:
        for values in prepared:
            values["serial"] = values["serial"] or None
            values["batch"] = batch
            item = FieldAsset(**values, status="AVAILABLE")
            db.add(item)
            db.flush()
            db.add(
                FieldAssetMovement(
                    asset_id=item.id,
                    from_branch_id=body.branch_id,
                    to_branch_id=body.branch_id,
                    actor_id=user.id,
                    quantity_delta=item.quantity,
                    reason="Excel branch stock received",
                )
            )
            audit(
                db,
                user,
                "Branch Stock Imported",
                item.id,
                new={
                    "branch_id": body.branch_id,
                    "category": item.category,
                    "quantity": item.quantity,
                    "batch": batch,
                },
                request=request,
            )
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(
            409, "Stock changed during import. Refresh and preview again."
        ) from None
    return {"applied": True, "already_imported": False, "count": len(prepared)}
