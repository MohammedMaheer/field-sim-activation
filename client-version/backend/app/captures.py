"""Encrypted capture, review and export workflow shared by both editions."""

import base64
import hashlib
import io
import json
from datetime import datetime
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, Request, Response, Query
from pydantic import BaseModel, Field
from sqlalchemy import select, func
from sqlalchemy.exc import IntegrityError
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from .db import KycCapture, DB, User, get_db, now
from .security import principal, permissions, require, assert_agent, visible_agents, cipher
from .services import audit
from .capture_ocr import extractor, inspect_image

router = APIRouter(prefix="/api/kyc-captures", tags=["KYC transaction captures"])


def access(db, user, write=False):
    require(db, user, "read")
    allowed = {"activation.write", "ekyc.write"}
    if not write:
        allowed.add("compliance.write")
    if not permissions(db, user) & allowed:
        raise HTTPException(403, "Your role cannot access transaction captures")


def payload(row):
    return (
        json.loads(cipher.decrypt(row.payload_encrypted.encode()))
        if row.payload_encrypted
        else {"rows": [], "lines": [], "history": []}
    )


def store(row, data):
    row.payload_encrypted = cipher.encrypt(json.dumps(data, ensure_ascii=False).encode()).decode()


def view(row, detail=True):
    result = {
        "id": row.id,
        "agent_id": row.agent_id,
        "source_reference": row.source_reference,
        "status": row.status,
        "version": row.version,
        "image_type": row.image_type,
        "image_hash": row.image_hash,
        "created_at": row.created_at,
        "updated_at": row.updated_at,
        "error": row.error,
    }
    if detail:
        result.update(payload(row))
    return result


def get_capture(db, user, capture_id, lock=False):
    access(db, user)
    q = select(KycCapture).where(KycCapture.id == capture_id)
    row = db.scalar(q.with_for_update() if lock else q)
    if not row:
        raise HTTPException(404, "Capture not found")
    assert_agent(db, user, row.agent_id)
    return row


def record(db, row, user, action, request=None, data=None):
    data = data if data is not None else payload(row)
    data.setdefault("history", []).append(
        {"action": action, "actor": user.name, "at": now().isoformat(), "status": row.status}
    )
    store(row, data)
    row.version += 1
    row.updated_at = now()
    audit(
        db,
        user,
        action,
        row.id,
        row.agent_id,
        new={"status": row.status, "version": row.version},
        request=request,
    )


class CaptureBody(BaseModel):
    agent_id: str
    operation_id: str = Field(min_length=16, max_length=80)
    source_reference: str = Field(min_length=2, max_length=120)
    image_base64: str = Field(max_length=5_333_336)


@router.post("", status_code=201)
def create(body: CaptureBody, request: Request, user=Depends(principal), db=Depends(get_db)):
    access(db, user, True)
    assert_agent(db, user, body.agent_id)
    try:
        data = base64.b64decode(body.image_base64, validate=True)
        image_type = inspect_image(data)
    except (ValueError, TypeError) as exc:
        raise HTTPException(422, str(exc))
    digest = hashlib.sha256(data).hexdigest()
    existing = db.scalar(select(KycCapture).where(KycCapture.operation_id == body.operation_id))

    def replay(row):
        if (
            row.creator_id != user.id
            or row.image_hash != digest
            or row.agent_id != body.agent_id
            or row.source_reference != body.source_reference.strip()
        ):
            raise HTTPException(409, "This upload identifier was already used for another capture")
        return view(row)

    if existing:
        return replay(existing)
    if not body.source_reference.strip():
        raise HTTPException(422, "Enter the source transaction reference")
    queued = db.scalar(
        select(func.count()).select_from(KycCapture).where(KycCapture.status == "QUEUED")
    )
    if queued >= 20:
        raise HTTPException(429, "OCR queue is busy. Keep the image and retry shortly.")
    row = KycCapture(
        agent_id=body.agent_id,
        creator_id=user.id,
        operation_id=body.operation_id,
        source_reference=body.source_reference.strip(),
        image_hash=digest,
        image_type=image_type,
        image_encrypted=cipher.encrypt(data).decode(),
        version=1,
        status="QUEUED",
    )
    db.add(row)
    try:
        db.flush()
        record(db, row, user, "KYC Screenshot Captured", request)
        db.commit()
    except IntegrityError:
        db.rollback()
        return replay(
            db.scalar(select(KycCapture).where(KycCapture.operation_id == body.operation_id))
        )
    return view(row)


@router.get("")
def listing(user=Depends(principal), db=Depends(get_db), limit: int = Query(50, ge=1, le=100),
            offset: int = Query(0, ge=0), search: str = Query("", max_length=120),
            status: Literal["", "QUEUED", "OCR_FAILED", "EXTRACTED", "VALIDATED", "SUBMITTED", "VERIFIED", "REJECTED"] = ""):
    access(db, user)
    q = select(KycCapture).where(KycCapture.agent_id.in_(visible_agents(db, user)))
    if search.strip():
        q = q.where(KycCapture.source_reference.icontains(search.strip(), autoescape=True))
    if status:
        q = q.where(KycCapture.status == status)
    q = q.order_by(KycCapture.created_at.desc(), KycCapture.id.desc()).offset(offset).limit(limit)
    return [view(row, False) for row in db.scalars(q)]


@router.get("/{capture_id}")
def detail(capture_id: str, user=Depends(principal), db=Depends(get_db)):
    return view(get_capture(db, user, capture_id))


@router.get("/{capture_id}/original")
def original(capture_id: str, user=Depends(principal), db=Depends(get_db)):
    row = get_capture(db, user, capture_id)
    return Response(
        cipher.decrypt(row.image_encrypted.encode()),
        media_type=row.image_type,
        headers={
            "Content-Disposition": 'attachment; filename="capture-'
            + row.id
            + ('.png"' if row.image_type == "image/png" else '.jpg"')
        },
    )


class VersionBody(BaseModel):
    version: int = Field(ge=1)


def version_check(row, body):
    if row.version != body.version:
        raise HTTPException(409, "This capture changed. Reload it before saving.")


class TransactionRow(BaseModel):
    reference: str = Field(min_length=2, max_length=120)
    customer: str = Field(default="", max_length=120)
    account: str = Field(default="", max_length=80)
    details: str = Field(min_length=2, max_length=1000)
    source_line: int | None = Field(default=None, ge=0, le=199)


class RowsBody(VersionBody):
    rows: list[TransactionRow] = Field(min_length=1, max_length=100)
    reason: str = Field(min_length=3, max_length=300)


@router.patch("/{capture_id}/rows")
def save_rows(
    capture_id: str, body: RowsBody, request: Request, user=Depends(principal), db=Depends(get_db)
):
    access(db, user, True)
    row = get_capture(db, user, capture_id, True)
    version_check(row, body)
    if row.status not in {"EXTRACTED", "VALIDATED", "REJECTED"}:
        raise HTTPException(409, "Rows cannot be edited at this stage")
    rows = [r.model_dump() for r in body.rows]
    refs = [r["reference"].strip().casefold() for r in rows]
    if any(not r["reference"].strip() or not r["details"].strip() for r in rows) or len(
        set(refs)
    ) != len(refs):
        raise HTTPException(422, "Each row needs a unique reference and transaction details")
    data = payload(row)
    if any(
        r["source_line"] is not None and r["source_line"] >= len(data.get("lines", []))
        for r in rows
    ):
        raise HTTPException(422, "Invalid source line")
    data.setdefault("revisions", []).append(
        {
            "rows": data.get("rows", []),
            "reason": body.reason,
            "actor": user.name,
            "at": now().isoformat(),
        }
    )
    data["rows"] = rows
    row.status = "VALIDATED"
    record(db, row, user, "KYC Rows Validated", request, data)
    db.commit()
    return view(row)


@router.post("/{capture_id}/submit")
def submit(
    capture_id: str,
    body: VersionBody,
    request: Request,
    user=Depends(principal),
    db=Depends(get_db),
):
    access(db, user, True)
    row = get_capture(db, user, capture_id, True)
    if row.status == "SUBMITTED":
        return view(row)
    version_check(row, body)
    if row.status != "VALIDATED":
        raise HTTPException(409, "Review and validate transaction rows before submission")
    row.status = "SUBMITTED"
    record(db, row, user, "KYC Submitted for Backend Review", request)
    db.commit()
    return view(row)


class ReviewBody(VersionBody):
    outcome: Literal["VERIFIED", "REJECTED"]
    reason: str = Field(min_length=5, max_length=300)


@router.post("/{capture_id}/review")
def review(
    capture_id: str, body: ReviewBody, request: Request, user=Depends(principal), db=Depends(get_db)
):
    require(db, user, "compliance.write")
    row = get_capture(db, user, capture_id, True)
    version_check(row, body)
    if row.status != "SUBMITTED":
        raise HTTPException(409, "Only submitted captures can be reviewed")
    if row.creator_id == user.id:
        raise HTTPException(403, "Another authorized reviewer must verify this capture")
    data = payload(row)
    data["review"] = {
        "outcome": body.outcome,
        "reason": body.reason,
        "reviewer": user.name,
        "at": now().isoformat(),
    }
    row.status, row.reviewer_id = body.outcome, user.id
    record(db, row, user, "KYC Backend " + body.outcome, request, data)
    db.commit()
    return view(row)


@router.post("/{capture_id}/retry")
def retry(
    capture_id: str,
    body: VersionBody,
    request: Request,
    user=Depends(principal),
    db=Depends(get_db),
):
    access(db, user, True)
    row = get_capture(db, user, capture_id, True)
    version_check(row, body)
    if row.status != "OCR_FAILED":
        raise HTTPException(409, "Only failed extraction can be retried")
    row.status, row.error = "QUEUED", ""
    record(db, row, user, "KYC OCR Retried", request)
    db.commit()
    return view(row)


@router.get("/{capture_id}/excel")
def excel(capture_id: str, request: Request, user=Depends(principal), db=Depends(get_db)):
    row = get_capture(db, user, capture_id)
    data = payload(row)
    if row.status not in {"VALIDATED", "SUBMITTED", "VERIFIED", "REJECTED"} or not data.get("rows"):
        raise HTTPException(409, "Validate at least one transaction row before export")
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Transactions"
    sheet.append(
        [
            "Transaction reference",
            "Customer",
            "Account / MSISDN",
            "Details",
            "Source line",
            "Source image reference",
            "AI extraction status",
            "Backend verification status",
        ]
    )
    for item in data["rows"]:
        values = [
            item["reference"],
            item["customer"],
            item["account"],
            item["details"],
            str(item["source_line"] + 1) if item["source_line"] is not None else "Manual",
            f"capture-{row.id}.{'jpg' if row.image_type == 'image/jpeg' else 'png'}",
            "EXTRACTED",
            row.status if row.status in {"VERIFIED", "REJECTED"} else "PENDING",
        ]
        sheet.append(values)
        for cell in sheet[sheet.max_row]:
            cell.data_type = "s"  # User-entered text must never become a spreadsheet formula.
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="761B3A")
    for col, width in zip("ABCDEFGH", [28, 28, 28, 70, 16, 48, 24, 28]):
        sheet.column_dimensions[col].width = width
    meta = workbook.create_sheet("Provenance")
    for key, value in [
        ("Capture", row.id),
        ("Source reference", row.source_reference),
        ("Original SHA-256", row.image_hash),
        ("Generated UTC", datetime.utcnow().isoformat()),
        ("OCR engine", "Tesseract on VPS"),
        ("Verification", row.status),
        ("Review", data.get("review", {}).get("reason", "Pending backend review")),
    ]:
        meta.append([key, value])
        meta.cell(meta.max_row, 2).data_type = "s"
    meta.column_dimensions["A"].width = 24
    meta.column_dimensions["B"].width = 80
    out = io.BytesIO()
    workbook.save(out)
    audit(db, user, "KYC Excel Exported", row.id, row.agent_id, request=request)
    db.commit()
    return Response(
        out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="kyc-{row.id}.xlsx"'},
    )


def process_capture():
    with DB() as db:
        row = db.scalar(
            select(KycCapture)
            .where(KycCapture.status == "QUEUED")
            .order_by(KycCapture.created_at)
            .with_for_update(skip_locked=True)
            .limit(1)
        )
        if not row:
            return
        user = db.get(User, row.creator_id)
        data = payload(row)
        try:
            extracted = extractor.extract(cipher.decrypt(row.image_encrypted.encode()))
            data.update({k: v for k, v in extracted.items() if k != "history"})
            row.status, row.error = "EXTRACTED", ""
        except ValueError:
            row.status, row.error = (
                "OCR_FAILED",
                "No readable text or unsupported image. Try a clearer PNG/JPEG screenshot.",
            )
        except Exception:
            row.status, row.error = (
                "OCR_FAILED",
                "VPS extraction could not complete. Retry or contact operations.",
            )
        record(db, row, user, "KYC OCR " + row.status, data=data)
        db.commit()
