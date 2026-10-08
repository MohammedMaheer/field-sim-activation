"""Sales management records independent of carrier and payment integrations."""

import base64
import csv
import hashlib
import io
import re
import calendar
from datetime import date, datetime, timedelta, timezone
from zoneinfo import ZoneInfo
from itertools import islice
from zipfile import ZipFile
from typing import Literal

from fastapi import APIRouter, Depends, HTTPException, Request
from pydantic import Field, model_validator, field_validator
from .validation import BusinessInput
from sqlalchemy import select, or_
from sqlalchemy.exc import IntegrityError
from openpyxl import load_workbook

from .db import (
    Agent, Branch, CallAttempt, NoSaleFeedback, Outlet, Role, SalesCallTask, SalesRecord,
    SalesTarget, Session, User, Customer, Sim, SimProgress, Movement, KycCapture, get_db, now, business_date,
)
from .security import agent_for_update, assert_agent, cipher, password_hash, permissions, principal
from .services import audit

router = APIRouter(prefix="/api/sales-management", tags=["Sales management"])
ORDER_TYPES = ("NEW", "MNP", "P2P", "HW", "ELIFE", "WASEL", "VISITOR")
SALE_STATUSES = ("IN_PROGRESS", "CLOSED", "CANCELLED")
CALL_STAGES = ("TELE_VERIFICATION", "WELCOME_CALL")
STAFF_ROLES = ("Sales Manager", "Tele Verification Officer", "Welcome Call Officer")
CALL_ORDER_TYPES = ("NEW", "MNP", "P2P", "HW", "ELIFE")
ROUTER_FULFILMENTS = ("DELIVERY", "ON_SPOT", "WITHOUT_ROUTER")


def router_required(order_type, details):
    # Preserve the mandatory serial for historical HW entries with no delivery choice.
    return order_type == "HW" and (details.get("router_fulfilment") or "ON_SPOT") == "ON_SPOT"


@router.get("/staff")
def staff(user=Depends(principal), db=Depends(get_db)):
    if db.get(Role, user.role_id).name not in {"Administrator", "Operations Manager"}:
        raise HTTPException(403, "Only administrators and backend staff can manage sales staff")
    users = db.scalars(select(User).join(Role).where(Role.name.in_(STAFF_ROLES)).order_by(User.name)).all()
    return [{"id": row.id, "name": row.name, "email": row.email,
             "role": db.get(Role, row.role_id).name,
             "branch_id": row.branch_id,
             "branch": "All branches" if db.get(Role, row.role_id).name == "Sales Manager" else db.get(Branch, row.branch_id).name if row.branch_id else "All branches"}
            for row in users]


class StaffCreate(BusinessInput):
    name: str = Field(min_length=2, max_length=120)
    email: str = Field(min_length=5, max_length=180)
    password: str = Field(min_length=10, max_length=72)
    role: Literal["Sales Manager", "Tele Verification Officer", "Welcome Call Officer"]
    branch_id: str = ""

    @model_validator(mode="after")
    def encoded_password_limit(self):
        if len(self.password.encode("utf-8")) > 72:
            raise ValueError("Password must contain at most 72 UTF-8 bytes")
        return self


@router.post("/staff", status_code=201)
def create_staff(body: StaffCreate, request: Request, user=Depends(principal), db=Depends(get_db)):
    if db.get(Role, user.role_id).name not in {"Administrator", "Operations Manager"}:
        raise HTTPException(403, "Only administrators and backend staff can manage sales staff")
    email = body.email.strip().lower()
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]+", email):
        raise HTTPException(422, "Enter a valid email address")
    if db.scalar(select(User.id).where(User.email == email)):
        raise HTTPException(409, "This sign-in email already exists")
    if body.branch_id:
        from .branch_lifecycle import active_branch
        active_branch(db, body.branch_id)
    role = db.scalar(select(Role).where(Role.name == body.role))
    if not role:
        raise HTTPException(409, "Run the current migration before adding staff")
    account = User(name=body.name.strip(), email=email,
                   password_hash=password_hash(body.password), role_id=role.id,
                   branch_id=body.branch_id or None)
    db.add(account)
    db.flush()
    audit(db, user, "Sales Staff Added", account.id,
          new={"role": body.role, "branch_id": account.branch_id}, request=request)
    db.commit()
    return {"id": account.id, "name": account.name, "email": account.email, "role": body.role}


class StaffAssignment(BusinessInput):
    role: Literal["Sales Manager", "Tele Verification Officer", "Welcome Call Officer"]
    branch_id: str = ""
    reason: str = Field(min_length=5, max_length=300)


@router.patch("/staff/{staff_id}")
def assign_staff(staff_id: str, body: StaffAssignment, request: Request, user=Depends(principal), db=Depends(get_db)):
    if db.get(Role, user.role_id).name not in {"Administrator", "Operations Manager"}:
        raise HTTPException(403, "Only administrators and backend staff can manage sales staff")
    person = db.get(User, staff_id)
    if not person or db.get(Role, person.role_id).name not in STAFF_ROLES:
        raise HTTPException(404, "Sales staff account not found")
    if body.branch_id:
        from .branch_lifecycle import active_branch
        active_branch(db, body.branch_id)
    old = {"role": db.get(Role, person.role_id).name, "branch_id": person.branch_id}
    new = {"role": body.role, "branch_id": body.branch_id or None}
    if old == new:
        raise HTTPException(409, "No assignment change to record")
    person.role_id = db.scalar(select(Role.id).where(Role.name == body.role))
    person.branch_id = body.branch_id or None
    for session in db.scalars(select(Session).where(Session.user_id == person.id)):
        session.revoked = True
    audit(db, user, "Sales Staff Reassigned", person.id, old=old, new=new,
          reason=body.reason.strip(), request=request)
    db.commit()
    return {"id": person.id, **new}


def assignment(db, agent_id):
    agent = db.get(Agent, agent_id)
    if not agent:
        raise HTTPException(404, "Agent not found")
    outlet = db.get(Outlet, agent.outlet_id)
    return agent, outlet


def assigned_manager(db, branch_id):
    people = db.scalars(select(User).join(Role).where(
        Role.name == "Sales Manager", User.branch_id == branch_id)).all()
    # Ambiguous or missing assignments must be flagged, never guessed.
    return people[0].id if len(people) == 1 else None


def can_record(db, user):
    if db.get(Role, user.role_id).name not in {"Field Agent", "Administrator", "Operations Manager"}:
        raise HTTPException(403, "Only sales agents and backend management can record sales or feedback")


def permitted(db, user, row):
    role = db.get(Role, user.role_id).name
    if role == "Field Agent":
        return db.get(Agent, row.agent_id).user_id == user.id
    if role == "Team Leader":
        return row.leader_id == user.id
    if role == "Sales Manager":
        return "read" in permissions(db, user)
    if role == "Branch Manager":
        return row.branch_id == user.branch_id
    if role in {"Tele Verification Officer", "Welcome Call Officer"}:
        return not user.branch_id or row.branch_id == user.branch_id
    return "read" in permissions(db, user)


def scoped(db, user, model):
    if "read" not in permissions(db, user):
        raise HTTPException(403, "Your role cannot view sales")
    role = db.get(Role, user.role_id).name
    query = select(model)
    if role == "Field Agent":
        agent = db.scalar(select(Agent).where(Agent.user_id == user.id))
        query = query.where(model.agent_id == (agent.id if agent else ""))
    elif role == "Team Leader":
        query = query.where(model.leader_id == user.id)
    elif role == "Branch Manager":
        query = query.where(model.branch_id == user.branch_id)
    elif "read" not in permissions(db, user):
        raise HTTPException(403, "Your role cannot view sales")
    return query


def sale_view(db, row):
    from .sr_verification import sale_state
    sr_state = sale_state(db, row)
    agent, outlet = assignment(db, row.agent_id)
    document = cipher.decrypt(row.document_encrypted.encode()).decode() if row.document_encrypted else ""
    return {
        "id": row.id, "created_at": row.created_at, "capture_id": row.capture_id,
        "agent_id": row.agent_id, "agent": db.get(User, agent.user_id).name,
        "leader_id": row.leader_id, "leader": (db.get(User, row.leader_id).name if db.get(User, row.leader_id) else "Not assigned") if row.leader_id else "Not assigned",
        "manager_id": row.manager_id, "sales_manager": db.get(User, row.manager_id).name if row.manager_id else "Not recorded",
        "employee_id": agent.employee_id,
        "assignment_warning": "Team leader not assigned" if not row.leader_id else "Sales Manager not recorded" if not row.manager_id else "",
        "branch_id": row.branch_id, "branch": db.get(Branch, row.branch_id).name if db.get(Branch, row.branch_id) else "Former branch",
        "outlet_id": row.outlet_id, "outlet": db.get(Outlet, row.outlet_id).name if db.get(Outlet, row.outlet_id) else "Former outlet",
        "order_type": row.order_type, "customer_name": row.customer_name,
        "document": "•••• " + document[-4:] if document else "Not recorded",
        "nationality": row.nationality or "Not recorded",
        "plan_name": row.plan_name, "request_id": row.request_id or "Not recorded",
        "status": row.status, "status_updated_at": row.status_updated_at,
        **row.details,
        "sr_status": sr_state["status"], "sr_verification": sr_state,
        "payment_record_status": row.details.get("payment_record_status") or "NOT_RECORDED",
        "router_fulfilment": (row.details.get("router_fulfilment") or "ON_SPOT") if row.order_type == "HW" else "",
        "stock_recording_status": "SERIAL_NOT_RECORDED" if not row.details.get("sim_serial") else "CONSUMED" if row.details.get("stock_deducted_sim_id") else "AWAITING_BACKEND_CONFIRMATION",
    }


def register_capture_sale(db, capture):
    """Snapshot assignment when an agent submits the new screenshot-order flow."""
    from .captures import payload
    data = payload(capture)
    intake = data.get("intake") or {}
    if intake.get("capture_mode") not in {"SCREENSHOT_ORDER", "SCREENSHOT_SALE"}:
        return
    if db.scalar(select(SalesRecord.id).where(SalesRecord.capture_id == capture.id)):
        return
    if intake.get("order_type") not in ORDER_TYPES:
        raise HTTPException(422, "Choose the order type before submitting the sale")
    if router_required(intake["order_type"], intake) and not intake.get("router_serial", "").strip():
        raise HTTPException(422, "Router serial is required for on-spot home wireless sales")
    agent, outlet = assignment(db, capture.agent_id)
    reference = (intake.get("order_reference") or "").strip() or None
    if reference and db.scalar(select(SalesRecord.id).where(SalesRecord.request_id == reference)):
        raise HTTPException(409, "This request ID is already recorded")
    document = (intake.get("document_number") or "").strip()
    record = SalesRecord(
        capture_id=capture.id, agent_id=agent.id, leader_id=agent.leader_id,
        manager_id=assigned_manager(db, outlet.branch_id),
        branch_id=outlet.branch_id, outlet_id=outlet.id,
        order_type=intake.get("order_type") or "UNSPECIFIED", customer_name=intake.get("name") or "Not recorded",
        document_encrypted=cipher.encrypt(document.encode()).decode() if document else "",
        nationality=intake.get("nationality") or "", plan_name=intake.get("package_name") or intake.get("plan_name") or "Not recorded",
        request_id=reference, status="IN_PROGRESS",
        details={"msisdn": intake.get("msisdn") or "", "source_reference": capture.source_reference,
                 "monthly_cost": intake.get("monthly_cost") or "", "prepayment": intake.get("prepayment") or "",
                 **{key: intake.get(key) or "" for key in ("account_number", "router_serial", "advance_transaction_number", "sr_number", "alternate_number")},
                 "router_fulfilment": intake.get("router_fulfilment", "ON_SPOT") if intake["order_type"] == "HW" else "",
                 "sim_serial": intake.get("sim_identifier") or "",
                 "assignment_effective_at": agent.assignment_effective_at.isoformat() if agent.assignment_effective_at else "",
                 "request_id_on_image": intake.get("order_reference") or ""},
    )
    record.details = {**record.details, "payment_record_status": "RECORDED" if intake.get("payment_image") else "NOT_RECORDED"}
    db.add(record)
    db.flush()
    customer = db.scalar(select(Customer).where(Customer.agent_id == agent.id, Customer.mobile == intake.get("msisdn", ""), Customer.name == record.customer_name))
    if not customer:
        customer = Customer(agent_id=agent.id, name=record.customer_name, mobile=intake.get("msisdn") or "", nationality=record.nationality)
        db.add(customer)
        db.flush()
    record.details = {**record.details, "customer_id": customer.id}
    create_call_tasks(db, record)


def change_sale_status(db, row, status, user, request):
    """Backend confirms an externally completed sale; never calls a carrier."""
    if status == "CLOSED" and row.capture_id:
        capture = db.get(KycCapture, row.capture_id)
        if capture.status != "VERIFIED":
            raise HTTPException(409, "Verify the submitted evidence before closing this sale")
    serial = row.details.get("sim_serial")
    if status == "CLOSED" and serial and not row.details.get("stock_deducted_sim_id"):
        sim = db.scalar(select(Sim).where(or_(Sim.serial == serial, Sim.iccid == serial)).with_for_update())
        if not sim or sim.agent_id != row.agent_id or db.get(Outlet, sim.outlet_id).branch_id != row.branch_id or sim.status not in {"AVAILABLE", "ASSIGNED TO AGENT", "RESERVED"}:
            raise HTTPException(409, "Resolve SIM ownership and availability before confirming this sale")
        previous = sim.status
        sim.status, sim.activated_at = "ACTIVATED", now()
        row.details = {**row.details, "stock_deducted_sim_id": sim.id}
        db.add(Movement(sim_id=sim.id, agent_id=row.agent_id, user_id=user.id, old_status=previous, new_status="ACTIVATED", reason="Backend confirmed external sale " + row.id))
        audit(db, user, "Sale SIM Consumed", sim.id, row.agent_id, new={"sale_id": row.id}, request=request)
        progress = db.scalar(select(SimProgress).where(SimProgress.capture_id == row.capture_id).with_for_update()) if row.capture_id else None
        if progress:
            progress.stage, progress.payment_status = "ACTIVATED", "VERIFIED"
    row.status, row.status_updated_at = status, now()
    update_call_dependencies(db, row)


def create_call_tasks(db, row):
    """Place submitted sales in the internal call queue once per stage."""
    if row.order_type not in CALL_ORDER_TYPES:
        return
    for stage, status in (("TELE_VERIFICATION", "PENDING"), ("WELCOME_CALL", "BLOCKED")):
        if not db.scalar(select(SalesCallTask.id).where(SalesCallTask.sale_id == row.id, SalesCallTask.stage == stage)):
            db.add(SalesCallTask(sale_id=row.id, stage=stage, status=status))


def update_call_dependencies(db, row):
    if row.order_type not in CALL_ORDER_TYPES:
        return
    tele = db.scalar(select(SalesCallTask).where(SalesCallTask.sale_id == row.id, SalesCallTask.stage == "TELE_VERIFICATION"))
    welcome = db.scalar(select(SalesCallTask).where(SalesCallTask.sale_id == row.id, SalesCallTask.stage == "WELCOME_CALL").with_for_update())
    if welcome and welcome.status not in {"COMPLETED", "SKIPPED"}:
        ready = row.status == "CLOSED" and tele and tele.status == "COMPLETED"
        if ready and welcome.status == "BLOCKED":
            welcome.status, welcome.updated_at = "PENDING", now()
        elif not ready and welcome.status != "BLOCKED":
            welcome.status, welcome.updated_at = "BLOCKED", now()


def record_capture_activation(db, capture, outcome):
    row = db.scalar(select(SalesRecord).where(SalesRecord.capture_id == capture.id))
    if row:
        row.status = "CLOSED" if outcome == "ACTIVATED" else "IN_PROGRESS"
        row.status_updated_at = now()
        update_call_dependencies(db, row)


class SaleCreate(BusinessInput):
    agent_id: str
    order_type: Literal["NEW", "MNP", "P2P", "HW", "ELIFE", "WASEL", "VISITOR"]
    customer_name: str = Field(min_length=2, max_length=120)
    document_number: str = Field(default="", max_length=80)
    nationality: str = Field(default="", max_length=80)
    plan_name: str = Field(min_length=2, max_length=160)
    request_id: str = Field(default="", max_length=120)
    account_number: str = Field(default="", max_length=120)
    sim_serial: str = Field(default="", max_length=100)
    router_serial: str = Field(default="", max_length=100)
    router_fulfilment: Literal["DELIVERY", "ON_SPOT", "WITHOUT_ROUTER"] = "ON_SPOT"
    advance_transaction_number: str = Field(default="", max_length=120)
    sr_number: str = Field(default="", max_length=120)
    alternate_number: str = Field(default="", max_length=40)
    msisdn: str = Field(default="", max_length=40)
    note: str = Field(default="", max_length=500)

    @field_validator("router_fulfilment", mode="before")
    @classmethod
    def legacy_router_default(cls, value):
        return "ON_SPOT" if isinstance(value, str) and not value.strip() else value

    @model_validator(mode="after")
    def conditional_fields(self):
        if not self.customer_name.strip() or not self.plan_name.strip():
            raise ValueError("Enter customer and plan names")
        if router_required(self.order_type, self.model_dump()) and not self.router_serial.strip():
            raise ValueError("Router serial is required for on-spot home wireless sales")
        return self


def filtered_sales_query(db, user, period="", order_type="", status="", branch_id="", agent_id="", leader_id="", from_date="", to_date=""):
    query = scoped(db, user, SalesRecord)
    if period:
        if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", period):
            raise HTTPException(422, "Period must use YYYY-MM")
        year, month = map(int, period.split("-"))
        dubai = ZoneInfo("Asia/Dubai")
        start = datetime(year, month, 1, tzinfo=dubai).astimezone(timezone.utc).replace(tzinfo=None)
        end = (datetime(year + 1, 1, 1, tzinfo=dubai) if month == 12 else datetime(year, month + 1, 1, tzinfo=dubai)).astimezone(timezone.utc).replace(tzinfo=None)
        query = query.where(SalesRecord.created_at >= start, SalesRecord.created_at < end)
    try:
        first = date.fromisoformat(from_date) if from_date else None
        last = date.fromisoformat(to_date) if to_date else None
        if last == date.max or first and last and first > last:
            raise ValueError()
    except ValueError:
        raise HTTPException(422, "Choose a valid date range") from None
    for day, is_end in ((first, False), (last, True)):
        if day:
            day = day + timedelta(days=1) if is_end else day
            boundary = datetime.combine(day, datetime.min.time(), ZoneInfo("Asia/Dubai")).astimezone(timezone.utc).replace(tzinfo=None)
            query = query.where(SalesRecord.created_at < boundary if is_end else SalesRecord.created_at >= boundary)
    for key, value in (("order_type", order_type), ("status", status), ("branch_id", branch_id), ("agent_id", agent_id), ("leader_id", leader_id)):
        if value:
            query = query.where(getattr(SalesRecord, key) == value)
    return query


@router.get("/sales")
def sales(user=Depends(principal), db=Depends(get_db), period: str = "", order_type: str = "", status: str = "", branch_id: str = "", agent_id: str = "", leader_id: str = "", from_date: str = "", to_date: str = ""):
    query = filtered_sales_query(db, user, period, order_type, status, branch_id, agent_id, leader_id, from_date, to_date)
    rows = db.scalars(query.order_by(SalesRecord.created_at.desc()).limit(1000)).all()
    return [sale_view(db, row) for row in rows]


@router.get("/sales/{sale_id}")
def sale_detail(sale_id: str, user=Depends(principal), db=Depends(get_db)):
    row = db.get(SalesRecord, sale_id)
    if "read" not in permissions(db, user) or not row or not permitted(db, user, row):
        raise HTTPException(404, "Sale not found")
    attempts = db.scalars(select(CallAttempt).where(CallAttempt.sale_id == row.id).order_by(CallAttempt.created_at)).all()
    tasks = db.scalars(select(SalesCallTask).where(SalesCallTask.sale_id == row.id)).all()
    tele = next((task for task in tasks if task.stage == "TELE_VERIFICATION"), None)
    return {**sale_view(db, row), "calls": [{"stage": task.stage, "status": call_state(row, task, tele), **call_timing(row, task.stage)} for task in tasks if row.order_type in CALL_ORDER_TYPES],
        "attempts": [{"stage": attempt.stage, "outcome": attempt.outcome, "remark": attempt.remark,
                      "at": attempt.created_at, "actor": db.get(User, attempt.actor_id).name} for attempt in attempts]}


@router.get("/performance")
def performance(period: str = "", user=Depends(principal), db=Depends(get_db), order_type: str = "", status: str = "", branch_id: str = "", agent_id: str = "", leader_id: str = "", from_date: str = "", to_date: str = ""):
    period = period or business_date(now()).strftime("%Y-%m")
    if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", period):
        raise HTTPException(422, "Period must use YYYY-MM")
    year, month = (int(part) for part in period.split("-"))
    dubai = ZoneInfo("Asia/Dubai")
    query = filtered_sales_query(db, user, period, order_type, status, branch_id, agent_id, leader_id, from_date, to_date)
    rows = db.scalars(query).all()
    closed = [row for row in rows if row.status == "CLOSED"]
    product = {name: sum(row.order_type == name for row in closed) for name in ORDER_TYPES}
    today = now().replace(tzinfo=timezone.utc).astimezone(dubai).date()
    today_rows = [row for row in rows if row.created_at.replace(tzinfo=timezone.utc).astimezone(dubai).date() == today]
    target_rows = targets(period, user, db)
    target_rows = [item for item in target_rows if
        (not agent_id or item["agent_id"] == agent_id)
        and (not branch_id or db.get(Outlet, db.get(Agent, item["agent_id"]).outlet_id).branch_id == branch_id)
        and (not leader_id or db.get(Agent, item["agent_id"]).leader_id == leader_id)
        and (not order_type or item["order_type"] == order_type)]
    by_agent = {}
    for target_row in target_rows:
        by_agent.setdefault(target_row["agent_id"], []).append(target_row)
    target = sum(next((item["monthly_target"] for item in items if item["order_type"] == "ALL"),
                      sum(item["monthly_target"] for item in items)) for items in by_agent.values())
    month_days = calendar.monthrange(year, month)[1]
    month_start, month_end = date(year, month, 1), date(year, month, month_days)
    elapsed_days = 0 if today < month_start else month_days if today > month_end else today.day
    remaining_days = month_days - elapsed_days
    achieved = sum(business_date(row.created_at) <= today for row in closed)
    daily = {}
    for row in rows:
        day = business_date(row.created_at).isoformat()
        item = daily.setdefault(day, {"date": day, "recorded": 0, "closed": 0, "cancelled": 0, "in_progress": 0})
        item["recorded"] += 1
        item[{"CLOSED": "closed", "CANCELLED": "cancelled"}.get(row.status, "in_progress")] += 1
    return {"period": period, "recorded": len(rows), "closed": len(closed),
            "in_progress": sum(row.status == "IN_PROGRESS" for row in rows),
            "cancelled": sum(row.status == "CANCELLED" for row in rows),
            "monthly_target": target, "remaining": max(target - len(closed), 0),
            "achievement_percent": round(len(closed) * 100 / target, 1) if target else None,
            "daily_by_product": {name: sum(row.order_type == name and row.status == "CLOSED" for row in today_rows) for name in ORDER_TYPES},
            "recorded_today": len(today_rows), "closed_today": sum(row.status == "CLOSED" for row in today_rows),
            "by_product": product, "mtd_achievement": achieved, "as_of": today.isoformat(),
            "elapsed_days": elapsed_days, "remaining_days": remaining_days,
            "crr": achieved / elapsed_days if elapsed_days else None,
            "drr": (target - achieved) / remaining_days if remaining_days and target_rows else None,
            "projection": None, "projection_status": "UNCONFIGURED",
            "daily_summary": [daily[day] for day in sorted(daily)]}


@router.post("/sales", status_code=201)
def create_sale(body: SaleCreate, request: Request, user=Depends(principal), db=Depends(get_db)):
    can_record(db, user)
    if "read" not in permissions(db, user):
        raise HTTPException(403, "Your role cannot record sales")
    if db.get(Role, user.role_id).name == "Sales Manager":
        raise HTTPException(403, "Sales Managers have reporting access")
    assert_agent(db, user, body.agent_id)
    agent, outlet = assignment(db, body.agent_id)
    from .branch_lifecycle import active_branch
    active_branch(db, outlet.branch_id)
    values = body.model_dump()
    reference = values.pop("request_id").strip() or None
    if reference and db.scalar(select(SalesRecord.id).where(SalesRecord.request_id == reference)):
        raise HTTPException(409, "This request ID is already recorded")
    name = values.pop("customer_name").strip()
    document = values.pop("document_number").strip()
    nationality = values.pop("nationality").strip()
    plan = values.pop("plan_name").strip()
    order_type = values.pop("order_type")
    values.pop("agent_id")
    row = SalesRecord(
        agent_id=agent.id, leader_id=agent.leader_id, branch_id=outlet.branch_id,
        manager_id=assigned_manager(db, outlet.branch_id),
        outlet_id=outlet.id, order_type=order_type, customer_name=name,
        document_encrypted=cipher.encrypt(document.encode()).decode() if document else "",
        nationality=nationality, plan_name=plan, request_id=reference,
        status="IN_PROGRESS", details={**{key: value.strip() for key, value in values.items()},
            "assignment_effective_at": agent.assignment_effective_at.isoformat() if agent.assignment_effective_at else ""},
    )
    db.add(row)
    try:
        db.flush()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "This request ID is already recorded") from None
    create_call_tasks(db, row)
    audit(db, user, "Sale Recorded", row.id, row.agent_id, new={"order_type": order_type, "status": row.status}, request=request)
    db.commit()
    return sale_view(db, row)


class SaleCorrection(BusinessInput):
    order_type: Literal["NEW", "MNP", "P2P", "HW", "ELIFE", "WASEL", "VISITOR"]
    router_serial: str = Field(default="", max_length=100)
    router_fulfilment: Literal["DELIVERY", "ON_SPOT", "WITHOUT_ROUTER"] = "ON_SPOT"
    reason: str = Field(min_length=5, max_length=300)

    @field_validator("router_fulfilment", mode="before")
    @classmethod
    def legacy_router_default(cls, value):
        return "ON_SPOT" if isinstance(value, str) and not value.strip() else value


@router.patch("/sales/{sale_id}")
def correct_sale(sale_id: str, body: SaleCorrection, request: Request, user=Depends(principal), db=Depends(get_db)):
    if db.get(Role, user.role_id).name not in {"Administrator", "Operations Manager"}:
        raise HTTPException(403, "Only backend management can correct a sale")
    row = db.get(SalesRecord, sale_id)
    if not row or not permitted(db, user, row):
        raise HTTPException(404, "Sale not found")
    if router_required(body.order_type, body.model_dump()) and not body.router_serial.strip():
        raise HTTPException(422, "Router serial is required for on-spot home wireless sales")
    old = {"order_type": row.order_type, "router_serial": row.details.get("router_serial", ""), "router_fulfilment": row.details.get("router_fulfilment") or "ON_SPOT"}
    row.order_type = body.order_type
    row.details = {**row.details, "router_serial": body.router_serial.strip(), "router_fulfilment": body.router_fulfilment}
    create_call_tasks(db, row)
    audit(db, user, "Sale Corrected", row.id, row.agent_id, old=old,
          new={"order_type": row.order_type, "router_serial": body.router_serial.strip(), "router_fulfilment": body.router_fulfilment},
          reason=body.reason.strip(), request=request)
    db.commit()
    return sale_view(db, row)


class FeedbackCreate(BusinessInput):
    agent_id: str
    customer_name: str = Field(default="", max_length=120)
    contact_number: str = Field(default="", max_length=40)
    product_suggested: str = Field(min_length=2, max_length=160)
    feedback: str = Field(min_length=3, max_length=1000)
    rejection_reason: str = Field(min_length=3, max_length=300)


@router.get("/feedback")
def feedback(user=Depends(principal), db=Depends(get_db)):
    rows = db.scalars(scoped(db, user, NoSaleFeedback).order_by(NoSaleFeedback.created_at.desc()).limit(1000)).all()
    return [{"id": row.id, "created_at": row.created_at,
             "agent_id": row.agent_id, "agent": db.get(User, db.get(Agent, row.agent_id).user_id).name,
             "branch_id": row.branch_id, "customer_name": row.customer_name,
             "contact": "•••• " + cipher.decrypt(row.contact_encrypted.encode()).decode()[-4:] if row.contact_encrypted else "Not recorded",
             "product_suggested": row.product_suggested, "feedback": row.feedback,
             "rejection_reason": row.rejection_reason} for row in rows]


@router.post("/feedback", status_code=201)
def create_feedback(body: FeedbackCreate, request: Request, user=Depends(principal), db=Depends(get_db)):
    can_record(db, user)
    if "read" not in permissions(db, user):
        raise HTTPException(403, "Your role cannot record feedback")
    if db.get(Role, user.role_id).name == "Sales Manager":
        raise HTTPException(403, "Sales Managers have reporting access")
    assert_agent(db, user, body.agent_id)
    agent, outlet = assignment(db, body.agent_id)
    row = NoSaleFeedback(
        agent_id=agent.id, leader_id=agent.leader_id, branch_id=outlet.branch_id,
        customer_name=body.customer_name.strip(),
        contact_encrypted=cipher.encrypt(body.contact_number.strip().encode()).decode() if body.contact_number.strip() else "",
        product_suggested=body.product_suggested.strip(), feedback=body.feedback.strip(),
        rejection_reason=body.rejection_reason.strip(),
    )
    db.add(row)
    db.flush()
    audit(db, user, "No Sale Feedback Recorded", row.id, row.agent_id, new={"product": row.product_suggested}, request=request)
    db.commit()
    return {"id": row.id, "recorded": True}


class TargetWrite(BusinessInput):
    agent_id: str
    period: str
    order_type: Literal["ALL", "NEW", "MNP", "P2P", "HW", "ELIFE", "WASEL", "VISITOR"] = "ALL"
    daily_target: int = Field(ge=0, le=1000)
    monthly_target: int = Field(ge=0, le=10000)


@router.get("/targets")
def targets(period: str = "", user=Depends(principal), db=Depends(get_db)):
    if "read" not in permissions(db, user):
        raise HTTPException(403, "Your role cannot view targets")
    query = select(SalesTarget)
    role = db.get(Role, user.role_id).name
    if role == "Field Agent":
        agent = db.scalar(select(Agent).where(Agent.user_id == user.id))
        query = query.where(SalesTarget.agent_id == (agent.id if agent else ""))
    elif role == "Team Leader":
        query = query.where(SalesTarget.agent_id.in_(select(Agent.id).where(Agent.leader_id == user.id)))
    elif role == "Branch Manager":
        query = query.where(SalesTarget.agent_id.in_(select(Agent.id).join(Outlet).where(Outlet.branch_id == user.branch_id)))
    if period:
        query = query.where(SalesTarget.period == period)
    rows = db.scalars(query.order_by(SalesTarget.period.desc(), SalesTarget.created_at.desc())).all()
    return [{"id": row.id, "agent_id": row.agent_id, "agent": db.get(User, db.get(Agent, row.agent_id).user_id).name,
             "period": row.period, "order_type": row.order_type,
             "daily_target": row.daily_target, "monthly_target": row.monthly_target} for row in rows]


@router.put("/targets")
def save_target(body: TargetWrite, request: Request, user=Depends(principal), db=Depends(get_db)):
    role = db.get(Role, user.role_id).name
    if role not in {"Administrator", "Operations Manager", "Team Leader"}:
        raise HTTPException(403, "Only management or the assigned team leader can set targets")
    if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", body.period):
        raise HTTPException(422, "Period must use YYYY-MM")
    assert_agent(db, user, body.agent_id)
    agent = agent_for_update(db, body.agent_id)
    # A transfer may have completed since the earlier scope check.
    assert_agent(db, user, body.agent_id)
    if role == "Team Leader" and (not agent or agent.leader_id != user.id):
        raise HTTPException(404, "Agent not in your team")
    if not agent or agent.employment_status != "ACTIVE":
        raise HTTPException(422, "Choose an active agent when setting targets")
    row = db.scalar(select(SalesTarget).where(SalesTarget.agent_id == body.agent_id,
        SalesTarget.period == body.period, SalesTarget.order_type == body.order_type).with_for_update())
    if not row:
        row = SalesTarget(agent_id=body.agent_id, period=body.period, order_type=body.order_type,
                          set_by=user.id)
        db.add(row)
    old = {"daily": row.daily_target, "monthly": row.monthly_target}
    row.daily_target, row.monthly_target, row.set_by = body.daily_target, body.monthly_target, user.id
    db.flush()
    audit(db, user, "Sales Target Set", row.id, body.agent_id, old=old,
          new={"daily": body.daily_target, "monthly": body.monthly_target}, request=request)
    db.commit()
    return {"id": row.id, "saved": True}


class CallWrite(BusinessInput):
    stage: Literal["TELE_VERIFICATION", "WELCOME_CALL"]
    outcome: Literal["REACHED", "NO_ANSWER", "RETRY", "PASSED", "FAILED"]
    remark: str = Field(min_length=3, max_length=1000)


def call_stage_access(db, user, stage):
    role = db.get(Role, user.role_id).name
    required = "call.tele.write" if stage == "TELE_VERIFICATION" else "call.welcome.write"
    granted = permissions(db, user)
    if not ((role in {"Administrator", "Operations Manager", "Compliance Officer"} and "compliance.write" in granted) or required in granted):
        raise HTTPException(403, "Your role cannot record this call stage")


def visible_call_rows(db, user):
    role = db.get(Role, user.role_id).name
    if role in {"Tele Verification Officer", "Welcome Call Officer"}:
        required = "call.tele.read" if role == "Tele Verification Officer" else "call.welcome.read"
        if required not in permissions(db, user):
            raise HTTPException(403, "Your role cannot view this queue")
        query = select(SalesRecord)
        if user.branch_id:
            query = query.where(SalesRecord.branch_id == user.branch_id)
    else:
        query = scoped(db, user, SalesRecord)
    return db.scalars(query.where(SalesRecord.order_type.in_(CALL_ORDER_TYPES)).order_by(SalesRecord.created_at.desc()).limit(1000)).all()


def task_stage_filter(db, user):
    role = db.get(Role, user.role_id).name
    if role == "Tele Verification Officer":
        return "TELE_VERIFICATION"
    if role == "Welcome Call Officer":
        return "WELCOME_CALL"
    return None


def call_timing(row, stage):
    sale_day = business_date(row.created_at)
    deferred = row.details.get("tele_deferred_until")
    due = sale_day + timedelta(days=2) if stage == "WELCOME_CALL" else sale_day
    if stage == "TELE_VERIFICATION" and deferred == (sale_day + timedelta(days=1)).isoformat():
        due = sale_day + timedelta(days=1)
    due_at = datetime.combine(due + timedelta(days=1), datetime.min.time(), ZoneInfo("Asia/Dubai")).astimezone(timezone.utc)
    return {"due_date": due.isoformat(), "due_at": due_at.isoformat(),
            "deferred": stage == "TELE_VERIFICATION" and due != sale_day,
            "overdue": business_date(now()) > due}


def call_state(row, task, tele):
    if row.status == "CANCELLED":
        return "CANCELLED"
    if task.stage == "WELCOME_CALL" and task.status not in {"COMPLETED", "SKIPPED"}:
        return "BLOCKED" if row.status != "CLOSED" or not tele or tele.status != "COMPLETED" else "PENDING" if task.status == "BLOCKED" else task.status
    return task.status


def call_task_view(db, row, task, tele):
    state = call_state(row, task, tele)
    timing = call_timing(row, task.stage)
    tele_timing = call_timing(row, "TELE_VERIFICATION")
    return {**sale_view(db, row), "id": task.id, "sale_id": row.id, "stage": task.stage,
            "status": state, "last_outcome": task.last_outcome or "Not started",
            **timing, "overdue": timing["overdue"] and state not in {"COMPLETED", "SKIPPED", "CANCELLED"},
            "activation_blocked": task.stage == "WELCOME_CALL" and row.status != "CLOSED",
            "sequence_warning": "Activation recorded before tele-verification" if row.status == "CLOSED" and not (tele and tele.status == "COMPLETED") and (not tele_timing["deferred"] or tele_timing["overdue"]) else "",
            "updated_at": task.updated_at}


@router.get("/call-tasks")
def call_tasks(user=Depends(principal), db=Depends(get_db)):
    rows = visible_call_rows(db, user)
    sale_ids = [row.id for row in rows]
    if not sale_ids:
        return []
    tasks = db.scalars(select(SalesCallTask).where(SalesCallTask.sale_id.in_(sale_ids))).all()
    tele_by_sale = {task.sale_id: task for task in tasks if task.stage == "TELE_VERIFICATION"}
    sales_by_id = {row.id: row for row in rows}
    restricted_stage = task_stage_filter(db, user)
    tasks = [task for task in tasks if not restricted_stage or task.stage == restricted_stage]
    tasks.sort(key=lambda task: (task.status not in {"PENDING", "FAILED"}, task.created_at))
    return [call_task_view(db, sales_by_id[task.sale_id], task, tele_by_sale.get(task.sale_id)) for task in tasks]


@router.get("/call-tasks/summary")
def call_task_summary(user=Depends(principal), db=Depends(get_db)):
    tasks = call_tasks(user, db)
    return {"actionable": sum(task["status"] in {"PENDING", "FAILED"} for task in tasks),
            "pending_tele": sum(task["stage"] == "TELE_VERIFICATION" and task["status"] == "PENDING" for task in tasks),
            "pending_welcome": sum(task["stage"] == "WELCOME_CALL" and task["status"] == "PENDING" for task in tasks),
            "blocked_welcome": sum(task["stage"] == "WELCOME_CALL" and task["status"] == "BLOCKED" for task in tasks),
            "overdue": sum(task["overdue"] for task in tasks),
            "sequence_warnings": len({task["sale_id"] for task in tasks if task["sequence_warning"]})}


@router.get("/calls")
def call_queue(user=Depends(principal), db=Depends(get_db)):
    rows = visible_call_rows(db, user)
    attempts = db.scalars(select(CallAttempt).where(CallAttempt.sale_id.in_([r.id for r in rows])).order_by(CallAttempt.created_at)).all()
    restricted = task_stage_filter(db, user)
    by_sale = {}
    for attempt in attempts:
        if restricted and restricted != attempt.stage:
            continue
        by_sale.setdefault(attempt.sale_id, {}).setdefault(attempt.stage, []).append({
            "id": attempt.id, "outcome": attempt.outcome, "remark": attempt.remark,
            "at": attempt.created_at, "actor": db.get(User, attempt.actor_id).name,
        })
    return [{**sale_view(db, row), "tele_verification": by_sale.get(row.id, {}).get("TELE_VERIFICATION", []),
             "welcome_call": by_sale.get(row.id, {}).get("WELCOME_CALL", [])} for row in rows]


@router.post("/sales/{sale_id}/calls", status_code=201)
def record_call(sale_id: str, body: CallWrite, request: Request, user=Depends(principal), db=Depends(get_db)):
    call_stage_access(db, user, body.stage)
    row = db.get(SalesRecord, sale_id)
    if not row or not permitted(db, user, row):
        raise HTTPException(404, "Sale not found")
    if row.status == "CANCELLED":
        raise HTTPException(409, "Calls cannot be recorded for a cancelled sale")
    if row.order_type not in CALL_ORDER_TYPES:
        raise HTTPException(409, "This product does not require these call stages")
    task = db.scalar(select(SalesCallTask).where(SalesCallTask.sale_id == row.id,
                     SalesCallTask.stage == body.stage).with_for_update())
    if not task:
        create_call_tasks(db, row)
        db.flush()
        task = db.scalar(select(SalesCallTask).where(SalesCallTask.sale_id == row.id,
                         SalesCallTask.stage == body.stage).with_for_update())
    tele = db.scalar(select(SalesCallTask).where(SalesCallTask.sale_id == row.id,
                     SalesCallTask.stage == "TELE_VERIFICATION"))
    if body.stage == "WELCOME_CALL" and (row.status != "CLOSED" or not tele or tele.status != "COMPLETED"):
        raise HTTPException(409, "Complete tele-verification and record activation before the welcome call")
    if task.status in {"COMPLETED", "SKIPPED"}:
        raise HTTPException(409, "This call stage is already complete")
    attempt = CallAttempt(sale_id=row.id, stage=body.stage, outcome=body.outcome,
                          remark=body.remark.strip(), actor_id=user.id)
    db.add(attempt)
    db.flush()
    task.last_outcome, task.updated_at = body.outcome, now()
    task.status = ("COMPLETED" if body.outcome == "PASSED" else "FAILED" if body.outcome == "FAILED" else "PENDING") if body.stage == "TELE_VERIFICATION" else ("COMPLETED" if body.outcome in {"PASSED", "REACHED"} else "FAILED" if body.outcome == "FAILED" else "PENDING")
    if body.stage == "TELE_VERIFICATION" and task.status == "COMPLETED":
        welcome = db.scalar(select(SalesCallTask).where(SalesCallTask.sale_id == row.id,
                            SalesCallTask.stage == "WELCOME_CALL").with_for_update())
        if welcome and welcome.status == "BLOCKED" and row.status == "CLOSED":
            welcome.status, welcome.updated_at = "PENDING", now()
    audit(db, user, "Sales Call Recorded", row.id, row.agent_id,
          new={"stage": body.stage, "outcome": body.outcome}, request=request)
    db.commit()
    return {"id": attempt.id, "recorded": True}


class CallSkip(BusinessInput):
    reason: str = Field(min_length=5, max_length=300)


@router.post("/sales/{sale_id}/calls/skip-tele")
def skip_tele(sale_id: str, body: CallSkip, request: Request, user=Depends(principal), db=Depends(get_db)):
    if db.get(Role, user.role_id).name not in {"Administrator", "Operations Manager"}:
        raise HTTPException(403, "Only backend management can waive tele-verification")
    row = db.get(SalesRecord, sale_id)
    if not row or not permitted(db, user, row):
        raise HTTPException(404, "Sale not found")
    raise HTTPException(409, "Tele-verification is required. Record a technical postponement to the next day instead.")


@router.post("/sales/{sale_id}/calls/defer-tele")
def defer_tele(sale_id: str, body: CallSkip, request: Request, user=Depends(principal), db=Depends(get_db)):
    call_stage_access(db, user, "TELE_VERIFICATION")
    row = db.scalar(select(SalesRecord).where(SalesRecord.id == sale_id).with_for_update())
    if not row or not permitted(db, user, row):
        raise HTTPException(404, "Sale not found")
    if row.status == "CANCELLED" or row.order_type not in CALL_ORDER_TYPES:
        raise HTTPException(409, "This sale cannot be postponed for tele-verification")
    sale_day, today = business_date(row.created_at), business_date(now())
    due_day = sale_day + timedelta(days=1)
    if today < sale_day or today > due_day:
        raise HTTPException(409, "Technical postponement is limited to the day after the sale")
    task = db.scalar(select(SalesCallTask).where(SalesCallTask.sale_id == sale_id,
                     SalesCallTask.stage == "TELE_VERIFICATION").with_for_update())
    if not task or task.status in {"COMPLETED", "SKIPPED"}:
        raise HTTPException(409, "Tele-verification cannot be postponed now")
    if row.details.get("tele_deferred_until"):
        raise HTTPException(409, "The next-day postponement has already been recorded")
    row.details = {**row.details, "tele_deferred_until": due_day.isoformat(),
                   "tele_defer_reason": body.reason.strip(), "tele_deferred_by": user.id,
                   "tele_deferred_at": now().isoformat()}
    task.updated_at = now()
    audit(db, user, "Tele-verification technically postponed", row.id, row.agent_id,
          new={"due_date": due_day.isoformat()}, reason=body.reason.strip(), request=request)
    db.commit()
    return {"status": task.status, "due_date": due_day.isoformat(), "deferred": True}


class StatusFile(BusinessInput):
    filename: str = Field(max_length=120)
    content_base64: str = Field(max_length=2_800_000)
    apply: bool = False


def parse_status_file(body, columns=("sale_id", "current_status", "new_status")):
    try:
        content = base64.b64decode(body.content_base64, validate=True)
        if len(content) > 2_000_000:
            raise ValueError("File exceeds 2 MB")
        if body.filename.lower().endswith(".csv"):
            rows = list(islice(csv.reader(io.StringIO(content.decode("utf-8-sig"))), 503))
        elif body.filename.lower().endswith(".xlsx"):
            with ZipFile(io.BytesIO(content)) as archive:
                members = archive.infolist()
                if len(members) > 100 or sum(member.file_size for member in members) > 20_000_000:
                    raise ValueError("Workbook too large")
            book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            rows = list(islice(book.active.values, 503))
            book.close()
        else:
            raise ValueError("Use CSV or XLSX")
    except Exception:
        raise HTTPException(422, "Could not read the file. Use CSV or Excel under 2 MB.") from None
    if not 1 <= len(rows) - 1 <= 500 or [str(v or "").strip().lower() for v in rows[0]][:len(columns)] != list(columns):
        raise HTTPException(422, "Use 1–500 rows with " + ",".join(columns) + " columns")
    return rows[1:]


@router.post("/status-file")
def status_file(body: StatusFile, request: Request, user=Depends(principal), db=Depends(get_db)):
    if db.get(Role, user.role_id).name not in {"Administrator", "Operations Manager"}:
        raise HTTPException(403, "Only backend management can update sales statuses")
    if body.apply:
        db.execute(select(Role).where(Role.name == "Administrator").with_for_update()).scalar_one()
    errors, changes, seen = [], [], set()
    for number, values in enumerate(parse_status_file(body), 2):
        if len(values) < 3:
            errors.append(f"Row {number}: expected sale_id, current_status and new_status")
            continue
        sale_id, current_status, status = [str(value or "").strip() for value in values[:3]]
        if sale_id in seen:
            errors.append(f"Row {number}: duplicate sale ID")
            continue
        seen.add(sale_id)
        row = db.get(SalesRecord, sale_id)
        if not row:
            errors.append(f"Row {number}: unmatched sale ID")
        elif current_status != row.status:
            errors.append(f"Row {number}: status changed since download; refresh the file")
        elif status not in SALE_STATUSES:
            errors.append(f"Row {number}: use IN_PROGRESS, CLOSED or CANCELLED")
        elif status == "CLOSED" and (row.order_type not in ORDER_TYPES or router_required(row.order_type, row.details) and not row.details.get("router_serial")):
            errors.append(f"Row {number}: correct the product and required router serial before closing")
        elif status == "CLOSED" and row.capture_id and db.get(KycCapture, row.capture_id).status != "VERIFIED":
            errors.append(f"Row {number}: backend evidence verification is pending")
        else:
            changes.append({"sale_id": sale_id, "from": row.status, "to": status})
    if body.apply and errors:
        raise HTTPException(422, {"errors": errors})
    if body.apply:
        for change in changes:
            row = db.get(SalesRecord, change["sale_id"])
            if row.status != change["to"]:
                change_sale_status(db, row, change["to"], user, request)
                audit(db, user, "Sale Status Imported", row.id, row.agent_id,
                      old={"status": change["from"]}, new={"status": row.status}, request=request)
        db.commit()
    return {"changes": changes, "errors": errors, "applied": body.apply and not errors,
            "file_hash": hashlib.sha256(base64.b64decode(body.content_base64)).hexdigest()}


TARGET_COLUMNS = ("employee_id", "period", "order_type", "current_daily", "current_monthly", "daily_target", "monthly_target")


def table_download(columns, rows, filename, format="csv"):
    from fastapi.responses import Response
    safe = [["'" + str(value) if str(value).startswith(("=", "+", "-", "@")) else value for value in row] for row in rows]
    if format == "xlsx":
        from openpyxl import Workbook
        book = Workbook()
        book.active.append(list(columns))
        for row in safe:
            book.active.append(row)
        book.active.freeze_panes = "A2"
        book.active.auto_filter.ref = book.active.dimensions
        out = io.BytesIO()
        book.save(out)
        content, media = out.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:
        out = io.StringIO()
        writer = csv.writer(out)
        writer.writerow(columns)
        writer.writerows(safe)
        content, media = out.getvalue(), "text/csv"
    return Response(content, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename}.{format}"'})


@router.get("/targets/template")
def target_template(period: str = "", format: Literal["csv", "xlsx"] = "xlsx", user=Depends(principal), db=Depends(get_db)):
    if db.get(Role, user.role_id).name not in {"Administrator", "Operations Manager", "Team Leader"}:
        raise HTTPException(403, "Your role cannot upload targets")
    period = period or business_date(now()).strftime("%Y-%m")
    if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", period):
        raise HTTPException(422, "Period must use YYYY-MM")
    from .security import visible_agents
    rows = []
    for agent in db.scalars(select(Agent).where(Agent.id.in_(visible_agents(db, user)), Agent.employment_status == "ACTIVE")):
        existing = db.scalars(select(SalesTarget).where(SalesTarget.agent_id == agent.id, SalesTarget.period == period)).all()
        for row in existing:
            rows.append([agent.employee_id, period, row.order_type, row.daily_target, row.monthly_target, row.daily_target, row.monthly_target])
        if not existing:
            rows.append([agent.employee_id, period, "ALL", "", "", 0, 0])
    return table_download(TARGET_COLUMNS, rows, "sales-targets", format)


@router.post("/targets/file")
def target_file(body: StatusFile, request: Request, user=Depends(principal), db=Depends(get_db)):
    if db.get(Role, user.role_id).name not in {"Administrator", "Operations Manager", "Team Leader"}:
        raise HTTPException(403, "Your role cannot upload targets")
    from .security import visible_agents
    allowed = set(visible_agents(db, user))
    if body.apply:
        db.execute(select(Role).where(Role.name == "Administrator").with_for_update()).scalar_one()
    errors, changes, seen = [], [], set()
    for number, values in enumerate(parse_status_file(body, TARGET_COLUMNS), 2):
        try:
            if len(values) < 7:
                raise ValueError("expected all seven target columns")
            employee, period, product = [str(v or "").strip() for v in values[:3]]
            key = (employee, period, product)
            if key in seen:
                raise ValueError("duplicate agent, month and product")
            seen.add(key)
            agent_query = select(Agent).where(Agent.employee_id == employee, Agent.id.in_(allowed), Agent.employment_status == "ACTIVE")
            agent = db.scalar(agent_query)
            if agent and body.apply:
                agent = agent_for_update(db, agent.id)
            if not agent or agent.employment_status != "ACTIVE" or agent.id not in visible_agents(db, user):
                raise ValueError("agent is unavailable or outside your access")
            if not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", period) or product not in ("ALL", *ORDER_TYPES):
                raise ValueError("invalid month or product")
            def integer(value, limit):
                text = str(value).strip()
                if not re.fullmatch(r"\d+", text) or int(text) > limit:
                    raise ValueError("targets must be whole numbers within the allowed limit")
                return int(text)
            daily, monthly = integer(values[5], 1000), integer(values[6], 10000)
            row = db.scalar(select(SalesTarget).where(SalesTarget.agent_id == agent.id, SalesTarget.period == period, SalesTarget.order_type == product).with_for_update())
            previous = [row.daily_target, row.monthly_target] if row else [None, None]
            supplied = [None if v is None or str(v).strip() == "" else integer(v, 10000) for v in values[3:5]]
            if previous != supplied:
                raise ValueError("target changed since download; download a fresh file")
            changes.append({"agent_id": agent.id, "employee_id": employee, "period": period, "order_type": product,
                            "from": previous, "daily_target": daily, "monthly_target": monthly})
        except ValueError as error:
            errors.append(f"Row {number}: {error}")
    if body.apply and errors:
        raise HTTPException(422, {"errors": errors})
    if body.apply:
        for change in changes:
            row = db.scalar(select(SalesTarget).where(SalesTarget.agent_id == change["agent_id"], SalesTarget.period == change["period"], SalesTarget.order_type == change["order_type"]))
            if not row:
                row = SalesTarget(agent_id=change["agent_id"], period=change["period"], order_type=change["order_type"], set_by=user.id)
                db.add(row)
            row.daily_target, row.monthly_target, row.set_by = change["daily_target"], change["monthly_target"], user.id
            try:
                db.flush()
            except IntegrityError:
                db.rollback()
                raise HTTPException(409, "Targets changed while applying. Download a fresh file.") from None
            audit(db, user, "Sales Target Imported", row.id, row.agent_id, old={"targets": change["from"]},
                  new={"daily": row.daily_target, "monthly": row.monthly_target}, request=request)
        audit(db, user, "Target File Applied", body.filename, new={"rows": len(changes), "hash": hashlib.sha256(base64.b64decode(body.content_base64)).hexdigest()}, request=request)
        db.commit()
    return {"changes": changes, "errors": errors, "applied": body.apply and not errors}


@router.get("/export")
def export_sales(format: Literal["csv", "xlsx"] = "csv", period: str = "", order_type: str = "", status: str = "", branch_id: str = "", agent_id: str = "", leader_id: str = "", from_date: str = "", to_date: str = "", user=Depends(principal), db=Depends(get_db)):
    query = filtered_sales_query(db, user, period, order_type, status, branch_id, agent_id, leader_id, from_date, to_date)
    rows = [sale_view(db, row) for row in db.scalars(query.order_by(SalesRecord.created_at.desc()))]
    columns = ["sale_id", "current_status", "new_status", "created_at", "agent", "leader", "sales_manager", "branch", "outlet", "order_type", "customer_name", "document", "nationality", "plan_name", "account_number", "request_id", "sim_serial", "router_serial", "advance_transaction_number", "sr_number", "alternate_number", "msisdn", "employee_id", "assignment_effective_at", "monthly_cost", "prepayment", "tele_status", "tele_remark", "tele_last_attempt", "welcome_status", "welcome_remark", "welcome_last_attempt"]
    attempts = db.scalars(select(CallAttempt).where(CallAttempt.sale_id.in_([r["id"] for r in rows])).order_by(CallAttempt.created_at, CallAttempt.id)).all()
    latest = {(attempt.sale_id, attempt.stage): attempt for attempt in attempts}
    tasks = {(task.sale_id, task.stage): task for task in db.scalars(select(SalesCallTask).where(SalesCallTask.sale_id.in_([r["id"] for r in rows])))}
    output = []
    for row in rows:
        values = [row["id"], row["status"], row["status"]] + [row.get(key) or "Not recorded" for key in columns[3:-6]]
        for stage in CALL_STAGES:
            attempt, task = latest.get((row["id"], stage)), tasks.get((row["id"], stage))
            values += [task.status if task else "Not recorded", attempt.remark if attempt else "Not recorded", attempt.created_at if attempt else "Not recorded"]
        output.append(values)
    return table_download(columns, output, "sales-records", format)
