"""Daily carrier-report reconciliation, independent of external activation."""
import base64
import csv
import hashlib
import io
import json
import re
from datetime import date, datetime, timedelta
from itertools import islice
from zipfile import ZipFile

import jwt
from fastapi import APIRouter, Depends, HTTPException, Request, Response
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import DateTime, ForeignKey, JSON, String, Text, UniqueConstraint, or_, select, text
from sqlalchemy.orm import Mapped, mapped_column
from sqlalchemy.exc import OperationalError

from .db import Entity, SalesRecord, KycCapture, User, Role, Agent, Branch, get_db, now, business_date
from .security import SECRET, cipher, permissions, principal, visible_agents
from .services import audit

router = APIRouter(prefix="/api/sales-management/sr", tags=["SR verification"])
WRITERS = {"Administrator", "Operations Manager", "Compliance Officer"}


class SrBatch(Entity):
    __tablename__ = "sr_batches"
    business_date: Mapped[str] = mapped_column(String(10), index=True)
    fingerprint: Mapped[str] = mapped_column(String(64), unique=True)
    filename: Mapped[str] = mapped_column(String(180))
    actor_id: Mapped[str] = mapped_column(ForeignKey("users.id"))
    summary: Mapped[dict] = mapped_column(JSON, default=dict)
    report_encrypted: Mapped[str] = mapped_column(Text)


class SrCheck(Entity):
    __tablename__ = "sr_checks"
    sale_id: Mapped[str] = mapped_column(ForeignKey("sales_records.id"), unique=True)
    status: Mapped[str] = mapped_column(String(32), index=True)
    business_date: Mapped[str] = mapped_column(String(10), index=True)
    batch_id: Mapped[str] = mapped_column(ForeignKey("sr_batches.id"))
    verified_at: Mapped[datetime] = mapped_column(DateTime, default=now)
    reason: Mapped[str] = mapped_column(String(300), default="")


class SrNotice(Entity):
    __tablename__ = "sr_notices"
    __table_args__ = (UniqueConstraint("user_id", "sale_id", "batch_id", name="uq_sr_notice_recipient"),)
    user_id: Mapped[str] = mapped_column(ForeignKey("users.id"), index=True)
    sale_id: Mapped[str] = mapped_column(ForeignKey("sales_records.id"), index=True)
    batch_id: Mapped[str] = mapped_column(ForeignKey("sr_batches.id"))
    status: Mapped[str] = mapped_column(String(32))
    reason: Mapped[str] = mapped_column(String(300), default="")


def role(db, user):
    return db.get(Role, user.role_id).name


def access(db, user, write=False):
    allowed = WRITERS if write else WRITERS | {"Sales Manager"}
    granted = permissions(db, user)
    if role(db, user) not in allowed or "read" not in granted or (write and "compliance.write" not in granted):
        raise HTTPException(403, "Only backend staff can reconcile SR reports" if write else "Your role cannot view SR report imports")


def sale_state(db, sale):
    check = db.scalar(select(SrCheck).where(SrCheck.sale_id == sale.id))
    return {
        "status": check.status if check else "PENDING_SR_VERIFICATION",
        "business_date": business_date(sale.created_at).isoformat(),
        "batch_id": check.batch_id if check else None,
        "verified_at": check.verified_at if check else None,
        "reason": check.reason if check else "Daily SR report not checked",
        "sale_status": sale.status,
    }


def normal_sr(value):
    value = str(value or "").strip()
    if not value:
        return ""
    if not re.fullmatch(r"[A-Za-z0-9][A-Za-z0-9 _./-]{0,119}", value):
        raise ValueError("SR numbers must be text references, not formulas")
    return re.sub(r"\s+", "", value).casefold()


class ReportBody(BaseModel):
    business_date: str = Field(default="", max_length=10)
    period: str = Field(default="", max_length=7)
    filename: str = Field(min_length=4, max_length=180)
    content_base64: str = Field(min_length=4, max_length=2_800_000)
    apply: bool = False
    preview_token: str = Field(default="", max_length=4096)


def report_date(value):
    try:
        day = date.fromisoformat(value)
        if day.isoformat() != value or day > business_date():
            raise ValueError()
        return day
    except ValueError:
        raise HTTPException(422, "Choose the sale's business date, not a future date") from None


def report_range(day_value="", period=""):
    if period:
        if day_value or not re.fullmatch(r"20\d{2}-(0[1-9]|1[0-2])", period):
            raise HTTPException(422, "Choose either a sales date or a month in YYYY-MM format")
        import calendar
        year, month = map(int, period.split('-'))
        first = date(year, month, 1)
        if first > business_date_today():
            raise HTTPException(422, "Choose a current or past sales month")
        return first, min(date(year, month, calendar.monthrange(year,month)[1]), business_date_today())
    day = report_date(day_value) if day_value else business_date_today()
    return day, day


def read_report(body):
    try:
        raw = base64.b64decode(body.content_base64, validate=True)
        if len(raw) > 2_000_000:
            raise ValueError()
        if body.filename.lower().endswith(".csv"):
            # A small CSV can contain hundreds of thousands of empty rows.
            # Stop at the first over-limit row instead of allocating them all.
            values = list(islice(csv.reader(io.StringIO(raw.decode("utf-8-sig"))), 5002))
        elif body.filename.lower().endswith(".xlsx"):
            with ZipFile(io.BytesIO(raw)) as archive:
                members = archive.infolist()
                if len(members) > 100 or sum(item.file_size for item in members) > 20_000_000:
                    raise ValueError()
            workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=False)
            try:
                if workbook.active.max_row > 5001 or workbook.active.max_column > 50:
                    raise ValueError()
                values = list(workbook.active.iter_rows(values_only=True))
            finally:
                workbook.close()
        else:
            raise ValueError()
    except (ValueError, TypeError, UnicodeError, OSError, KeyError):
        raise HTTPException(422, "Upload a valid CSV or Excel .xlsx report, up to 5,000 SRs") from None
    except Exception:
        raise HTTPException(422, "The SR report could not be read") from None
    if not values or len(values) > 5001:
        raise HTTPException(422, "Keep the SR report within 5,000 data rows")
    if any(len(row) > 50 for row in values):
        raise HTTPException(422, "Keep the SR report within 50 columns")
    aliases = {"sr": "sr_number", "sr number": "sr_number", "sr no": "sr_number", "sr id": "sr_number", "service request number": "sr_number", "status": "status", "sale status": "status", "request id": "request_id", "request number": "request_id"}
    columns = [aliases.get(re.sub(r"[_\s]+", " ", str(value or "").strip().casefold()), "") for value in values[0]]
    if columns.count("sr_number") != 1 or any(columns.count(key) > 1 for key in ("status", "request_id")):
        raise HTTPException(422, "The report must have one SR number column; status and request ID are optional")
    entries, errors, seen = {}, [], set()
    status_aliases = {"": "", "closed": "CLOSED", "activated": "CLOSED", "cancelled": "CANCELLED", "canceled": "CANCELLED", "in progress": "IN_PROGRESS", "in_progress": "IN_PROGRESS"}
    for number, values_row in enumerate(values[1:], 2):
        if not any(str(value or "").strip() for value in values_row):
            continue
        row = {key: values_row[index] if index < len(values_row) else "" for index, key in enumerate(columns) if key}
        value = row.get("sr_number")
        if not isinstance(value, str):
            if isinstance(value, int) and not isinstance(value, bool) and 0 <= value < 10**15:
                value = str(value)
            else:
                errors.append(f"Row {number}: store SR numbers as text to preserve all digits")
                continue
        try:
            sr = normal_sr(value)
            if not sr or sr in seen:
                errors.append(f"Row {number}: {'duplicate' if sr else 'missing'} SR number")
                continue
            status = status_aliases.get(str(row.get("status") or "").strip().casefold())
            request_id = str(row.get("request_id") or "").strip()
            if status is None or len(request_id) > 120 or request_id.startswith(("=", "+", "@")):
                errors.append(f"Row {number}: use CLOSED, CANCELLED or IN_PROGRESS and a text request ID")
                continue
            seen.add(sr)
            entries[sr] = {"sr_number": value.strip(), "status": status, "request_id": request_id,
                           "row_number": number}
        except ValueError as error:
            errors.append(f"Row {number}: {error}")
    if not entries and not errors:
        errors.append("Add at least one SR number before reconciliation")
    # Row positions aid review but do not change an identical report's identity.
    # Retain fingerprints used before row metadata was added.
    identity = {key: {field: row[field] for field in ("sr_number", "status", "request_id")}
                for key, row in entries.items()}
    fingerprint = hashlib.sha256(((body.period or body.business_date) + "\n" + json.dumps(identity, sort_keys=True)).encode()).hexdigest()
    return entries, errors, fingerprint


def day_sales(db, day, lock=False, last=None):
    # Backend-only import considers every sale on the selected business date.
    from datetime import timezone
    from zoneinfo import ZoneInfo
    start = datetime.combine(day, datetime.min.time(), ZoneInfo("Asia/Dubai")).astimezone(timezone.utc).replace(tzinfo=None)
    end = datetime.combine((last or day) + timedelta(days=1), datetime.min.time(), ZoneInfo("Asia/Dubai")).astimezone(timezone.utc).replace(tzinfo=None)
    query = select(SalesRecord).where(SalesRecord.created_at >= start, SalesRecord.created_at < end).order_by(SalesRecord.id)
    try:
        return list(db.scalars(query.with_for_update(nowait=True) if lock else query))
    except OperationalError as error:
        if getattr(error.orig,"sqlstate",None) not in {"55P03","40P01"}:
            raise
        db.rollback()
        raise HTTPException(409,"Sales are being updated. Preview and apply the SR report again shortly.") from None


def comparison_for(sale, entries, *, report_available=True):
    """Show exact evidence and diagnostics without promoting a request-ID hint to an SR match."""
    recorded_sr = sale.details.get("sr_number") or ""
    recorded_request = (sale.request_id or "").strip()
    try:
        key = normal_sr(recorded_sr)
    except ValueError:
        key = ""
    entry = entries.get(key) if report_available else None
    related = []
    if report_available and not entry and recorded_request:
        related = [dict(row) for row in entries.values()
                   if row.get("request_id") and row["request_id"].strip().casefold() == recorded_request.casefold()][:5]
    report = dict(entry) if entry else None
    if not report_available:
        sr_result = "NOT_CHECKED"
    elif not key:
        sr_result = "NOT_RECORDED"
    else:
        sr_result = "MATCHED" if entry else "NOT_IN_REPORT"
    if not entry:
        request_result = "NOT_COMPARED"
    elif not entry.get("request_id"):
        request_result = "NOT_PROVIDED_IN_REPORT"
    elif not recorded_request:
        request_result = "NOT_RECORDED"
    else:
        request_result = "MATCHED" if entry["request_id"].casefold() == recorded_request.casefold() else "MISMATCH"
    return {"recorded_sr": recorded_sr or None, "recorded_request_id": recorded_request or None,
            "recorded_status": sale.status, "report_row": report, "related_report_rows": related,
            "sr_result": sr_result, "request_id_result": request_result,
            "report_available": report_available}


def changes_for(db, sales, entries):
    counts = {}
    for sale in sales:
        try:
            sr = normal_sr(sale.details.get("sr_number"))
        except ValueError:
            sr = ""
        if sr:
            counts[sr] = counts.get(sr, 0) + 1
    changes = []
    for sale in sales:
        sr_original = sale.details.get("sr_number") or ""
        try:
            sr = normal_sr(sr_original)
        except ValueError:
            sr = ""
        entry = entries.get(sr)
        status, reason = "MATCHED", "SR matched the daily report"
        if not sr:
            status, reason = "PENDING_SR_VERIFICATION", "SR number not recorded on the sale"
        elif counts[sr] > 1:
            status, reason = "MISMATCH", "The same SR appears on multiple sales; resolve the duplicate"
        elif not entry:
            status, reason = "MISMATCH", "SR not found in the selected report"
        elif entry["request_id"] and entry["request_id"].casefold() != (sale.request_id or "").strip().casefold():
            status, reason = "MISMATCH", "SR matched but the report's request ID differs"
        sale_status = sale.status
        if status == "MATCHED" and entry["status"]:
            target = entry["status"]
            if target != "CLOSED" or not sale.capture_id or db.get(KycCapture, sale.capture_id).status == "VERIFIED":
                sale_status = target
            else:
                reason = "SR matched; independent evidence review is still pending"
        current = sale_state(db, sale)
        agent = db.get(Agent, sale.agent_id)
        changes.append({"sale_id": sale.id, "customer_name": sale.customer_name, "branch_id": sale.branch_id,
                        "agent": db.get(User, agent.user_id).name, "branch": db.get(Branch, sale.branch_id).name,
                        "sr_number": sr_original, "from": current["status"], "to": status, "reason": reason,
                        "sale_status_from": sale.status, "sale_status_to": sale_status,
                        "comparison": comparison_for(sale, entries)})
    return changes


def state_hash(db, sales, changes):
    snapshot = [(row.id, row.status_updated_at.isoformat(), row.details, sale_state(db, row)) for row in sales]
    return hashlib.sha256(json.dumps({"sales": snapshot, "changes": changes}, sort_keys=True, default=str).encode()).hexdigest()


@router.get("/template")
def template(business_date: str = "", period: str = "", user=Depends(principal), db=Depends(get_db)):
    access(db, user)
    day, last = report_range(business_date, period)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Monthly SR report" if period else "Daily SR report"
    sheet.append(["sr_number", "status", "request_id"])
    from .exports import add_metadata, filename, style_sheet
    from openpyxl.worksheet.datavalidation import DataValidation
    style_sheet(sheet, ["sr_number", "status", "request_id"], machine_headers=True)
    for column in ("A", "C"):
        sheet.column_dimensions[column].width = 30
        for index in range(2, 5002):
            sheet[f"{column}{index}"].number_format = "@"
    validation = DataValidation(type="list", formula1='"CLOSED,CANCELLED,IN_PROGRESS"', allow_blank=True)
    validation.errorTitle = "Choose a supported status"
    validation.error = "Use CLOSED, CANCELLED or IN_PROGRESS."
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add("B2:B5001")
    add_metadata(workbook, "SR reconciliation template", {"from_date": day, "to_date": last, "maximum_srs": 5000}, ["Keep SR number and request ID as text, preserving leading zeros. Upload the carrier report for the selected day or month. An SR match does not replace independent backend evidence review."])
    output = io.BytesIO()
    workbook.save(output)
    return Response(output.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename("SR report template", "xlsx", day.isoformat() if day else "")}"', "Cache-Control": "private, no-store"})


@router.get("/batches")
def batches(user=Depends(principal), db=Depends(get_db)):
    access(db, user)
    # History needs counts only; do not load up to 100 encrypted reports and
    # thousands of comparison snapshots merely to render this compact list.
    query = select(SrBatch.id, SrBatch.business_date, SrBatch.filename, SrBatch.created_at,
                   User.name.label("actor"), SrBatch.summary["report_period"].as_string().label("report_period"),
                   *(SrBatch.summary[key].as_integer().label(key) for key in ("matched", "mismatch", "pending")))\
        .join(User, User.id == SrBatch.actor_id).order_by(SrBatch.created_at.desc()).limit(100)
    return [{**dict(row), **{key: row[key] or 0 for key in ("matched", "mismatch", "pending")}}
            for row in db.execute(query).mappings()]


def stored_entries(batch):
    try:
        rows = json.loads(cipher.decrypt(batch.report_encrypted.encode()))
        return rows if isinstance(rows, dict) and all(isinstance(row, dict) for row in rows.values()) else None
    except Exception:
        return None


@router.get("/results")
def results(business_date: str = "", period: str = "", branch_id: str = "", user=Depends(principal), db=Depends(get_db)):
    """Current per-sale verification, including the report evidence behind each verdict."""
    access(db, user)
    day, last = report_range(business_date, period)
    allowed = set(visible_agents(db, user))
    sales = [sale for sale in day_sales(db, day, last=last) if sale.agent_id in allowed and
             (not branch_id or sale.branch_id == branch_id)]
    truncated = len(sales) > 5000
    sales = sales[:5000]
    checks = {row.sale_id: row for row in db.scalars(select(SrCheck).where(SrCheck.sale_id.in_([sale.id for sale in sales])))}
    batches_by_id = {row.id: row for row in db.scalars(select(SrBatch).where(SrBatch.id.in_({row.batch_id for row in checks.values()})))}
    saved = {ident: {row["sale_id"]: row for row in batch.summary.get("comparisons", [])}
             for ident, batch in batches_by_id.items()}
    entries = {ident: stored_entries(batch) for ident, batch in batches_by_id.items()}
    items = []
    for sale in sales:
        check = checks.get(sale.id)
        batch = batches_by_id.get(check.batch_id) if check else None
        snapshot = saved.get(check.batch_id, {}).get(sale.id) if check else None
        agent = db.get(Agent, sale.agent_id)
        comparison = (snapshot or {}).get("comparison") or comparison_for(
            sale, entries.get(batch.id) or {} if batch else {},
            report_available=bool(batch and entries.get(batch.id) is not None))
        items.append({"sale_id": sale.id, "customer_name": sale.customer_name,
            "agent": db.get(User, agent.user_id).name, "branch_id": sale.branch_id,
            "branch": db.get(Branch, sale.branch_id).name, "sr_number": sale.details.get("sr_number") or "",
            "to": check.status if check else "PENDING_SR_VERIFICATION",
            "reason": check.reason if check else "Daily SR report not checked",
            "comparison": comparison, "comparison_source": "SAVED_VERIFICATION" if snapshot else "CURRENT_SALE",
            "record_changed": bool(snapshot and (comparison.get("recorded_sr") != (sale.details.get("sr_number") or None)
                or comparison.get("recorded_request_id") != ((sale.request_id or "").strip() or None))),
            "batch_id": batch.id if batch else None, "filename": batch.filename if batch else None,
            "verified_at": check.verified_at if check else None})
    return {"items": items, "business_date": day.isoformat(), "truncated": truncated,
            "summary": {"matched": sum(row["to"] == "MATCHED" for row in items),
                        "mismatch": sum(row["to"] == "MISMATCH" for row in items),
                        "pending": sum(row["to"] == "PENDING_SR_VERIFICATION" for row in items)}}


def business_date_today():
    return business_date()


@router.get("/batches/{batch_id}")
def batch_detail(batch_id: str, branch_id: str = "", user=Depends(principal), db=Depends(get_db)):
    access(db, user)
    batch = db.get(SrBatch, batch_id)
    if not batch:
        raise HTTPException(404, "SR report import not found")
    allowed = set(visible_agents(db, user))
    sales = [sale for sale in day_sales(db, date.fromisoformat(batch.business_date), last=date.fromisoformat(batch.summary.get("range_end") or batch.business_date))
             if sale.agent_id in allowed and (not branch_id or sale.branch_id == branch_id)]
    identifiers = {sale.id for sale in sales}
    rows = batch.summary.get("comparisons")
    historical = isinstance(rows, list)
    if not historical:
        entries = stored_entries(batch)
        if entries is None:
            raise HTTPException(503, "This stored report cannot be read. Upload the report again to compare it.")
        rows = changes_for(db, sales, entries)
    rows = [row for row in rows if row.get("sale_id") in identifiers]
    return {"id": batch.id, "business_date": batch.business_date, "report_period": batch.summary.get("report_period"), "filename": batch.filename,
            "created_at": batch.created_at, "changes": rows,
            "comparison_source": "SAVED_VERIFICATION" if historical else "CURRENT_SALE",
            "summary": {"matched": sum(row["to"] == "MATCHED" for row in rows),
                        "mismatch": sum(row["to"] == "MISMATCH" for row in rows),
                        "pending": sum(row["to"] == "PENDING_SR_VERIFICATION" for row in rows)}}


@router.get("/email-status")
def email_status(user=Depends(principal), db=Depends(get_db)):
    access(db, user)
    from .email_delivery import EmailOutbox, smtp_ready
    from sqlalchemy import func
    counts = dict(db.execute(select(EmailOutbox.status, func.count()).group_by(EmailOutbox.status)).all())
    return {"configured": smtp_ready(), "statuses": counts,
            "queued": counts.get("QUEUED", 0) + counts.get("RETRY", 0),
            "accepted": counts.get("SENT", 0), "failed": counts.get("FAILED", 0), "skipped": counts.get("SKIPPED", 0)}


@router.post("/file")
def reconcile(body: ReportBody, request: Request, user=Depends(principal), db=Depends(get_db)):
    access(db, user, True)
    day, last = report_range(body.business_date, body.period)
    entries, errors, fingerprint = read_report(body)
    if body.apply and errors:
        raise HTTPException(422, {"errors": errors})
    if body.apply and db.bind.dialect.name == "postgresql":
        if not db.scalar(text("SELECT pg_try_advisory_xact_lock(82831017)")):
            raise HTTPException(409,"Another SR import is running. Retry shortly.")
    sales = day_sales(db, day, body.apply, last=last)
    changes = changes_for(db, sales, entries)
    summary = {"matched": sum(item["to"] == "MATCHED" for item in changes),
               "mismatch": sum(item["to"] == "MISMATCH" for item in changes),
               "pending": sum(item["to"] == "PENDING_SR_VERIFICATION" for item in changes)}
    snapshot = state_hash(db, sales, changes)

    def validate_preview():
        try:
            proof = jwt.decode(body.preview_token, SECRET, algorithms=["HS256"], audience="relay-sr-preview")
            if proof["sub"] != user.id or proof["file"] != fingerprint or proof["state"] != snapshot or proof["day"] != day.isoformat():
                raise ValueError()
        except (ValueError, KeyError, jwt.PyJWTError):
            raise HTTPException(409, "The report or sales changed. Preview the file again before applying.") from None

    if body.apply and body.preview_token:
        validate_preview()
    existing = db.scalar(select(SrBatch).where(or_(
        SrBatch.fingerprint == fingerprint,
        SrBatch.summary["report_fingerprint"].as_string() == fingerprint,
    )).order_by(SrBatch.created_at.desc()).limit(1))
    # An identical report is only a replay when it still describes every sale.
    # Late submissions, corrected SRs and subsequent reports must be checked
    # again rather than being hidden by the file's original import identifier.
    if existing and not errors:
        checks = {check.sale_id: check for check in db.scalars(
            select(SrCheck).where(SrCheck.sale_id.in_([sale.id for sale in sales]))
        )}
        unchanged = all(
            item["sale_id"] in checks
            and checks[item["sale_id"]].batch_id == existing.id
            and checks[item["sale_id"]].status == item["to"]
            and checks[item["sale_id"]].reason == item["reason"]
            and item["sale_status_from"] == item["sale_status_to"]
            for item in changes
        )
        if unchanged:
            return {"applied": True, "already_applied": True, "batch_id": existing.id,
                    "errors": [], "changes": changes, "summary": summary, "preview_token": ""}
    if body.apply:
        if not body.preview_token:
            validate_preview()
        batch_fingerprint = hashlib.sha256((fingerprint + "\n" + snapshot).encode()).hexdigest() if existing else fingerprint
        batch = SrBatch(business_date=day.isoformat(), fingerprint=batch_fingerprint, filename=body.filename,
                        actor_id=user.id, summary={**summary, "report_period":body.period, "range_end":last.isoformat(), "report_fingerprint": fingerprint, "comparisons": changes},
                        report_encrypted=cipher.encrypt(json.dumps(entries).encode()).decode())
        db.add(batch)
        db.flush()
        from .sales_management import change_sale_status
        from .email_delivery import queue_sr_email
        by_id = {sale.id: sale for sale in sales}
        for item in changes:
            sale = by_id[item["sale_id"]]
            if sale.status != item["sale_status_to"]:
                change_sale_status(db, sale, item["sale_status_to"], user, request)
                audit(db, user, "Sale Status Reconciled", sale.id, sale.agent_id, old={"status": item["sale_status_from"]}, new={"status": sale.status}, reason="Daily SR report", request=request)
            check = db.scalar(select(SrCheck).where(SrCheck.sale_id == sale.id))
            old_status = check.status if check else "PENDING_SR_VERIFICATION"
            if check is None:
                check = SrCheck(sale_id=sale.id)
                db.add(check)
            check.status, check.business_date, check.batch_id = item["to"], business_date(sale.created_at).isoformat(), batch.id
            check.verified_at, check.reason = now(), item["reason"]
            # Do not resend an unchanged mismatch whenever the same day is rechecked.
            first_check = old_status == "PENDING_SR_VERIFICATION" and item["from"] == "PENDING_SR_VERIFICATION" and check.created_at is None
            if old_status != check.status or first_check:
                agent = db.get(Agent, sale.agent_id)
                for recipient in {agent.user_id, sale.leader_id} - {None}:
                    db.add(SrNotice(user_id=recipient, sale_id=sale.id, batch_id=batch.id, status=check.status, reason=check.reason))
                    if check.status != "MATCHED":
                        queue_sr_email(db, sale, recipient, batch.id, check.status, check.reason)
            audit(db, user, "Sale SR Reconciled", sale.id, sale.agent_id,
                  old={"sr_status": old_status}, new={"sr_status": check.status, "batch_id": batch.id}, reason=check.reason, request=request)
        audit(db, user, "Daily SR Report Applied", batch.id, new={"business_date": day.isoformat(), **summary}, request=request)
        db.commit()
        return {"applied": True, "errors": [], "changes": changes, "summary": summary, "batch_id": batch.id, "preview_token": ""}
    proof = jwt.encode({"aud": "relay-sr-preview", "sub": user.id, "file": fingerprint,
                        "state": snapshot, "day": day.isoformat(), "exp": now() + timedelta(minutes=15)}, SECRET, algorithm="HS256")
    return {"applied": False, "errors": errors, "changes": changes, "summary": summary, "preview_token": proof}
