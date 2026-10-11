"""Consistent, typed and safe downloads for authorized application data."""

import csv
import io
import json
import re
from datetime import date, datetime, timezone
from decimal import Decimal
from pathlib import Path
from xml.sax.saxutils import escape

from fastapi.responses import Response
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

BRAND = "Qanawat Connect Solutions LLC"
BURGUNDY = "761B3A"
MISSING = "Not recorded"
ASSETS = Path(__file__).with_name("assets")
LOGO = ASSETS / "qanawat-logo.png"
XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
TEXT_KEYS = {"id", "sale_id", "iccid", "sim_serial", "serial", "sim_identifier", "msisdn", "phone", "phone_number", "employee_id", "document", "document_number", "account", "account_number", "request_id", "sr_id", "sr_number", "router_serial", "advance_transaction_number", "alternate_number", "reference"}
MONEY_KEYS = {"amount", "amount_aed", "monthly_cost", "monthly_charge", "prepayment", "total_paid", "price", "plan_price", "commission", "incentive", "component_amount"}


def require_scoped_export(db, user):
    """Sales agents can download their own scoped data; management needs report access."""
    from .db import Role
    from .security import require
    role = db.get(Role, user.role_id)
    require(db, user, "read" if role and role.name == "Field Agent" else "report.read")


def filename(title, extension, reference="", generated=None):
    """ASCII, bounded filenames cannot inject download response headers."""
    stamp = (generated or datetime.now(timezone.utc)).strftime("%Y-%m-%d")
    parts = ["qanawat", title, reference, stamp]
    stem = "-".join(filter(None, (re.sub(r"[^a-z0-9]+", "-", str(part).lower()).strip("-")[:64] for part in parts)))
    return f"{stem}.{extension}"


def label(key):
    replacements = {"agent": "Sales agent", "sales_agent": "Sales agent", "agents": "Sales agents", "leader": "Team leader", "outlet": "Branch", "msisdn": "Phone number (MSISDN)", "sr_number": "SR number", "sr_id": "SR number", "iccid": "ICCID", "aht": "Average handling time (minutes)", "ekyc_rate": "eKYC rate (%)", "low_stock": "Low stock", "monthly_cost": "Monthly charge (AED)", "prepayment": "Order prepayment (AED)", "amount": "Amount (AED)", "component_amount": "Component amount (AED)"}
    if key in replacements:
        return replacements[key]
    if str(key).startswith("agent_"):
        return "Sales agent " + label(str(key)[6:])
    if str(key).startswith("outlet_"):
        return "Branch " + label(str(key)[7:])
    words = str(key).replace("_", " ").split()
    return " ".join(word.upper() if word.lower() in {"id", "sr", "sim", "iccid", "crr", "drr", "vat", "aed"} else word.capitalize() for word in words)


def clean_text(value):
    # Spreadsheet XML cannot contain control characters; preserve useful line breaks.
    return re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", str(value))


def csv_value(value):
    if value is None:
        return ""
    if isinstance(value, (int, float, Decimal)) and not isinstance(value, bool):
        return value
    if isinstance(value, datetime):
        return value.isoformat(sep=" ") + (" UTC" if value.tzinfo is None else "")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, (dict, list, tuple)):
        value = json.dumps(value, ensure_ascii=False, default=str)
    text = clean_text(value)
    # Check whitespace too, since spreadsheet apps trim it before formula evaluation.
    return "'" + text if text.lstrip().startswith(("=", "+", "-", "@")) or text.startswith(("\t", "\r")) else text


def _typed(value, key):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.astimezone(timezone.utc).replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (dict, list, tuple)):
        return clean_text(json.dumps(value, ensure_ascii=False, default=str))
    if key in TEXT_KEYS or key.endswith("_id") or key.endswith("_number") or key.endswith("_serial"):
        return clean_text(value)
    if isinstance(value, str):
        if key.endswith("_at") or key in {"date", "at", "tele_last_attempt", "welcome_last_attempt", "assignment_effective_at"}:
            try:
                return _typed(datetime.fromisoformat(value.replace("Z", "+00:00")), key)
            except ValueError:
                pass
        if key in MONEY_KEYS and re.fullmatch(r"[+-]?\d+(?:\.\d+)?", value.strip()):
            return float(value)
        return clean_text(value)
    return value


def style_sheet(sheet, columns, *, machine_headers=False):
    sheet.sheet_view.showGridLines = False
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.print_title_rows = "1:1"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape" if len(columns) > 5 else "portrait"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth, sheet.page_setup.fitToHeight = 1, 0
    sheet.oddHeader.center.text = f"&B{BRAND}&B"
    sheet.oddHeader.center.size = 10
    sheet.oddFooter.left.text = "Confidential | Authorized export"
    sheet.oddFooter.right.text = "Page &P of &N"
    thin = Side(style="hair", color="DDD5DF")
    for cell in sheet[1]:
        cell.value = clean_text(cell.value)
        cell.data_type = "s"
        cell.font = Font(name="Calibri", color="FFFFFF", bold=True, size=11)
        cell.fill = PatternFill("solid", fgColor=BURGUNDY)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 34
    for number, column in enumerate(columns, 1):
        letter = get_column_letter(number)
        key = str(column).lower().replace(" ", "_")
        lengths = [len(str(cell.value or "")) for cell in list(sheet.iter_cols(min_col=number, max_col=number))[0][:200]]
        width = max(16, min(48, max(lengths or [16]) + 2))
        if any(word in key for word in ("note", "remark", "detail", "reason", "value", "benefit")):
            width = max(width, 38)
        sheet.column_dimensions[letter].width = width
        for cells in sheet.iter_rows(min_row=2, min_col=number, max_col=number):
            cell = cells[0]
            cell.font = Font(name="Calibri", size=10, color="202A3D")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7F2F8")
            if isinstance(cell.value, str):
                # An XLSX text cell is never executable, without altering recorded values.
                cell.data_type = "s"
            if key in TEXT_KEYS or key.endswith(("_id", "_number", "_serial")):
                cell.number_format = "@"
            elif key in MONEY_KEYS or "aed" in key:
                cell.number_format = '#,##0.00;[Red](#,##0.00);0.00'
            elif isinstance(cell.value, datetime):
                cell.number_format = "dd mmm yyyy hh:mm"
            elif isinstance(cell.value, date):
                cell.number_format = "dd mmm yyyy"
    for number in range(2, sheet.max_row + 1):
        line_count = max((max(1, len(str(cell.value or "")) // max(12, int(sheet.column_dimensions[cell.column_letter].width)) + 1, str(cell.value or "").count("\n") + 1) for cell in sheet[number]), default=1)
        sheet.row_dimensions[number].height = min(409, max(25, 15 * line_count))


def add_metadata(book, title, metadata=None, notes=None):
    sheet = book.create_sheet("Export details")
    sheet.sheet_view.showGridLines = False
    sheet.column_dimensions["A"].width = 26
    sheet.column_dimensions["B"].width = 90
    sheet.merge_cells("A1:B1")
    sheet["A1"] = BRAND
    sheet["A1"].font = Font(name="Calibri", size=18, bold=True, color=BURGUNDY)
    sheet.row_dimensions[1].height = 32
    sheet.append(["Report", title])
    sheet.append(["Generated UTC", datetime.now(timezone.utc).strftime("%d %b %Y %H:%M UTC")])
    sheet.append(["Date filter timezone", "Asia/Dubai (UAE)"])
    for key, value in (metadata or {}).items():
        sheet.append([label(key), value if value is not None and str(value).strip() else MISSING])
    for note in notes or []:
        sheet.append(["Note", note])
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.value = _typed(cell.value, "metadata")
            if isinstance(cell.value, str):
                cell.data_type = "s"
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            cell.font = Font(name="Calibri", size=11, bold=cell.column == 1, color="202A3D")
        sheet.row_dimensions[row[0].row].height = min(409, max(30, 15 * (len(str(row[1].value or "")) // 85 + 1)))
    if LOGO.exists():
        image = ExcelImage(str(LOGO))
        image.width, image.height = 250, 71
        sheet.add_image(image, f"A{sheet.max_row + 3}")
    return sheet


def workbook_bytes(columns, rows, title, *, metadata=None, machine_headers=False, notes=None, sheet_name="Records"):
    book = Workbook()
    sheet = book.active
    sheet.title = sheet_name[:31]
    sheet.append(list(columns) if machine_headers else [label(key) for key in columns])
    for row in rows:
        values = [row.get(key, "") for key in columns] if isinstance(row, dict) else row
        sheet.append([_typed(value, str(columns[index]).lower().replace(" ", "_")) for index, value in enumerate(values)])
    style_sheet(sheet, columns, machine_headers=machine_headers)
    add_metadata(book, title, {"records": len(rows), **(metadata or {})}, notes)
    book.properties.creator = BRAND
    book.properties.title = title
    book.properties.subject = "Authorized operational records"
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def table_download(columns, rows, title, format="csv", *, metadata=None, machine_headers=True, notes=None):
    if format == "xlsx":
        data, media = workbook_bytes(columns, rows, title, metadata=metadata, machine_headers=machine_headers, notes=notes), XLSX_MEDIA
    elif format == "pdf":
        records = [row if isinstance(row, dict) else dict(zip(columns, row)) for row in rows]
        subtitle = " | ".join(f"{label(key)}: {value}" for key, value in (metadata or {}).items() if value is not None)
        data, media = table_pdf_bytes(title, records, subtitle, columns=columns), "application/pdf"
    else:
        output = io.StringIO(newline="")
        writer = csv.writer(output)
        writer.writerow([csv_value(column) for column in (columns if machine_headers else [label(key) for key in columns])])
        for row in rows:
            values = [row.get(key, "") for key in columns] if isinstance(row, dict) else row
            writer.writerow([csv_value(value) for value in values])
        data, media = output.getvalue().encode("utf-8-sig"), "text/csv; charset=utf-8"
    return Response(data, media_type=media, headers={"Content-Disposition": f'attachment; filename="{filename(title, format)}"', "Cache-Control": "private, no-store", "X-Content-Type-Options": "nosniff"})


def _pdf_styles():
    from reportlab.lib import colors
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    regular, bold = "Helvetica", "Helvetica-Bold"
    paths = [Path("/usr/share/fonts/truetype/dejavu"), Path("C:/Windows/Fonts")]
    for path in paths:
        normal_file = path / ("arial.ttf" if path.name == "Fonts" else "DejaVuSans.ttf")
        bold_file = path / ("arialbd.ttf" if path.name == "Fonts" else "DejaVuSans-Bold.ttf")
        if normal_file.exists() and bold_file.exists():
            if "QanawatSans" not in pdfmetrics.getRegisteredFontNames():
                pdfmetrics.registerFont(TTFont("QanawatSans", str(normal_file)))
                pdfmetrics.registerFont(TTFont("QanawatSans-Bold", str(bold_file)))
                pdfmetrics.registerFontFamily("QanawatSans", normal="QanawatSans", bold="QanawatSans-Bold")
            regular, bold = "QanawatSans", "QanawatSans-Bold"
            break
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = bold if style.name in {"Title", "Heading1", "Heading2", "Heading3"} else regular
        style.textColor = colors.HexColor("#202A3D")
    styles["Normal"].fontSize, styles["Normal"].leading = 8, 11
    styles["Title"].fontSize, styles["Title"].leading = 17, 22
    styles["Title"].textColor = colors.HexColor("#761B3A")
    styles["Heading2"].fontSize, styles["Heading2"].leading = 10, 14
    return styles


def _paragraph(value, style):
    from reportlab.platypus import Paragraph
    text = MISSING if value is None or str(value).strip() == "" else str(value)
    if re.search(r"[\u0600-\u06ff]", text):
        try:
            import arabic_reshaper
            from bidi.algorithm import get_display
            text = get_display(arabic_reshaper.reshape(text))
        except ImportError:
            pass
    return Paragraph(escape(clean_text(text)).replace("\n", "<br/>"), style)


def _pdf_header(story, title, subtitle, styles, width):
    from reportlab.platypus import Image, Spacer, Table, TableStyle
    if LOGO.exists():
        logo = Image(str(LOGO), width=156, height=44.5)
        heading = Table([[logo, _paragraph(title, styles["Title"])]], colWidths=[170, width - 170])
        heading.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        story.append(heading)
    else:
        story.append(_paragraph(BRAND + " | " + title, styles["Title"]))
    if subtitle:
        story.append(_paragraph(subtitle, styles["Normal"]))
    story.append(Spacer(1, 12))


def _pdf_footer(canvas, doc):
    width, _ = doc.pagesize
    canvas.setFont("Helvetica", 7)
    canvas.setFillColorRGB(0.35, 0.38, 0.44)
    canvas.drawString(doc.leftMargin, 18, BRAND + " | Confidential")
    canvas.drawRightString(width - doc.rightMargin, 18, f"Page {doc.page}")


def table_pdf_bytes(title, rows, subtitle="", columns=None):
    """Every field is retained; wide tables use readable repeated column groups."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.platypus import SimpleDocTemplate, PageBreak, Spacer, Table, TableStyle
    styles, output = _pdf_styles(), io.BytesIO()
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=30, rightMargin=30, topMargin=25, bottomMargin=32, title=title, author=BRAND)
    keys = list(columns) if columns is not None else list(dict.fromkeys(key for row in rows for key in row))
    width, story = doc.width, []
    _pdf_header(story, title, subtitle, styles, width)
    if not rows or not keys:
        story.append(_paragraph("No records match the selected filters.", styles["Normal"]))
    else:
        anchor = next((key for key in ("reference", "request_id", "sr_number", "employee_id", "name", "iccid", "serial", "sale_id", "Field") if key in keys), keys[0])
        remaining = [key for key in keys if key != anchor]
        bands = [[anchor, *remaining[start:start + 5]] for start in range(0, len(remaining), 5)] or [[anchor]]
        for band_number, band in enumerate(bands, 1):
            if band_number > 1:
                story.append(PageBreak())
                _pdf_header(story, title, subtitle, styles, width)
            if len(bands) > 1:
                story += [_paragraph(f"Fields {band_number} of {len(bands)} | {len(rows)} records", styles["Heading2"]), Spacer(1, 6)]
            cells = [[_paragraph(label(key), styles["Heading2"]) for key in band]]
            for row in rows:
                cells.append([_paragraph(csv_value(row.get(key)) if not isinstance(row.get(key), str) else row.get(key), styles["Normal"]) for key in band])
            table = Table(cells, repeatRows=1, colWidths=[width / len(band)] * len(band), hAlign="LEFT", splitInRow=1)
            table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDE1E9")), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8F5F9")]), ("VALIGN", (0, 0), (-1, -1), "TOP"), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6), ("LINEBELOW", (0, 0), (-1, 0), 1, colors.HexColor("#761B3A")), ("LINEBELOW", (0, 1), (-1, -1), 0.3, colors.HexColor("#DDD5DF"))]))
            story.append(table)
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return output.getvalue()


def receipt_pdf_bytes(title, sections, subtitle=""):
    """Compact sectioned receipt, retaining all recorded and missing information."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.platypus import SimpleDocTemplate, Spacer, Table, TableStyle
    styles, output = _pdf_styles(), io.BytesIO()
    styles["Normal"].fontSize, styles["Normal"].leading = 7.5, 10
    styles["Heading2"].spaceBefore = 5
    styles["Heading2"].spaceAfter = 4
    doc = SimpleDocTemplate(output, pagesize=A4, leftMargin=30, rightMargin=30, topMargin=24, bottomMargin=30, title=title, author=BRAND)
    story = []
    _pdf_header(story, title, subtitle, styles, doc.width)
    for section in sections:
        story.append(_paragraph(section.get("title", "Details"), styles["Heading2"]))
        fields = section.get("fields") or []
        cells = []
        for start in range(0, len(fields), 2):
            pair = fields[start:start + 2]
            cells.append([_paragraph(field.get("label", "Details"), styles["Normal"]) if item % 2 == 0 else _paragraph(field.get("value"), styles["Normal"]) for field in pair for item in range(2)] + (["", ""] if len(pair) == 1 else []))
        if cells:
            table = Table(cells, colWidths=[doc.width * .19, doc.width * .31, doc.width * .19, doc.width * .31], hAlign="LEFT", splitInRow=1)
            table.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "TOP"), ("BACKGROUND", (0, 0), (0, -1), colors.HexColor("#F4EDF3")), ("BACKGROUND", (2, 0), (2, -1), colors.HexColor("#F4EDF3")), ("GRID", (0, 0), (-1, -1), .3, colors.HexColor("#DED4DE")), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            story.append(table)
        story.append(Spacer(1, 4))
    doc.build(story, onFirstPage=_pdf_footer, onLaterPages=_pdf_footer)
    return output.getvalue()
