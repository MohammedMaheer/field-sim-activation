"""Atomic administrator import of SIM stock from an Excel workbook."""

import base64
from io import BytesIO
from zipfile import ZipFile
from fastapi import APIRouter, Depends, HTTPException, Request
from fastapi.responses import StreamingResponse
from openpyxl import Workbook, load_workbook
from pydantic import BaseModel, Field
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError
from .db import Branch, Movement, Outlet, Role, Sim, get_db
from .security import principal
from .services import audit
from .stock_categories import validate_category

router = APIRouter(prefix="/api/inventory", tags=["Inventory"])


def admin(db, user):
    if db.get(Role, user.role_id).name != "Administrator":
        raise HTTPException(403, "Only administrators can import SIM stock")


@router.get("/bulk-template")
def template(user=Depends(principal), db=Depends(get_db)):
    admin(db, user)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "SIM stock"
    sheet.append(["ICCID", "SIM Serial", "SIM Type", "Business Category"])
    for col, width in {"A": 28, "B": 25, "C": 20, "D": 25}.items():
        sheet.column_dimensions[col].width = width
    data = BytesIO()
    workbook.save(data)
    data.seek(0)
    return StreamingResponse(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                             headers={"Content-Disposition": 'attachment; filename="sim-stock-template.xlsx"'})


class ImportBody(BaseModel):
    branch_id: str
    content_base64: str = Field(min_length=20, max_length=1500000)
    reason: str = Field(min_length=5, max_length=300)


@router.post("/bulk", status_code=201)
def import_stock(body: ImportBody, request: Request, user=Depends(principal), db=Depends(get_db)):
    admin(db, user)
    branch = db.get(Branch, body.branch_id)
    if not branch:
        raise HTTPException(422, "Select a valid branch")
    from .branch_lifecycle import active_branch
    active_branch(db, branch.id)
    if len(body.reason.strip()) < 5:
        raise HTTPException(422, "Enter a meaningful stock reason")
    outlet = db.scalar(select(Outlet).where(Outlet.branch_id == branch.id).order_by(Outlet.created_at))
    if not outlet:
        raise HTTPException(409, "This branch has no stock assignment")
    try:
        binary = base64.b64decode(body.content_base64, validate=True)
        if len(binary) > 1000000 or binary[:2] != b"PK":
            raise ValueError()
        with ZipFile(BytesIO(binary)) as archive:
            parts = archive.infolist()
            if len(parts) > 30 or sum(part.file_size for part in parts) > 8000000:
                raise ValueError()
        book = load_workbook(BytesIO(binary), read_only=True, data_only=True)
        sheet = book.active
        rows = sheet.iter_rows(values_only=True)
        header = next(rows, None)
        if not header or [str(v or "").strip().casefold() for v in header[:3]] != ["iccid", "sim serial", "sim type"]:
            raise ValueError()
        if len(header) > 3 and str(header[3] or "").strip().casefold() != "business category":
            raise ValueError()
        entries = []
        for index, row in enumerate(rows, 2):
            if index > 10000:
                raise HTTPException(422, "Keep the import under 10,000 rows")
            if not any(value is not None and str(value).strip() for value in row):
                continue
            if len(entries) >= 500 or len(row) < 3 or any(not isinstance(v, str) for v in row[:3]):
                raise HTTPException(422, f"Row {index}: use text values; maximum 500 SIMs per file")
            iccid, serial, sim_type = [v.strip() for v in row[:3]]
            if not (3 <= len(iccid) <= 60 and 3 <= len(serial) <= 60 and sim_type in {"Physical", "eSIM"}):
                raise HTTPException(422, f"Row {index}: check ICCID, serial and SIM type")
            category = str(row[3] or "Not recorded").strip() if len(row) > 3 else "Not recorded"
            validate_category(category)
            entries.append((iccid, serial, sim_type, category))
        book.close()
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(422, "Use the Excel template with valid SIM text values") from exc
    if not entries:
        raise HTTPException(422, "The Excel file contains no SIM stock")
    iccids = [entry[0] for entry in entries]
    serials = [entry[1] for entry in entries]
    if len(iccids) != len(set(iccids)) or len(serials) != len(set(serials)):
        raise HTTPException(409, "The Excel file contains duplicate SIM identifiers")
    if db.scalar(select(Sim.id).where((Sim.iccid.in_(iccids)) | (Sim.serial.in_(serials))).limit(1)):
        raise HTTPException(409, "A SIM identifier already exists")
    try:
        created = []
        for iccid, serial, sim_type, category in entries:
            sim = Sim(iccid=iccid, serial=serial, sim_type=sim_type, business_category=category, status="AVAILABLE", outlet_id=outlet.id)
            db.add(sim)
            db.flush()
            db.add(Movement(sim_id=sim.id, user_id=user.id, old_status="WAREHOUSE", new_status="AVAILABLE",
                            reason=body.reason.strip()))
            audit(db, user, "SIM Imported", sim.id, new={"iccid": iccid, "serial": serial,
                  "sim_type": sim_type, "business_category": category, "branch": branch.name}, reason=body.reason.strip(), request=request)
            created.append(sim.id)
        db.commit()
    except IntegrityError:
        db.rollback()
        raise HTTPException(409, "A SIM identifier already exists") from None
    return {"imported": len(created), "branch": branch.name}
